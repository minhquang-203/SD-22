# -*- coding: utf-8 -*-
"""
Viết lại PhamMinhQuang_Testcase_DATN.xlsx theo nghiệp vụ (use case + chính sách cửa hàng).

Nguyên tắc:
- Cột Kết quả mong đợi = những gì cửa hàng/khách PHẢI thấy (không mô tả hàm Java).
- Không lấy expected từ code. Nếu code lệch thì case Fail — đó là phát hiện lỗi.
- Phạm vi: đúng 8 chức năng đã làm, không test toàn hệ thống.
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "PhamMinhQuang_Testcase_DATN.xlsx"

THIN = Border(
    left=Side(style="thin", color="8EA9DB"),
    right=Side(style="thin", color="8EA9DB"),
    top=Side(style="thin", color="8EA9DB"),
    bottom=Side(style="thin", color="8EA9DB"),
)
TITLE_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
OK_FILL = PatternFill("solid", fgColor="E2EFDA")
ROW_A = PatternFill("solid", fgColor="FFFFFF")
ROW_B = PatternFill("solid", fgColor="D6E3F8")
TITLE_FONT = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Times New Roman", size=11)
BOLD = Font(name="Times New Roman", size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ---------------------------------------------------------------------------
# Quy tắc nghiệp vụ (oracle). Nguồn: đặc tả use case UC-06/08/09/10/12/17/18/21/25
# + chính sách cửa hàng công bố (đổi trả 7 ngày, seal còn nguyên).
# ---------------------------------------------------------------------------
RULES = [
    ("NV-SALE-01", "UC-25.1", "Đợt giảm giá",
     "Giá khuyến mãi chỉ bán khi đợt đang bật và nằm trong khoảng ngày bắt đầu 00:00 đến hết ngày kết thúc. Trước giờ mở bán khách vẫn thấy giá gốc."),
    ("NV-SALE-02", "UC-25.1", "Đợt giảm giá",
     "Mức giảm 1–100%. Tiền sau giảm làm tròn đến đồng. Mọi kênh (website, giỏ, quầy, đặt hàng) cùng một giá."),
    ("NV-SALE-03", "UC-25.1", "Đợt giảm giá",
     "Được gắn sản phẩm khi đợt còn Sắp diễn ra để chuẩn bị giờ mở bán. Đợt ngừng / hết hạn / đã ẩn thì không thêm hàng, khách về giá gốc."),
    ("NV-SALE-04", "UC-25.1", "Đợt giảm giá",
     "Hai đợt chồng lên cùng món: khách trả mức thấp nhất (có lợi hơn). Đơn đã đặt giữ nguyên giá lúc đặt."),
    ("NV-SALE-05", "UC-25.1", "Đợt giảm giá",
     "Sản phẩm đang giảm giá vẫn hiện ở danh mục và trang Khuyến mãi với giá sau giảm — không ẩn khỏi cửa hàng."),
    ("NV-SALE-06", "UC-25.2/25.3", "Đợt giảm giá",
     "Đang chạy: được ngừng ngay (giá gốc trở lại), không sửa mức % để tránh giá nhảy khi khách đang xem. Sắp diễn ra được sửa. Hết hạn không sửa."),
    ("NV-SALE-07", "UC-25", "Đợt giảm giá",
     "Chỉ Quản lý / Chủ cửa hàng tạo–sửa–ngừng–xóa đợt. Nhân viên quầy và khách không."),
    ("NV-VC-01", "UC-25.4", "Phiếu giảm giá",
     "Ba loại: phần trăm, tiền mặt, miễn ship. Miễn ship chỉ cho đơn giao hàng, không dùng tại quầy."),
    ("NV-VC-02", "UC-25.4", "Phiếu giảm giá",
     "Mã chỉ dùng khi còn hạn, còn lượt (lượt là của cả cửa hàng), đủ giá trị đơn tối thiểu, đang bật. Nhập mã không phân biệt hoa thường."),
    ("NV-VC-03", "UC-25.4", "Phiếu giảm giá",
     "Phần trăm tính trên tiền hàng (sau giá đợt giảm nếu có), có trần; không tính trên phí ship. Tiền mặt không vượt tiền hàng. Miễn ship đúng phí vận chuyển, có trần."),
    ("NV-VC-04", "UC-25.4", "Phiếu giảm giá",
     "Đặt thành công trừ 1 lượt. Hủy đơn hoặc thanh toán cổng thất bại hoàn lại 1 lượt. Hai đơn cùng lúc khi còn 1 lượt: chỉ 1 đơn dùng được."),
    ("NV-VC-05", "UC-25.4", "Phiếu giảm giá",
     "Mã cá nhân chỉ khách được gán mới dùng được, kể cả khi người khác biết mã. Tại quầy phải chọn đúng khách."),
    ("NV-VC-06", "UC-25", "Phiếu giảm giá",
     "Chỉ Quản lý / Chủ phát hành mã. Khách phải đăng nhập mới xem/áp mã của mình. Khách vãng lai không dùng voucher."),
    ("NV-OC-01", "UC-06", "Đặt hàng",
     "Khách đăng nhập đặt từ giỏ tài khoản. Khách vãng lai đặt bằng họ tên, SĐT, email, địa chỉ; nhận mã tra cứu qua email; không gắn tài khoản."),
    ("NV-OC-02", "UC-06/14", "Đặt hàng",
     "Mỗi dòng 1–10 sản phẩm và không vượt tồn. Hết hàng / ngừng bán / giá không hợp lệ thì không tạo đơn, kho không đổi."),
    ("NV-OC-03", "UC-06", "Đặt hàng",
     "Giá dòng lấy tại lúc đặt (đã gồm đợt giảm). Voucher trừ bước sau. Thành tiền = hàng − giảm + ship, không âm."),
    ("NV-OC-04", "UC-06/08", "Đặt hàng",
     "Bắt buộc địa chỉ, phường/xã, SĐT hợp lệ, COD hoặc VNPay. Đơn 0 đồng không đi VNPay."),
    ("NV-OC-05", "UC-06", "Đặt hàng",
     "Bấm Đặt hàng hai lần / mất mạng gửi lại: chỉ một đơn, kho không trừ hai lần. Không lấy hàng trong giỏ của người khác."),
    ("NV-OC-06", "UC-06", "Đặt hàng",
     "COD: xóa dòng đã mua khỏi giỏ, đơn hiện ngay. VNPay: giữ giỏ và ẩn đơn đến khi trả thành công trong 15 phút."),
    ("NV-PAY-01", "UC-08", "Thanh toán",
     "COD: ghi nhận đã thu khi giao thành công, không thu sớm lúc xác nhận."),
    ("NV-PAY-02", "UC-08", "Thanh toán",
     "VNPay: sai chữ ký không đổi tiền/kho. Lệch số tiền không thành công, hủy đơn, trả kho và hoàn lượt mã. Hết 15 phút chưa trả: hủy mềm, trả kho, hoàn mã."),
    ("NV-PAY-03", "UC-08", "Thanh toán",
     "Cổng báo thành công sau khi đơn đã hủy: không hồi sinh đơn; cửa hàng được báo để đối soát/hoàn tay nếu tiền đã trừ tài khoản khách."),
    ("NV-PAY-04", "UC-08", "Thanh toán",
     "Ghi nhận thanh toán cổng chỉ một lần: callback/IPN lặp không nhân tiền, kho, mail. Đơn đã thu hoặc đã kết thúc không tạo lần trả mới trên cùng hóa đơn."),
    ("NV-TC-01", "UC-21", "Trạng thái đơn",
     "Luồng giao: Chờ xác nhận → Đã xác nhận → Đang chuẩn bị → Đang giao → Đã giao. Không lùi bước. Không nhảy từ Chờ xác nhận sang Đã giao. Đơn trả hàng không đi PATCH luồng giao thường."),
    ("NV-TC-02", "UC-17/21", "Trạng thái đơn",
     "Đơn VNPay chưa thanh toán không được xác nhận/giao (được hủy). Khách tự hủy trước khi giao; đang giao thì không tự hủy trên website."),
    ("NV-TC-03", "UC-17/21", "Trạng thái đơn",
     "Hủy đơn đã trừ kho thì trả đúng lô. Lịch sử trạng thái chỉ ghi thêm, không sửa/xóa. Người thao tác là nhân viên đang đăng nhập."),
    ("NV-TC-04", "UC-09", "Trạng thái đơn",
     "Khách chỉ xem đơn của mình. Đơn VNPay chưa trả không hiện trong danh sách. Khách vãng lai tra cứu bằng liên kết/email; sai thông tin không lộ đơn; chặn dò nhiều lần."),
    ("NV-TB-01", "UC-12", "Thông báo",
     "COD: báo cửa hàng (và mail khách) lúc đặt xong. VNPay: cùng lúc đó sau khi thanh toán thành công, không báo khi còn chờ/hết hạn."),
    ("NV-TB-02", "UC-12/18", "Thông báo",
     "Đổi trạng thái giao: chuông khách. Duyệt/từ chối/hoàn trả hàng: thông báo riêng, không lẫn với “cập nhật đơn”. Mail lỗi không làm hỏng đơn."),
    ("NV-TB-03", "UC-12", "Thông báo",
     "Khách không đọc/xóa thông báo của người khác. Chuông quản trị không lẫn thông báo khách."),
    ("NV-SR-01", "UC-09/17", "Lọc đơn",
     "Lọc kết hợp điều kiện (và), không hiện đơn VNPay chờ trả. Khách vãng lai tra cứu được; nhân viên mới được tìm trên trang quản trị."),
    ("NV-RT-01", "UC-10", "Trả hàng",
     "Chỉ đơn online đã giao, trong 7 ngày kể từ khi nhận, đúng chủ đơn, chưa có yêu cầu đang mở. Tối thiểu 2 ảnh (tối đa 6, đúng ảnh). COD bắt buộc số tài khoản nhận tiền hợp lệ."),
    ("NV-RT-02", "UC-10/18", "Trả hàng",
     "Trả cả đơn, không trả từng món. Mỗi đơn một yêu cầu. Bị từ chối thì không gửi lại trên hệ thống — liên hệ CSKH."),
    ("NV-RT-03", "UC-18", "Trả hàng",
     "Cửa hàng duyệt/từ chối khi đang chờ. Sau duyệt khách tạo vận đơn hoàn. Nhận hàng mới được phân lô tốt/lỗi và nhập kho. Hoàn tiền chỉ sau khi đã nhận và kiểm hàng."),
    ("NV-RT-04", "UC-18", "Trả hàng",
     "Đã nhận hàng và nhập kho thì không được từ chối hoàn tiền (khách không mất cả hàng lẫn tiền). Phải hoàn hoặc trả hàng lại cho khách."),
    ("NV-RT-05", "UC-18", "Trả hàng",
     "Số hoàn > 0 và không vượt số đã thu. VNPay hoàn về kênh đã trả; COD về số tài khoản khách khai. Nhân viên không tự ý hoàn 0 đồng hoặc vượt số thu."),
    ("NV-RT-06", "UC-10", "Trả hàng",
     "Khách không xem/sửa yêu cầu của người khác (kể cả số tài khoản). Khách vãng lai không gửi trả hàng trên hệ thống."),
    ("NV-RT-07", "UC-18/21", "Trả hàng",
     "Khách duyệt xong nhưng không gửi hàng quá thời hạn cửa hàng quy định: hủy yêu cầu, đơn trở về Đã giao — không kẹt Trả hàng."),
    ("NV-RT-08", "UC-18", "Trả hàng",
     "Hàng lỗi không nhập lại tồn bán được. Số lượng phân lô phải khớp số đã bán. Hai lần gửi yêu cầu cùng lúc trên một đơn chỉ nhận một yêu cầu."),
]

# Case sẽ Fail trên bản code hiện tại nếu chạy đúng oracle nghiệp vụ (không phải fail vì test sai).
KNOWN_DRIFT = [
    ("SALE15", "NV-SALE-03", "Sắp diễn ra phải gắn được sản phẩm để chuẩn bị mở bán."),
    ("SALE19", "NV-SALE-05", "Sản phẩm đang KM phải hiện ở danh mục với giá sau giảm, không bị ẩn."),
    ("SALE25", "NV-SALE-07", "API tạo/sửa đợt phải chặn nhân viên quầy, không chỉ chặn giao diện."),
    ("VC16", "NV-VC-02", "Khách nhập sale20 hay SALE20 đều phải áp được mã."),
    ("VC20", "NV-VC-01", "Danh sách mã tại quầy không được lệch số trang vì vẫn đếm mã miễn ship."),
    ("VC24", "NV-VC-06", "API phát hành voucher phải chặn nhân viên quầy giống giao diện."),
    ("OC16", "NV-OC-02", "Gọi API vượt 10 sản phẩm/dòng phải bị từ chối như quy định giỏ hàng."),
    ("OC30", "NV-OC-04", "SĐT sai định dạng phải bị từ chối cả phía máy chủ, không chỉ form."),
    ("OC39", "NV-OC-05", "Hai lần đặt không kèm mã chống trùng vẫn không được tạo hai đơn."),
    ("TC03", "NV-TC-01", "Chỉ được bỏ bước Đang chuẩn bị khi đã có vận đơn; không coi mọi bước “số lớn hơn” là hợp lệ."),
    ("TC43", "NV-TC-03", "Không cho sửa/xóa lịch sử trạng thái."),
    ("TC44", "NV-TC-03", "Nhân viên trên chứng từ phải là người đang đăng nhập."),
    ("RT05", "NV-RT-01", "Hết 7 ngày kể từ khi nhận thì không gửi được yêu cầu trả."),
    ("RT07", "NV-RT-01", "Máy chủ cũng phải giới hạn số ảnh và chỉ nhận file ảnh."),
    ("RT15", "NV-RT-08", "Đơn cũ không có lô vẫn phải chọn tốt/lỗi, không nhập hết vào tồn bán."),
    ("RT17", "NV-RT-05", "Hoàn 0 đồng hoặc lớn hơn số đã thu phải bị từ chối."),
    ("RT18", "NV-RT-04", "Không cho từ chối hoàn sau khi đã nhận hàng và nhập kho."),
    ("RT20", "NV-RT-06", "Khách A không được đọc danh sách/STK trả hàng của khách B."),
    ("RT26", "NV-RT-07", "Yêu cầu treo không gửi hàng phải được đóng, đơn về Đã giao."),
    ("RT27", "NV-RT-08", "Hai yêu cầu trả cùng lúc trên một đơn chỉ được nhận một."),
]


def nv_tag(tid: str) -> str:
    p = re.match(r"[A-Z]+", tid)
    key = p.group(0) if p else ""
    table = {
        "SALE": "NV-SALE",
        "VC": "NV-VC",
        "OC": "NV-OC",
        "PAY": "NV-PAY",
        "TC": "NV-TC",
        "TB": "NV-TB",
        "SR": "NV-SR",
        "RT": "NV-RT",
    }
    return table.get(key, "NV")


def clean_tech(old: str, tid: str) -> str:
    t = (old or "").replace("\n", " ").strip()
    t = re.sub(r"^NV-[A-Z]+(?:-\d+)?\s*·\s*", "", t)
    t = re.sub(r"\s*\((?:R|D)\d+[^)]*\)", "", t)
    t = re.sub(r"\s+[RD]\d+\b", "", t)
    t = re.sub(r"\s*\(Gap[^)]*\)", "", t, flags=re.I)
    t = re.sub(r"\s*\(catalog\)", "", t, flags=re.I)
    t = t.replace(" + gap page", "").replace(" (Gap)", "").replace("/Gap", "")
    t = t.replace("FE vs BE gap", "phân quyền").replace("Gap ", "")
    t = re.sub(r"\s{2,}", " ", t).strip(" ·")
    return f"{nv_tag(tid)} · {t}" if t else nv_tag(tid)


def sanitize_exp(text: str) -> str:
    """Gỡ dấu vết đọc code; giữ ý nghiệp vụ nếu câu vốn đã rõ."""
    if not text:
        return text
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"(?i)\s*[—\-–]?\s*ghi nhận[^.\n]*", "", t)
    t = re.sub(r"(?i)không fail[^.\n]*", "", t)
    t = re.sub(r"(?i)expected\s*=\s*hành vi code hiện tại[^.\n]*", "", t)
    t = re.sub(r"(?i)đúng code hiện tại[^.\n]*", "", t)
    t = re.sub(r"(?i)\(GAP[^)]*\)", "", t)
    t = re.sub(r"(?i)GAP[^\n.]*", "", t)
    t = t.replace("@NotBlank ", "").replace("@NotEmpty ", "").replace("@NotNull ", "")
    t = t.replace("@Min(1) ", "").replace("@Email ", "")
    t = t.replace("ApiException ", "")
    t = re.sub(r"\bPaginationUtil\b", "sắp xếp mặc định", t)
    t = re.sub(r"\s{2,}", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip(" \n/·")


# title, loai, phan, exp  — None = giữ / chỉ làm sạch
PATCH: dict[str, dict] = {}


def P(tid, *, exp, title=None, loai=None, phan=None, tech=None):
    PATCH[tid] = {"exp": exp, "title": title, "loai": loai, "phan": phan, "tech": tech}


# ===== Đợt giảm giá =====
P("SALE01", exp="Tạo được đợt. Trong ngày bắt đầu, đợt ở trạng thái Đang chạy. Trên danh sách có dòng mới, số đợt đang chạy tăng 1. Thông báo tạo thành công.")
P("SALE02", exp="Đợt bắt đầu ngày mai ở Sắp diễn ra. Trang Khuyến mãi và giá bán chưa đổi — khách vẫn mua giá gốc đến giờ mở bán.")
P("SALE03", exp="Giảm 1% và 100% đều tạo được. 1%: giá mới = giá gốc trừ 1% (làm tròn đến đồng). 100%: khách không phải trả tiền hàng (0đ). Ngoài 1–100% thuộc SALE04.")
P("SALE04", exp="Thiếu mã/tên/ngày/%, % ngoài 1–100, ngày bắt đầu trong quá khứ, ngày kết thúc trước ngày bắt đầu, trùng mã, mã quá ngắn/có dấu cách: không tạo đợt, báo rõ từng lỗi.")
P("SALE05", exp="Sửa tên, mức %, ngày kết thúc của đợt Sắp diễn ra thành công. Khách vẫn chưa mua giá khuyến mãi.")
P("SALE06", title="Không sửa mức giảm khi đợt đang chạy hoặc đã hết hạn",
  exp="Đợt đang chạy hoặc đã ngừng nhưng còn trong hạn: không đổi mức % (tránh giá nhảy khi khách đang xem); được ngừng/bật lại. Đợt hết hạn: không sửa. Dữ liệu cũ giữ nguyên.")
P("SALE07", exp="Sau khi ngừng, đợt không còn áp. Mọi kênh (Khuyến mãi, chi tiết SP, giỏ, quầy, đặt hàng) về giá gốc ngay. Ngừng lần nữa: báo đợt đã ngừng.")
P("SALE08", exp="Bật lại khi còn hạn: giá khuyến mãi trở lại mọi kênh. Bật khi đang bật: báo đã áp dụng.")
P("SALE09", exp="Hủy hộp thoại xóa: đợt còn trên danh sách. Xác nhận xóa: đợt biến khỏi danh sách quản trị và không còn bán giá KM; lịch sử/đơn cũ không bị xóa cứng.")
P("SALE10", exp="Thao tác trên đợt không tồn tại hoặc đã xóa: báo không tìm thấy. Không đổi dữ liệu khác.")
P("SALE11", exp="Từng tab (đang chạy / sắp diễn ra / hết hạn / ngừng) đúng tập đợt còn hiệu lực. Tìm theo mã hoặc tên, không phân biệt hoa thường. Không khớp: danh sách trống, không lỗi. Trạng thái lạ: báo không hợp lệ.")
P("SALE12", exp="Có đủ số liệu tổng / đang chạy / sắp diễn ra / đã kết thúc. Đổi sắp xếp và sang trang không trùng dòng. Bấm một dòng vào đúng trang chi tiết.")
P("SALE13", exp="Trang chi tiết hiện mã, mức %, khoảng ngày, số sản phẩm, giá cũ/mới. Đổi lưới/danh sách và tìm sản phẩm hoạt động. Đợt không có hàng: hướng dẫn thêm sản phẩm.")
P("SALE14", exp="Gắn biến thể khi đợt đang chạy: giá sau giảm = giá gốc trừ đúng %, làm tròn đến đồng (ví dụ 199.000 giảm 15% → 169.150). Dòng xuất hiện trên chi tiết đợt. Nhân viên không tự nhập giá sau giảm — hệ thống tính.")
P("SALE15", title="Gắn sản phẩm khi Sắp diễn ra được; ngừng/hết hạn thì không",
  loai="Negative",
  exp="Đợt Sắp diễn ra: gắn được sản phẩm để chuẩn bị giờ mở bán (chưa bán giá KM). Đợt đã ngừng hoặc hết hạn: không gắn, báo rõ. Không tạo dòng hàng trên đợt không còn chạy.")
P("SALE16", exp="Gắn trùng đúng biến thể trong cùng đợt: báo đã có, không nhân dòng. Biến thể khác của cùng sản phẩm thì gắn được.")
P("SALE17", exp="Đợt đang chạy: không gỡ sản phẩm (tránh giá nhảy). Đợt Sắp diễn ra: gỡ được, khách chưa từng mua giá KM. Hủy xác nhận: không gỡ. Dòng không thuộc đợt: báo không hợp lệ.")
P("SALE18", exp="Trang Khuyến mãi chỉ hiện sản phẩm thuộc đợt đang chạy, có giá sau giảm và giá gốc gạch ngang. Hàng sắp diễn ra / ngừng / hết hạn không có. Khách chưa đăng nhập vẫn xem được. Không có đợt chạy: trang trống rõ ràng.")
P("SALE19", title="Sản phẩm đang khuyến mãi vẫn hiện ở danh mục với giá sau giảm",
  loai="Positive",
  exp="Danh mục / tìm kiếm vẫn hiện sản phẩm đang giảm giá, giá bán = giá sau giảm (không ẩn sản phẩm khỏi cửa hàng). Trang Khuyến mãi cũng hiện. Đơn đã đặt trước đó không bị đổi giá.")
P("SALE20", exp="Chi tiết sản phẩm, giỏ, quầy và lúc đặt hàng cùng một giá sau giảm. Không giảm hai lần trên cùng đợt. Voucher (nếu có) trừ trên tổng tiền hàng sau giá đợt.")
P("SALE21", exp="Hai đợt cùng món cùng lúc: khách trả mức thấp hơn. Tắt đợt sâu hơn thì còn mức còn lại. Đơn đã đặt giữ giá lúc đặt.")
P("SALE22", exp="Trong ngày kết thúc, đến hết ngày vẫn bán giá KM. Sang ngày hôm sau: hết khuyến mãi, mọi kênh về giá gốc.")
P("SALE23", exp="Ngừng hoặc hết hạn hoặc ẩn đợt: không còn giá KM trên Khuyến mãi, danh mục, chi tiết, giỏ.")
P("SALE24", exp="Biến thể ngừng bán hoặc sản phẩm đã ẩn: không hiện khuyến mãi, không bán giá KM.")
P("SALE25", title="Chỉ Quản lý/Chủ được tạo–sửa đợt giảm giá",
  loai="Negative",
  exp="Nhân viên quầy không vào trang quản trị đợt và không tạo/sửa/xóa qua API. Khách chỉ xem trang Khuyến mãi, không phát hành đợt.")
P("SALE26", exp="Thiếu mã sản phẩm, sản phẩm không tồn tại, đợt không tồn tại: không gắn hàng, báo rõ, không tạo dòng dở.")

# ===== Voucher =====
P("VC01", exp="Tạo phiếu % bắt đầu hôm nay: đang hoạt động, mã viết hoa, hiện trên danh sách, số phiếu đang chạy tăng.")
P("VC02", exp="Tạo được phiếu tiền mặt và miễn ship. Miễn ship không dùng số % khi tính; tiền mặt không bắt buộc trần %. Cả hai đang hoạt động, đúng loại trên danh sách.")
P("VC03", exp="Phiếu bắt đầu ngày mai ở Sắp diễn ra. Khách và quầy chưa thấy mã; đặt hàng với mã này bị từ chối vì chưa tới hạn.")
P("VC04", exp="Thiếu trường, %/tiền/số lượt không hợp lệ, đơn tối thiểu âm, ngày sai, trùng mã, mã quá ngắn: không tạo, báo rõ.")
P("VC05", exp="Sửa phiếu Sắp diễn ra thành công. Mã không đổi lúc sửa. Khách vẫn chưa dùng được.")
P("VC06", title="Không đổi giá trị phiếu khi đang chạy hoặc đã hết hạn",
  exp="Phiếu đang chạy / đã ngừng nhưng còn hạn: không đổi mức giảm (đơn đang đặt không bị đổi giữa chừng). Hết hạn: không sửa. Dữ liệu cũ giữ nguyên.")
P("VC07", exp="Ngừng: mã biến khỏi chỗ chọn lúc đặt hàng và quầy; nhập tay bị từ chối. Bật lại khi còn hạn: hiện lại. Ngừng lần nữa: báo đã ngừng.")
P("VC08", exp="Hủy xóa: phiếu còn. Xác nhận: ẩn khỏi tìm kiếm, không áp được; đơn cũ đã dùng mã không bị xóa.")
P("VC09", exp="Lọc theo từ khóa, trạng thái, loại ra đúng tập phiếu còn hiệu lực. Loại không hợp lệ: báo lỗi. Không khớp: danh sách trống.")
P("VC10", exp="Thống kê đang chạy / lượt đã dùng / tiền đã giảm / sắp hết hạn. Sang trang và sắp xếp không trùng dòng. Số lần đã dùng khớp số đơn đã áp mã.")
P("VC11", exp="Giảm 20% trên 200.000đ tiền hàng → 40.000đ. Có trần 30.000đ → chỉ giảm 30.000đ. Thành toán = hàng − giảm + ship. Không tính % trên phí ship.")
P("VC12", exp="Phiếu 30.000đ trên đơn 100.000đ → giảm 30.000đ. Đơn 20.000đ → chỉ giảm 20.000đ (không vượt tiền hàng). Ship cộng sau.")
P("VC13", exp="Miễn ship đúng bằng phí vận chuyển; nếu có trần thì khách trả phần vượt trần. Đơn ghi đã dùng mã, trừ 1 lượt.")
P("VC14", exp="Tại quầy không hiện và không áp mã miễn ship. Báo mã chỉ dành cho đơn giao hàng. Không trừ lượt.")
P("VC15", exp="Tiền hàng 1 đồng dưới mức tối thiểu: không áp, báo chưa đạt. Đúng mức (sau giá đợt, chưa cộng ship): áp được.")
P("VC16", title="Không áp mã chưa tới hạn / hết hạn / ngừng / hết lượt / sai mã; hoa thường vẫn đúng mã",
  loai="Negative",
  exp="Chưa tới hạn, hết hạn, đang ngừng, hết lượt, mã không tồn tại: không giảm, không trừ lượt, không hiện trong danh sách mã khả dụng. Khách gõ sale20 hay SALE20 đều áp được cùng mã.")
P("VC17", exp="Đặt thành công: còn lại giảm 1 lượt, đơn gắn mã. Hủy đơn hoặc VNPay thất bại: hoàn 1 lượt.")
P("VC18", exp="Còn 1 lượt, hai đơn cùng lúc: đúng một đơn dùng được, đơn kia báo hết lượt. Không để lượt âm.")
P("VC19", exp="Khách đăng nhập thấy mã đang chạy. Đơn dưới tối thiểu thì không chọn được. Tìm, chọn, bỏ chọn cập nhật lại số tiền.")
P("VC20", title="Quầy chỉ hiện phiếu % và tiền mặt; số trang khớp mã thật sự dùng được",
  loai="Positive",
  exp="Quầy không liệt kê miễn ship. % và tiền mặt chọn được, tính giá đúng. Tổng số / phân trang khớp số mã thực sự dùng tại quầy (không để trang trống vì vẫn đếm mã miễn ship).")
P("VC21", exp="Màn đặt hàng, quầy và hóa đơn hiện đúng mã và số tiền đã giảm. Thống kê lượt dùng / tiền giảm tăng.")
P("VC22", exp="Hết ngày kết thúc vẫn dùng được trong ngày đó. Sang ngày hôm sau: hết hạn, biến khỏi danh sách, nhập tay bị từ chối.")
P("VC23", exp="Ngừng, ẩn, hết hạn hoặc hết lượt: mã không còn ở web lẫn quầy; nhập tay bị từ chối.")
P("VC24", title="Chỉ khách đăng nhập mới xem/áp mã; chỉ Quản lý phát hành",
  loai="Negative",
  exp="Chưa đăng nhập không lấy danh sách mã và không đặt kèm voucher. Khách không tạo/sửa phiếu. Nhân viên quầy không vào trang phát hành và không gọi API tạo phiếu. Quản lý/Chủ làm được.")
P("VC25", exp="Phiếu cá nhân gán khách A: A dùng được. B thấy hoặc biết mã vẫn không áp được, báo mã dành cho khách được chỉ định.")
P("VC26", exp="Gán theo khoảng điểm: khách đủ điểm nhận mã, khách thiếu điểm không. Khi khách vừa đủ điểm, được gán. Phiếu công khai không gắn khoảng điểm — báo không hợp lệ.")
P("VC27", exp="Biết mã cá nhân của người khác vẫn không dùng được trên web. Quầy chưa chọn khách hoặc chọn sai khách: không áp.")
P("VC28", exp="Quầy: chưa chọn khách thì mã cá nhân không dùng. Chọn đúng khách A thì áp được. Chọn B thì không. Mã công khai vẫn hiện.")
P("VC29", exp="Phiếu công khai (kể cả phiếu cũ chưa gắn phạm vi) mọi khách đăng nhập dùng được khi đủ điều kiện.")

# ===== Checkout =====
P("OC01", exp="Tạo đơn giao hàng, chờ xác nhận. Đúng sản phẩm/số lượng, trừ kho theo lô gần hết hạn trước, xóa dòng đã chọn khỏi giỏ, báo cửa hàng và gửi mail khách. Thành tiền = hàng − giảm + ship.")
P("OC02", exp="Chỉ 2 dòng đã chọn vào đơn; dòng còn lại vẫn trong giỏ.")
P("OC03", exp="Đơn chờ thanh toán cổng, trừ kho và trừ lượt mã (nếu có), giỏ chưa xóa, có đường dẫn trả tiền. Đơn chưa hiện trong danh sách cho đến khi trả thành công.")
P("OC04", exp="Giỏ trống: không tạo đơn, báo giỏ trống.")
P("OC05", exp="Không chọn sản phẩm: không tạo đơn, yêu cầu chọn ít nhất một món.")
P("OC06", exp="Dòng không còn trong giỏ: không tạo đơn, báo sản phẩm không hợp lệ.")
P("OC07", exp="Sản phẩm đã ẩn: không tạo đơn, kho/giỏ/mã không đổi.")
P("OC08", exp="Biến thể ngừng bán: không tạo đơn, báo không còn bán.")
P("OC09", exp="Hết hàng: không tạo đơn, kho không âm.")
P("OC10", exp="Giá bán không hợp lệ (≤ 0, không KM): không tạo đơn.")
P("OC11", exp="Số lượng 1, còn hàng: đặt thành công.")
P("OC12", exp="Số lượng 0: không đặt, báo số lượng không hợp lệ.")
P("OC13", exp="Số lượng âm: bị từ chối ngay trên giỏ và lúc đặt.")
P("OC14", exp="Mua đúng bằng tồn: thành công, tồn về 0.")
P("OC15", exp="Mua hơn tồn 1: không đặt, báo hết hàng.")
P("OC16", title="Mỗi dòng tối đa 10 sản phẩm — kể cả gọi thẳng API",
  loai="Negative",
  exp="Số lượng 11–… bị từ chối (cùng quy định giỏ: tối đa 10/dòng và không vượt tồn). Không tạo đơn.")
P("OC17", exp="Giá đã đổi sau khi thêm giỏ: đơn lấy giá tại lúc đặt, không lấy giá cũ trên giỏ.")
P("OC18", exp="Có đợt giảm đang chạy: đơn giá dòng = giá sau giảm.")
P("OC19", exp="Mã % hợp lệ: trừ đúng số tiền, giảm 1 lượt.")
P("OC20", exp="Mã tiền mặt hợp lệ: trừ đúng, không vượt tiền hàng, giảm 1 lượt.")
P("OC21", exp="Mã miễn ship: giảm đúng phí vận chuyển (theo trần nếu có), thành tiền đúng.")
P("OC22", exp="Mã hết hạn: không đặt kèm giảm giá, báo mã không dùng được.")
P("OC23", exp="Mã chưa tới ngày: không áp, báo chưa tới hạn.")
P("OC24", exp="Hết lượt: không áp, báo hết lượt.")
P("OC25", exp="Tiền hàng dưới tối thiểu của mã: không áp mã.")
P("OC26", exp="Vừa đợt giảm vừa voucher: giá dòng đã KM, rồi mới trừ mã trên tổng hàng. Không nhân % hai lần.")
P("OC27", title="Cùng khách dùng lại mã khi cửa hàng còn lượt",
  loai="Positive",
  exp="Lượt là của cả cửa hàng (không giới hạn 1 lần/khách). Cùng người dùng lại được đến khi hết lượt toàn hệ thống; hết lượt thì đơn sau bị từ chối.")
P("OC28", exp="Thiếu địa chỉ giao: không tạo đơn, yêu cầu nhập địa chỉ.")
P("OC29", exp="Thiếu phường/xã: không tạo đơn, yêu cầu chọn địa chỉ giao đủ.")
P("OC30", exp="SĐT không phải số điện thoại Việt Nam hợp lệ: không tạo đơn. Chặn cả trên form và khi gửi thẳng máy chủ.")
P("OC31", exp="Để trống tên/SĐT lúc đặt: lấy từ hồ sơ khách đang đăng nhập.")
P("OC32", exp="Thiếu cách thanh toán: không tạo đơn.")
P("OC33", exp="Cách thanh toán khác COD/VNPay: từ chối.")
P("OC34", exp="Đơn 0 đồng không được chọn VNPay; phải COD (hoặc không tạo đơn VNPay).")
P("OC35", exp="Chưa đăng nhập không đặt theo giỏ tài khoản. Khách vãng lai đi luồng đặt vãng lai (OC53).")
P("OC36", exp="Không đặt được hàng trong giỏ của người khác.")
P("OC37", exp="Bấm Đặt hàng hai lần nhanh: chỉ gửi/nhận một đơn.")
P("OC38", exp="Hai lần gửi cùng khóa chống trùng: chỉ một đơn; lần sau nhận lại đơn đã tạo, kho không trừ hai lần.")
P("OC39", title="Hai lần đặt liên tiếp không được tạo hai đơn",
  loai="Negative",
  exp="Khách bấm/gửi hai lần gần như cùng lúc (dù không gửi khóa chống trùng): vẫn chỉ một đơn, kho không trừ hai lần.")
P("OC40", exp="Còn đơn VNPay chưa trả rồi đặt VNPay mới: đơn cũ hủy, đơn mới được tạo; kho/mã xử lý đúng một đơn còn hiệu lực.")
P("OC41", exp="Hai khách tranh món cuối: một thành công, một hết hàng; tồn không âm.")
P("OC42", exp="Hai khách tranh lượt mã cuối: một dùng được, một hết lượt.")
P("OC43", exp="Lỗi giữa chừng lúc đặt: không có đơn dở, kho và lượt mã như trước.")
P("OC44", exp="Tổng tồn nhìn đủ nhưng lô thực tế không đủ xuất: không tạo đơn, kho không đổi.")
P("OC45", exp="COD thành công: dòng đã mua biến khỏi giỏ, dòng khác còn.")
P("OC46", exp="VNPay chưa trả: giỏ giữ nguyên đến khi thanh toán thành công.")
P("OC47", exp="Kho giảm đúng số đã bán, ưu tiên lô gần hết hạn; hóa đơn ghi rõ lấy từ lô nào.")
P("OC48", exp="Mỗi món một dòng; thành tiền dòng = đơn giá × số lượng; tổng hàng cộng đúng; khách trả = hàng − giảm + ship, không âm.")
P("OC49", exp="COD sinh giao dịch chờ thu đúng số khách phải trả.")
P("OC50", exp="Đặt COD xong: cửa hàng có đơn mới, khách có mail nếu có email. Không gửi được mail thì đơn vẫn tạo, không mất đơn.")
P("OC51", exp="Hết 15 phút chưa trả VNPay: đơn hủy, trả kho, hoàn lượt mã.")
P("OC52", exp="Hãng vận chuyển không báo phí: vẫn đặt được, phí ship dùng mức dự phòng cửa hàng đã đặt, khách thấy phí trước khi xác nhận.")
P("OC53", exp="Khách vãng lai đặt COD: đơn không gắn tài khoản, có email và mã tra cứu, trừ kho, mail về email vừa nhập, cửa hàng thấy đơn mới.")
P("OC54", exp="Vãng lai chọn VNPay: có đường dẫn trả tiền, đơn ẩn đến khi trả thành công, giỏ trên máy chưa xóa.")
P("OC55", exp="Thiếu sản phẩm / email / tên / SĐT / địa chỉ / cách trả: không tạo đơn, kho không đổi, báo thiếu gì.")
P("OC56", exp="Khách vãng lai không áp voucher (cần tài khoản để kiểm tra lượt/mã cá nhân). Tiền giảm = 0, không trừ lượt.")
P("OC57", exp="Vãng lai gửi lại cùng khóa chống trùng: một đơn, kho trừ một lần.")
P("OC58", exp="Không bắt đăng nhập. Sau COD: hướng dẫn kiểm tra email và liên kết tra cứu đơn. Không hiện nút đơn của tài khoản.")

# ===== Thanh toán (chỉ nghiệp vụ; tạo đơn/timeout/guest/POS để sheet OC/TC) =====
P("PAY01", exp="Giao COD thành công: mới ghi nhận đã thu. Không tạo thêm giao dịch. Số tiền không đổi.")
P("PAY02", exp="Chưa giao xong: vẫn chờ thu, chưa ghi đã nhận tiền.")
P("PAY03", exp="Ghi nhận đã thu lần hai: vẫn một giao dịch đã thu, không nhân tiền.")
P("PAY04", exp="Trả thành công đúng hạn, đúng tiền, đúng chữ ký: đơn hiện, giỏ trừ món, mail + báo cửa hàng. Không hủy đơn.")
P("PAY05", exp="Chữ ký cổng sai: không đổi thanh toán, đơn, kho, mã, giỏ. Khách thấy thanh toán không hợp lệ.")
P("PAY06", exp="Số tiền lệch: không ghi nhận đã thu; giao dịch thất bại; hủy đơn chờ trả, trả kho và hoàn 1 lượt mã; giữ bản ghi để đối soát.")
P("PAY07", exp="Hủy/thất bại trên cổng: đơn hủy, trả kho, hoàn mã, giỏ nguyên, báo thất bại. Đơn không hiện như đã đặt xong.")
P("PAY08", exp="Cổng báo thành công sau khi đơn đã hủy: đơn không sống lại, không trừ kho lần nữa. Cửa hàng được báo để đối soát/hoàn tay nếu tiền khách đã bị trừ.")
P("PAY09", exp="Cổng gửi lại báo thành công khi đã ghi nhận: không trừ giỏ/mail/kho hai lần; vẫn một giao dịch đã thu.")

# ===== Trạng thái =====
P("TC01", exp="Chờ xác nhận → Đã xác nhận: đổi trạng thái, có mốc lịch sử, khách nhận thông báo.")
P("TC02", exp="Đi đủ Chờ xác nhận → Đã xác nhận → Đang chuẩn bị → Đang giao → Đã giao. Timeline đủ mốc, không lỗi.")
P("TC03", title="Đã xác nhận sang Đang giao chỉ khi đã có vận đơn — không nhảy cóc tùy ý",
  loai="Negative",
  exp="Được chuyển Đã xác nhận → Đang giao nếu đã tạo vận đơn (bỏ bước chuẩn bị khi hàng đã xuất). Chưa có vận đơn thì không nhảy. Không dùng quy tắc “số thứ tự lớn hơn là được”. Từ Chờ xác nhận sang Đã giao thuộc TC49.")
P("TC04", exp="Không lùi Đang giao về Đã xác nhận. Trạng thái giữ nguyên, báo không chuyển được.")
P("TC05", exp="Đơn đã giao không chuyển lại Đang giao.")
P("TC06", exp="Đơn đã hủy không mở lại Chờ xác nhận.")
P("TC07", exp="Không chuyển Đã giao → Trả hàng bằng nút trạng thái thường. Phải đi luồng duyệt trả hàng.")
P("TC08", exp="Chọn đúng trạng thái hiện tại: không ghi thêm mốc giả, báo không có gì để chuyển.")
P("TC09", exp="Cửa hàng hủy khi đang giao (theo quy định nội bộ, thường kèm hủy vận đơn): đơn hủy, trả kho nếu đã trừ.")
P("TC10", exp="Khách không tự hủy khi đơn đang giao. Báo liên hệ cửa hàng.")
P("TC11", exp="Khách hủy khi còn chuẩn bị hàng: đơn hủy, trả kho, hoàn mã; nếu đã thu thì có yêu cầu hoàn tiền.")
P("TC12", exp="VNPay chưa trả: không xác nhận/giao. Giữ nguyên.")
P("TC13", exp="VNPay đã trả: xác nhận được.")
P("TC14", exp="VNPay chưa trả vẫn hủy được (hủy chờ thanh toán).")
P("TC15", exp="Khách xem đúng đơn của mình: trạng thái, vận chuyển, mốc mới nhất. Không thấy lô nội bộ.")
P("TC16", exp="Khách A không xem đơn khách B (không lộ là đơn có tồn tại).")
P("TC17", exp="Chưa đăng nhập không xem “đơn của tôi”.")
P("TC18", exp="Đơn không tồn tại: báo không tìm thấy.")
P("TC19", exp="Trạng thái không thuộc quy trình: từ chối, không đổi đơn.")
P("TC20", exp="Thiếu trạng thái mới: không đổi đơn.")
P("TC21", exp="Hai nhân viên đổi một đơn cùng lúc: không ra trạng thái mâu thuẫn; không nhân thông báo/kho.")
P("TC22", exp="Khách đặt hai lần cùng khóa chống trùng: một đơn chờ xác nhận, kho trừ một lần.")
P("TC23", exp="Gửi trùng lệnh chuyển trạng thái: lần hai không ghi lịch sử trùng.")
P("TC24", exp="Đơn quầy thanh toán tiền mặt xong: đã giao/hoàn tất, trừ kho, cộng điểm nếu cửa hàng áp dụng.")
P("TC25", exp="Sau vài lần chuyển: đủ mốc theo thời gian, có người thao tác và ghi chú.")
P("TC26", exp="Thời điểm mốc khớp lúc thao tác, không trống, tăng dần.")
P("TC27", exp="VNPay chưa tới 15 phút: chưa tự hủy, chưa hiện như đơn đã đặt xong.")
P("TC28", exp="Quá 15 phút chưa trả: tự hủy, thanh toán thất bại, trả kho/mã.")
P("TC29", exp="Đổi trạng thái giao: khách có chuông và thấy nhãn mới, không cần tải lại nếu đang mở trang.")
P("TC30", exp="Duyệt trả hàng: đơn sang Trả hàng theo luồng trả, khách được báo trả hàng (không lẫn thông báo giao thường).")
P("TC31", exp="Giao xong: có chuông/cập nhật trạng thái; không bắt buộc mail “đã giao” nếu chính sách chỉ báo trên website.")
P("TC32", exp="Hủy đơn đã trừ kho: cộng đúng số về đúng lô.")
P("TC33", exp="Hãng báo đã giao: đơn Đã giao; COD chuyển sang đã thu.")
P("TC34", exp="Hãng hủy vận đơn: đơn hủy, trả kho.")
P("TC35", exp="Mã vận chuyển lạ: không đổi trạng thái, hệ thống không sập.")
P("TC36", exp="Khách thấy “Đã giao”, không thấy mã kỹ thuật.")
P("TC37", exp="Khi nhân viên đổi trạng thái, trang khách cập nhật nhãn mới.")
P("TC38", exp="Đặt đơn mới vẫn tạo bình thường (chờ xác nhận, trừ kho, mail).")
P("TC39", exp="Đổi trạng thái giao COD đến đã giao: mới ghi đã thu, số tiền khớp.")
P("TC40", exp="Chỉ xác nhận đơn: kho không trừ lần hai (đã trừ lúc đặt).")
P("TC41", exp="Mỗi lần đổi trạng thái: đúng một thông báo, không spam.")
P("TC42", title="Nhân viên được từ chối đơn chờ xác nhận, phải ghi lý do",
  loai="Positive",
  exp="Nhân viên từ chối đơn chờ xác nhận được (đúng UC xử lý đơn), khách thấy lý do. Khách không gọi được thao tác này.")
P("TC43", title="Không sửa hoặc xóa lịch sử trạng thái",
  loai="Negative",
  exp="Không có cách hợp lệ để sửa/xóa mốc đã ghi. Lịch sử chỉ thêm, phục vụ đối soát.")
P("TC44", title="Người thao tác trên đơn là nhân viên đang đăng nhập",
  loai="Negative",
  exp="Không nhận “hộ” mã nhân viên khác từ dữ liệu gửi lên. Chứng từ ghi đúng người đang đăng nhập.")
P("TC45", exp="Danh sách đơn của tôi không chứa đơn VNPay chưa trả.")
P("TC46", exp="Liên kết tra cứu đúng: xem được đơn, không cần đăng nhập. Sai/thiếu mã: không ra đơn. Chỉ xem, không hủy/trả.")
P("TC47", exp="Đúng mã + đúng email: xem được. Sai email hoặc sai mã: cùng thông báo không tìm thấy — không lộ đơn.")
P("TC48", exp="Tra cứu công khai quá số lần cho phép trong khoảng thời gian ngắn: bị tạm khóa, không dò được đơn.")

# ===== Thông báo =====
P("TB01", exp="Đặt COD: cửa hàng có chuông đơn mới; khách nhận mail hóa đơn nếu có email, không cần chuông “đặt thành công” trên tài khoản. Đơn chờ xác nhận.")
P("TB02", exp="VNPay chưa trả hoặc hết hạn: không báo đơn mới. Trả thành công: cửa hàng có đơn mới, khách có mail như đơn đã đặt xong.")
P("TB03", exp="Nhân viên đổi trạng thái giao: khách có một thông báo cập nhật đơn, bấm vào xem được đơn. Cửa hàng không nhận thêm chuông “cập nhật” kiểu khách.")
P("TB04", exp="Hủy đơn: khách thấy đơn đã bị hủy. Một sự kiện → một thông báo.")
P("TB05", exp="Hãng đổi trạng thái giao: khách vẫn có thông báo như khi nhân viên đổi tay.")
P("TB06", exp="Duyệt trả hàng: không gửi thông báo “cập nhật đơn giao”. Có thông báo trả hàng riêng.")
P("TB07", exp="Đơn quầy không gắn khách: không chuông tài khoản khách, không mail đặt hàng web.")
P("TB08", exp="Chuông khách: số chưa đọc đúng. Đọc một cái giảm 1. Đọc tất cả về 0. Đọc lại không lỗi.")
P("TB09", exp="Chuông quản trị chỉ việc cửa hàng (đơn mới…). Bấm vào đi đúng hóa đơn. Không lẫn thông báo của khách.")
P("TB10", exp="Khách A không đánh dấu đã đọc thông báo của B hoặc của cửa hàng. A không thấy danh sách của B.")
P("TB11", exp="Bấm thông báo cập nhật đơn: vào trang đơn của tôi / tra cứu, đánh dấu đã đọc.")
P("TB12", exp="Thông báo mới hiện trên chuông ngay. Nhiều sự kiện sát nhau không bắn toast liên tục gây khó chịu.")
P("TB13", exp="Chưa đăng nhập không lấy chuông. Khách không lấy chuông quản trị. Nhân viên không lấy chuông khách.")
P("TB14", exp="Hơn 99 chưa đọc: hiện 99+. Số thật vẫn đếm đủ. Không còn chưa đọc: không badge.")
P("TB15", exp="Không gửi được mail / khách không có email: đơn vẫn tạo, chuông cửa hàng vẫn có.")
P("TB16", exp="Chuông chỉ hiện một số thông báo mới nhất cho gọn; số chưa đọc vẫn đếm đủ toàn bộ.")
P("TB17", title="VNPay thành công phải báo đơn đã thanh toán/đặt xong",
  loai="Positive",
  exp="Sau khi khách trả VNPay thành công, cửa hàng nhận thông báo hiểu được là có đơn mới đã thanh toán. Không im lặng. Không bắt buộc một tên kỹ thuật cụ thể.")
P("TB18", exp="Một lần đổi trạng thái → một thông báo khách. Gửi trùng lệnh không tạo thêm.")
P("TB19", exp="Vãng lai đặt COD: mail về email lúc đặt, không chuông tài khoản.")

# ===== Lọc / tìm =====
P("SR01", exp="Nhân viên tìm theo mã đơn, tên khách, tên nhân viên phụ trách hoặc cách thanh toán ra đúng tập, không lẫn đơn không khớp.")
P("SR02", exp="Lọc đơn giao hàng / đơn quầy / tất cả ra đúng loại. Không lẫn web vào quầy.")
P("SR03", exp="Lọc từng trạng thái ra đúng nhóm. Giá trị trạng thái không thuộc quy trình: báo không hợp lệ hoặc danh sách trống, không sập.")
P("SR04", exp="Cùng một ngày: ra đơn trong ngày đó. Từ ngày sau đến ngày trước: không ra hoặc báo khoảng ngày không hợp lệ. Giờ đầu/cuối ngày vẫn tính đúng trong ngày.")
P("SR05", exp="Kết hợp từ khóa + loại đơn + trạng thái + khoảng ngày: chỉ ra đơn thỏa tất cả điều kiện (và), không ra đơn chỉ khớp một phần.")
P("SR06", exp="Đơn VNPay chưa trả không hiện trên danh sách quản trị (chưa phải đơn đã đặt xong).")
P("SR07", exp="Tab Chờ xác nhận chỉ đơn chờ xác nhận; tab Tất cả không gồm đơn VNPay đang chờ trả. Số đếm trên tab khớp số dòng.")
P("SR08", exp="Mỗi trang đúng số dòng quy định, sắp theo ngày tạo mới trước, hai trang không trùng đơn.")
P("SR09", exp="Khách trên trang tra cứu lọc theo nhóm (đang xử lý / đang giao / đã giao / đã hủy / trả hàng) ra đúng đơn của mình.")
P("SR10", exp="Khách tìm theo mã đơn hoặc tên sản phẩm trong đơn: ra đúng đơn của mình chứa món đó.")
P("SR11", exp="Khách A không thấy và không mở đơn của B. Đơn VNPay chưa trả của A cũng ẩn.")
P("SR12", exp="Không khớp: danh sách trống, câu thông báo đúng ngữ cảnh (quản trị hoặc trang khách), không lỗi.")
P("SR13", exp="Đơn mới hoặc đổi trạng thái không xóa bộ lọc/từ khóa đang bật; đơn ra hoặc khỏi nhóm cho đúng.")
P("SR14", exp="Gõ online hay ONLINE cùng ra đơn giao hàng.")
P("SR15", exp="Đúng trạng thái nhưng ngoài khoảng ngày đang lọc: không ra.")
P("SR16", exp="Chưa đăng nhập và khách không tìm đơn kiểu quản trị. Nhân viên / Quản lý / Chủ được.")
P("SR17", exp="Trang tra cứu vãng lai không bắt đăng nhập. Đúng liên kết hoặc đúng mã + email thì xem được, chỉ xem (ẩn hủy / trả / đánh giá).")
P("SR18", exp="Sai email hoặc sai liên kết: không ra đơn, không lộ thông tin người khác.")
P("SR19", exp="Tra cứu công khai quá nhiều lần trong thời gian ngắn: tạm chặn, không dò được đơn.")

# ===== Trả hàng =====
P("RT01", exp="Đơn COD đã giao, còn hạn 7 ngày: gửi yêu cầu → duyệt → tạo vận đơn hoàn → nhận hàng → hoàn tiền về STK. Khách được báo từng bước. Chưa hoàn tiền lúc mới duyệt.")
P("RT02", exp="Đơn VNPay: không bắt STK. Hoàn về đúng kênh đã trả, không nhập mã giao dịch tay. Các bước duyệt/gửi hàng/nhận hàng như đơn COD.")
P("RT03", exp="Đơn đang giao, đã hủy, đơn quầy, hoặc chưa đăng nhập: không hiện nút trả hàng, gửi yêu cầu bị từ chối.")
P("RT04", exp="Đơn đã có yêu cầu: không gửi thêm, nút ẩn.")
P("RT05", title="Hết 7 ngày kể từ khi nhận thì không trả được",
  loai="Negative",
  exp="Đơn giao quá 7 ngày: ẩn nút trả hàng, gửi yêu cầu bị từ chối, báo hết thời hạn đổi trả. Không nhận yêu cầu.")
P("RT06", title="Bị từ chối thì không gửi lại trên hệ thống",
  loai="Negative",
  exp="Sau khi cửa hàng từ chối: không gửi yêu cầu mới trên website (đúng chính sách: liên hệ CSKH). Nút ẩn, báo đã có yêu cầu.")
P("RT07", exp="Dưới 2 ảnh: không gửi. 2–6 ảnh đúng loại: gửi được. File không phải ảnh hoặc vượt 6: bị từ chối cả trên form và khi gửi thẳng máy chủ.")
P("RT08", exp="COD: số tài khoản 6–20 chữ số (sau khi bỏ khoảng trắng) mới nhận. Ngắn/dài hơn: báo sai. Đơn VNPay không hỏi STK.")
P("RT09", exp="Thiếu lý do / ngân hàng / chủ tài khoản (COD): không gửi, báo ô bắt buộc.")
P("RT10", exp="Duyệt khi đang chờ: yêu cầu đã duyệt, đơn sang Trả hàng, khách được báo. Chưa tạo hoàn tiền. Duyệt lại: không hợp lệ.")
P("RT11", exp="Từ chối khi đang chờ: yêu cầu bị từ chối có lý do, đơn vẫn Đã giao, không hoàn kho, khách được báo. Không từ chối khi đã duyệt xong bước chờ.")
P("RT12", exp="Nhận hàng / tạo vận đơn / hoàn tiền khi chưa tới bước: bị từ chối, không nhảy cóc tới hoàn tất.")
P("RT13", exp="Sau duyệt: chọn ca lấy hàng hợp lệ thì có mã vận đơn hoàn. Ca hết hạn / thiếu địa chỉ: không tạo vận đơn. Khách không tạo vận đơn hộ người khác.")
P("RT14", exp="Hãng báo hàng đã về: chỉ cập nhật vận chuyển, chưa nhập kho và chưa hoàn tiền — nhân viên phải xác nhận đã nhận và kiểm hàng.")
P("RT15", title="Nhận hàng: phân đúng số lượng tốt/lỗi; hàng lỗi không vào tồn bán",
  exp="Tổng số mang về phải khớp số đã bán. Hàng tốt cộng tồn bán; hàng lỗi chỉ ghi lỗi, không bán lại. Thiếu/dư số lượng hoặc lô không thuộc đơn: không nhận. Đơn cũ không có lô: vẫn phải chọn tốt/lỗi, không mặc định nhập hết vào tồn bán.")
P("RT16", exp="COD: bắt buộc mã giao dịch chuyển khoản rồi mới hoàn tất. VNPay: hoàn qua cổng, tự có mã tham chiếu. Khách được báo hoàn thành công.")
P("RT17", title="Số hoàn phải > 0 và không vượt số đã thu",
  loai="Negative",
  exp="Hoàn 0 đồng hoặc lớn hơn số khách đã trả: bị từ chối, yêu cầu chưa hoàn tất. Nhân viên không tự sửa số vượt trần.")
P("RT18", title="Không từ chối hoàn tiền sau khi đã nhận hàng và nhập kho",
  loai="Negative",
  exp="Đã nhận hàng, đã nhập kho: không cho từ chối hoàn. Phải hoàn tiền hoặc trả hàng lại cho khách. Không để khách mất cả hàng lẫn tiền.")
P("RT19", exp="Hoàn hoặc từ chối lần hai khi đã xử lý: bị từ chối, không gọi cổng hai lần, không tạo thêm phiếu hoàn.")
P("RT20", title="Khách không xem yêu cầu trả hàng / STK của người khác",
  loai="Negative",
  exp="Khách A không lấy được danh sách hay số tài khoản của khách B. Xem/tạo trên đơn không phải của mình: không tìm thấy. Chưa đăng nhập không gửi trả hàng.")
P("RT21", title="Khách không hoàn tiền; nhân viên được giao thì hoàn được",
  loai="Negative",
  exp="Khách không gọi thao tác hoàn tiền. Nhân viên/Quản lý/Chủ xử lý đơn thì hoàn được (đúng UC-18).")
P("RT22", exp="Gửi yêu cầu: cửa hàng có chuông. Duyệt/từ chối/nhận/hoàn: khách có thông báo/mail đúng việc. Trang chi tiết hiện đủ các bước; hàng về shop chưa được coi như đã kiểm xong.")
P("RT23", exp="Lọc danh sách trả hàng / hoàn tiền theo trạng thái và mã đơn ra đúng nhóm. Không khớp: trống, không sập.")
P("RT24", title="Trả cả đơn; hoàn tiền hàng, phí ship theo lý do",
  loai="Positive",
  exp="Không chọn từng món — trả cả đơn (chính sách cửa hàng). Số hoàn mặc định = tiền hàng đã thu. Phí ship chỉ hoàn khi lỗi thuộc cửa hàng; khách đổi ý thì không hoàn ship.")
P("RT25", exp="Đơn vãng lai tra cứu được nhưng không gửi trả hàng trên website (cần tài khoản chủ đơn). Nút ẩn; gửi API bị từ chối.")
P("RT26", exp="Yêu cầu đã duyệt nhưng khách không gửi hàng quá thời hạn cửa hàng: đóng yêu cầu, đơn trở về Đã giao, chưa nhập kho, chưa hoàn tiền. Khách được báo. Không kẹt Trả hàng.")
P("RT27", exp="Hai lần gửi yêu cầu trả trên cùng đơn gần như cùng lúc: chỉ một yêu cầu được nhận, lần kia báo đã có yêu cầu.")
P("TC49", exp="Từ Chờ xác nhận không nhảy thẳng Đã giao. Đơn giữ Chờ xác nhận, khách không nhận thông báo đã giao.")

NEW_CASES = {
    "Cập nhật trạng thái đơn hàng": [
        [
            "TC49",
            "NV-TC-01 · Chuyển trạng thái (âm)",
            "Không nhảy từ Chờ xác nhận sang Đã giao",
            "Đơn online đang Chờ xác nhận, đã thanh toán nếu là VNPay",
            "Trạng thái mới = Đã giao",
            "Nhân viên chọn Đã giao (hoặc gửi yêu cầu đổi trạng thái)",
            "Bị từ chối. Đơn vẫn Chờ xác nhận. Khách không nhận thông báo đã giao. Kho không xử lý như đơn hoàn tất.",
            "Cao",
            "Negative",
            "Nghiệp vụ",
            "Auth: Bearer <token NV>\nPATCH /api/hoa-don/{id}/status\n{\"trangThai\":\"HOAN_THANH\"}",
        ]
    ],
    "Trả hàng hoàn tiền": [
        [
            "RT26",
            "NV-RT-07 · Chuyển trạng thái + thời hạn",
            "Khách không gửi hàng hoàn quá hạn — đóng yêu cầu, đơn về Đã giao",
            "Yêu cầu đã duyệt (hoặc đang hoàn hàng), quá số ngày cửa hàng quy định mà chưa nhận được hàng",
            "Hết hạn xử lý (ví dụ 7 ngày kể từ duyệt) ; chưa nhận hàng",
            "Tới hạn: chạy việc định kỳ hoặc thao tác đóng yêu cầu treo",
            "Yêu cầu đóng/hủy vì quá hạn. Đơn trở về Đã giao (không kẹt Trả hàng). Chưa nhập kho, chưa hoàn tiền. Khách được báo liên hệ nếu vẫn muốn trả.",
            "Cao",
            "Negative",
            "Nghiệp vụ",
            "Mốc thời hạn theo chính sách cửa hàng. Kiểm tra đơn không còn TRA_HANG.",
        ],
        [
            "RT27",
            "NV-RT-08 · Đồng thời",
            "Hai lần gửi yêu cầu trả trên cùng đơn — chỉ nhận một",
            "Đơn đã giao, còn hạn 7 ngày, chưa có yêu cầu",
            "Hai yêu cầu gần như cùng lúc, đủ ảnh",
            "Gửi hai lần liên tiếp / hai thiết bị",
            "Chỉ một yêu cầu được tạo. Lần kia báo đã có yêu cầu. Không hai hồ sơ trả trên một đơn.",
            "Cao",
            "Negative",
            "Nghiệp vụ",
            "POST /api/online/orders/{id}/tra-hang (×2, cùng khách)",
        ],
    ],
}


def style_cell(cell, value, *, header=False, title=False, center=False, fill=None, bold=False):
    cell.value = value
    cell.border = THIN
    cell.alignment = CENTER if center or header or title else LEFT
    if title:
        cell.font = TITLE_FONT
        cell.fill = TITLE_FILL
    elif header:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.number_format = "@"
    else:
        cell.font = BOLD if bold else DATA_FONT
        if fill is not None:
            cell.fill = fill
        cell.number_format = "@"


def write_intro(wb):
    # --- Phạm vi ---
    ws = wb.create_sheet("0. Phạm vi & nguyên tắc", 0)
    ws.merge_cells("A1:F1")
    style_cell(ws["A1"], "KIỂM THỬ CHỨC NĂNG ĐÃ LÀM — ORACLE LÀ NGHIỆP VỤ, KHÔNG PHẢI CODE", title=True)
    for c in range(1, 7):
        ws.cell(1, c).fill = TITLE_FILL
        ws.cell(1, c).border = THIN
    ws.row_dimensions[1].height = 28

    headers = ["Mục", "Nội dung"]
    for c, h in enumerate(headers, 1):
        style_cell(ws.cell(2, c), h, header=True)
    rows = [
        ("Phạm vi",
         "Chỉ 8 chức năng đã làm: đợt giảm giá, phiếu giảm giá, đặt hàng từ giỏ (kể cả vãng lai), thanh toán COD/VNPay, cập nhật trạng thái đơn, thông báo đơn, lọc/tra cứu đơn, trả hàng & hoàn tiền."),
        ("Không nằm trong bộ này",
         "Quiz loại da, chat AI, POS đầy đủ, sản phẩm/danh mục, nhập kho–NCC, đánh giá, phân quyền nhân sự, báo cáo. Không phải test toàn hệ thống."),
        ("Oracle (chuẩn đúng/sai)",
         "Cột Kết quả mong đợi viết từ use case UC-06, 08, 09, 10, 12, 17, 18, 21, 25 và chính sách cửa hàng (đổi trả 7 ngày, trả cả đơn, miễn ship chỉ đơn giao). Không chép hành vi hàm/API."),
        ("Cách đọc Pass/Fail",
         "Pass = hệ thống làm đúng nghiệp vụ. Fail = code lệch chính sách — giữ nguyên expected, ghi vào Kết quả thực tế. Không sửa expected cho khớp code."),
        ("Độ phủ thể hiện thế nào",
         "Mỗi case truy vết quy tắc NV-xx + use case (sheet 1–2). Đủ hướng: hợp lệ, không hợp lệ, biên, luồng trạng thái, bảng quyết định, phân quyền, chạy đồng thời. Gom substep cùng một lớp lỗi, không nhân case trùng."),
        ("Kỹ thuật",
         "Phân vùng tương đương, giá trị biên, bảng quyết định, chuyển trạng thái, đoán lỗi, cặp điều kiện, kiểm thử đồng thời. Ghi ở cột Kỹ thuật kèm mã quy tắc."),
        ("Cột Kết quả thực tế",
         "Trống để khi chạy điền Pass / Fail / Blocked / Skip. Vàng là ô cần điền."),
        ("Nếu chạy trên bản hiện tại",
         "Một số case ở sheet “Lệch nghiệp vụ” có thể Fail — đó là phát hiện, phục vụ sửa code hoặc giải trình bảo vệ, không phải test viết sai."),
    ]
    for i, (a, b) in enumerate(rows):
        fill = ROW_A if i % 2 == 0 else ROW_B
        style_cell(ws.cell(3 + i, 1), a, fill=fill, bold=True)
        style_cell(ws.cell(3 + i, 2), b, fill=fill)
        ws.row_dimensions[3 + i].height = 48
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    # --- Quy tắc ---
    ws2 = wb.create_sheet("1. Quy tắc nghiệp vụ", 1)
    ws2.merge_cells("A1:D1")
    style_cell(ws2["A1"], "QUY TẮC NGHIỆP VỤ — NGUỒN CỦA CỘT KẾT QUẢ MONG ĐỢI", title=True)
    for c in range(1, 5):
        ws2.cell(1, c).fill = TITLE_FILL
        ws2.cell(1, c).border = THIN
    ws2.row_dimensions[1].height = 28
    for c, h in enumerate(["Mã", "Use case", "Nhóm", "Quy tắc (oracle)"], 1):
        style_cell(ws2.cell(2, c), h, header=True)
    for i, (ma, uc, nhom, nd) in enumerate(RULES):
        fill = ROW_A if i % 2 == 0 else ROW_B
        style_cell(ws2.cell(3 + i, 1), ma, fill=fill, center=True, bold=True)
        style_cell(ws2.cell(3 + i, 2), uc, fill=fill, center=True)
        style_cell(ws2.cell(3 + i, 3), nhom, fill=fill, center=True)
        style_cell(ws2.cell(3 + i, 4), nd, fill=fill)
        ws2.row_dimensions[3 + i].height = 42
    ws2.column_dimensions["A"].width = 16
    ws2.column_dimensions["B"].width = 16
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 100
    ws2.auto_filter.ref = f"A2:D{2 + len(RULES)}"
    ws2.freeze_panes = "A3"
    ws2.sheet_view.showGridLines = False

    # --- Ma trận ---
    ws3 = wb.create_sheet("2. Ma trận phủ", 2)
    ws3.merge_cells("A1:I1")
    style_cell(ws3["A1"], "MA TRẬN PHỦ THEO CHỨC NĂNG ĐÃ LÀM (KHÔNG PHẢI TOÀN HỆ THỐNG)", title=True)
    for c in range(1, 10):
        ws3.cell(1, c).fill = TITLE_FILL
        ws3.cell(1, c).border = THIN
    ws3.row_dimensions[1].height = 28
    return ws3


def fill_matrix(ws3, counts: list[tuple]):
    heads = ["Chức năng", "Use case", "Số case", "Hợp lệ", "Không hợp lệ", "Biên",
             "Đồng thời / phân quyền", "Kỹ thuật chính", "Quy tắc"]
    for c, h in enumerate(heads, 1):
        style_cell(ws3.cell(2, c), h, header=True)
    widths = [28, 18, 12, 12, 14, 10, 22, 36, 22]
    for i, w in enumerate(widths, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for i, row in enumerate(counts):
        fill = ROW_A if i % 2 == 0 else ROW_B
        for c, val in enumerate(row, 1):
            style_cell(ws3.cell(3 + i, c), val, fill=fill, center=c not in (1, 8, 9))
        ws3.row_dimensions[3 + i].height = 36
    note_r = 4 + len(counts)
    ws3.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=9)
    style_cell(
        ws3.cell(note_r, 1),
        "Cách bảo vệ: độ phủ cao = mỗi quy tắc NV-xx có case hợp lệ và case phá quy tắc (biên/âm/đồng thời), "
        "truy vết được về use case — không phải số ID càng lớn càng tốt, không test ngoài phạm vi đã làm.",
        fill=WARN_FILL,
    )
    ws3.row_dimensions[note_r].height = 40
    ws3.freeze_panes = "A3"
    ws3.sheet_view.showGridLines = False


def write_drift(wb):
    ws = wb.create_sheet("3. Lệch nghiệp vụ (nếu Fail)", 3)
    ws.merge_cells("A1:D1")
    style_cell(
        ws["A1"],
        "CASE MÀ BẢN CODE HIỆN TẠI CÓ THỂ FAIL — EXPECTED VẪN GIỮ THEO NGHIỆP VỤ",
        title=True,
    )
    for c in range(1, 5):
        ws.cell(1, c).fill = TITLE_FILL
        ws.cell(1, c).border = THIN
    ws.row_dimensions[1].height = 28
    for c, h in enumerate(["ID", "Quy tắc", "Vì sao Fail là đúng", "Cách hiểu khi bảo vệ"], 1):
        style_cell(ws.cell(2, c), h, header=True)
    explain = "Đây là lỗi/lệch chính sách, không phải test sai. Nêu hướng sửa hoặc nhận hạn chế."
    for i, (tid, nv, why) in enumerate(KNOWN_DRIFT):
        fill = WARN_FILL
        style_cell(ws.cell(3 + i, 1), tid, fill=fill, center=True, bold=True)
        style_cell(ws.cell(3 + i, 2), nv, fill=fill, center=True)
        style_cell(ws.cell(3 + i, 3), why, fill=fill)
        style_cell(ws.cell(3 + i, 4), explain, fill=fill)
        ws.row_dimensions[3 + i].height = 36
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70
    ws.column_dimensions["D"].width = 55
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False


def classify_loai(loai: str) -> str:
    t = (loai or "").lower()
    if "gap" in t:
        t = t.replace("/gap", "").replace("gap", "negative").strip(" /")
        if not t:
            t = "negative"
        return t[:1].upper() + t[1:] if t else "Negative"
    return loai


def apply_row(ws, r, tid: str):
    p = PATCH.get(tid, {})
    old_tech = str(ws.cell(r, 2).value or "")
    old_exp = str(ws.cell(r, 7).value or "")
    old_loai = str(ws.cell(r, 10).value or "")

    tech = p.get("tech") or clean_tech(old_tech, tid)
    ws.cell(r, 2).value = tech
    ws.cell(r, 2).alignment = LEFT
    ws.cell(r, 2).font = DATA_FONT

    if p.get("title"):
        ws.cell(r, 3).value = p["title"]
        ws.cell(r, 3).alignment = LEFT

    exp = p.get("exp") or sanitize_exp(old_exp)
    exp = re.sub(r"^Theo NV-[A-Z]+\.\s*", "", exp)
    exp = f"Theo {nv_tag(tid)}.\n{exp}"
    ws.cell(r, 7).value = exp
    ws.cell(r, 7).alignment = LEFT
    ws.cell(r, 7).font = DATA_FONT

    loai = p.get("loai") or classify_loai(old_loai)
    ws.cell(r, 10).value = loai
    if p.get("phan"):
        ws.cell(r, 11).value = p["phan"]


def copy_style_from(src, dest):
    dest.font = copy.copy(src.font)
    dest.border = copy.copy(src.border)
    dest.fill = copy.copy(src.fill)
    dest.alignment = copy.copy(src.alignment)
    dest.number_format = src.number_format


def append_cases(ws, rows: list[list]):
    last = ws.max_row
    # skip trailing empty
    while last > 2 and not str(ws.cell(last, 1).value or "").strip():
        last -= 1
    src_row = last
    for i, data in enumerate(rows):
        r = last + 1 + i
        dest = data[:7] + [None] + data[7:11]
        body = re.sub(r"^Theo NV-[A-Z]+\.\s*", "", str(dest[6] or ""))
        dest[6] = f"Theo {nv_tag(str(data[0]))}.\n{body}"
        fill = ROW_A if (r % 2 == 0) else ROW_B
        for c, val in enumerate(dest, 1):
            cell = ws.cell(r, c)
            copy_style_from(ws.cell(src_row, c), cell)
            cell.value = val
            cell.alignment = CENTER if c in (1, 9, 10, 11) else LEFT
            cell.fill = WARN_FILL if c == 8 else fill
            cell.font = BOLD if c == 1 else DATA_FONT
            cell.border = THIN
            cell.number_format = "@"
        ws.row_dimensions[r].height = 72
    # extend filter and DV
    new_last = last + len(rows)
    ws.auto_filter.ref = f"A2:L{new_last}"
    for dv in ws.data_validations.dataValidation:
        dv.add(f"H3:H{new_last}")


def count_sheet(ws, prefix: str):
    pat = re.compile(rf"^{re.escape(prefix)}\d+$")
    n = pos = neg = bva = extra = 0
    for r in range(3, ws.max_row + 1):
        tid = str(ws.cell(r, 1).value or "").strip()
        if not pat.match(tid):
            continue
        n += 1
        loai = str(ws.cell(r, 10).value or "")
        tech = str(ws.cell(r, 2).value or "").lower()
        phan = str(ws.cell(r, 11).value or "").lower()
        low = loai.lower()
        if "boundary" in low or "biên" in low:
            bva += 1
        elif "negative" in low or "không hợp lệ" in low:
            neg += 1
        else:
            pos += 1
        if any(k in (low + tech + phan) for k in ("concurrency", "security", "đồng thời", "phân quyền", "auth")):
            extra += 1
    return n, pos, neg, bva, extra


def main():
    if not XLSX.exists():
        raise SystemExit(f"Không thấy {XLSX}")
    wb = load_workbook(XLSX)

    # remove old intro sheets if re-run
    for name in list(wb.sheetnames):
        if name.startswith("0. ") or name.startswith("1. ") or name.startswith("2. ") or name.startswith("3. "):
            del wb[name]

    prefixes = {
        "Đợt giảm giá": "SALE",
        "Phiếu giảm giá": "VC",
        "Tạo đơn từ giỏ hàng": "OC",
        "Thanh toán đơn hàng": "PAY",
        "Cập nhật trạng thái đơn hàng": "TC",
        "Thông báo đơn hàng": "TB",
        "Lọc tìm kiếm đơn hàng": "SR",
        "Trả hàng hoàn tiền": "RT",
    }

    updated = 0
    for sheet, prefix in prefixes.items():
        ws = wb[sheet]
        extra = NEW_CASES.get(sheet) or []
        existing_ids = set()
        for r in range(3, ws.max_row + 1):
            tid = str(ws.cell(r, 1).value or "").strip()
            if re.match(rf"^{prefix}\d+$", tid):
                existing_ids.add(tid)
                apply_row(ws, r, tid)
                updated += 1
        to_add = [row for row in extra if row[0] not in existing_ids]
        if to_add:
            append_cases(ws, to_add)
            updated += len(to_add)

    ws3 = write_intro(wb)
    write_drift(wb)

    matrix = [
        ("Đợt giảm giá", "UC-25.1–25.3", *count_sheet(wb["Đợt giảm giá"], "SALE")[:5],
         "Biên %, thời hạn, chồng đợt, phân quyền", "NV-SALE-01…07"),
        ("Phiếu giảm giá", "UC-25.4–25.6", *count_sheet(wb["Phiếu giảm giá"], "VC")[:5],
         "3 loại mã, lượt, tối thiểu, mã cá nhân", "NV-VC-01…06"),
        ("Đặt hàng từ giỏ", "UC-06, UC-14", *count_sheet(wb["Tạo đơn từ giỏ hàng"], "OC")[:5],
         "Tồn kho, giá lúc đặt, chống trùng, vãng lai", "NV-OC-01…06"),
        ("Thanh toán", "UC-08", *count_sheet(wb["Thanh toán đơn hàng"], "PAY")[:5],
         "COD thu lúc giao; VNPay success/fail/chữ ký/lệch tiền; đối soát; replay", "NV-PAY-01…04"),
        ("Trạng thái đơn", "UC-09, 17, 21", *count_sheet(wb["Cập nhật trạng thái đơn hàng"], "TC")[:5],
         "Không nhảy cóc, hủy, GHN, lịch sử", "NV-TC-01…04"),
        ("Thông báo đơn", "UC-12", *count_sheet(wb["Thông báo đơn hàng"], "TB")[:5],
         "COD vs VNPay, trả hàng, mail lỗi", "NV-TB-01…03"),
        ("Lọc / tra cứu đơn", "UC-09, 17", *count_sheet(wb["Lọc tìm kiếm đơn hàng"], "SR")[:5],
         "Lọc AND, ẩn chờ trả, tra cứu vãng lai", "NV-SR-01"),
        ("Trả hàng & hoàn tiền", "UC-10, 18", *count_sheet(wb["Trả hàng hoàn tiền"], "RT")[:5],
         "7 ngày, nhận hàng rồi hoàn, không IDOR", "NV-RT-01…08"),
    ]
    # count_sheet returns 5-tuple already unpacked with *[:5] — wait I used *count[:5] which is n,pos,neg,bva,extra
    fill_rows = []
    total = [0, 0, 0, 0, 0]
    for row in matrix:
        name, uc, n, pos, neg, bva, extra, tech, nv = row
        fill_rows.append((name, uc, n, pos, neg, bva, extra, tech, nv))
        for i, v in enumerate((n, pos, neg, bva, extra)):
            total[i] += v
    fill_rows.append(("TỔNG (8 chức năng)", "—", *total, "Truy vết NV-xx + UC", "Sheet 1"))
    fill_matrix(ws3, fill_rows)

    wb.save(XLSX)
    print(f"Updated {updated} cases -> {XLSX}")
    print("Total in matrix:", total[0])


if __name__ == "__main__":
    main()
