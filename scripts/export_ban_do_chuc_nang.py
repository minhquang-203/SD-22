# -*- coding: utf-8 -*-
"""Xuất bản đồ chức năng bảo vệ đồ án ra Excel (mỗi chức năng 1 sheet)."""
from openpyxl import Workbook
from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chartsheet import Chartsheet
from copy import copy

OUT = r"D:\Desktop\SD-22\Ban_do_chuc_nang_bao_ve.xlsx"

# --- colors ---
NAVY = "1B3A4B"
TEAL = "0F766E"
TEAL2 = "115E59"
GOLD = "B45309"
SLATE = "334155"
ROW_A = "F8FAFC"
ROW_B = "F1F5F9"
SEC_API = "0F766E"
SEC_MODEL = "1D4ED8"
SEC_BE = "7C3AED"
SEC_FE = "BE185D"
SEC_DB = "B45309"
SEC_FLOW = "0369A1"
WHITE = "FFFFFF"
INDEX_FILL = "ECFDF5"

thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
thick_bottom = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="medium", color="0F766E"),
)

font_title = Font(name="Calibri", size=16, bold=True, color=WHITE)
font_sub = Font(name="Calibri", size=11, italic=True, color="CCFBF1")
font_sec = Font(name="Calibri", size=12, bold=True, color=WHITE)
font_head = Font(name="Calibri", size=10, bold=True, color=WHITE)
font_cell = Font(name="Calibri", size=10, color=SLATE)
font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
font_index_title = Font(name="Calibri", size=18, bold=True, color=WHITE)
font_link = Font(name="Calibri", size=11, bold=True, color=TEAL, underline="single")
wrap = Alignment(wrap_text=True, vertical="center")
wrap_left = Alignment(wrap_text=True, vertical="center", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")


def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def merge_title(ws, cols, row, value, fill_hex, font, height=28):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, value)
    cell.font = font
    cell.fill = fill(fill_hex)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill(fill_hex)
        ws.cell(row, c).border = Border()
    ws.row_dimensions[row].height = height
    return row + 1


def section(ws, cols, row, title, color):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, title)
    cell.font = font_sec
    cell.fill = fill(color)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill(color)
    ws.row_dimensions[row].height = 22
    return row + 1


def headers(ws, row, cols_titles, color="1E293B"):
    for i, t in enumerate(cols_titles, 1):
        cell = ws.cell(row, i, t)
        cell.font = font_head
        cell.fill = fill(color)
        cell.alignment = center
        cell.border = thin
    ws.row_dimensions[row].height = 20
    return row + 1


def data_row(ws, row, values, zebra=True):
    bg = ROW_A if (row % 2 == 0) else ROW_B
    max_lines = 1
    for i, v in enumerate(values, 1):
        cell = ws.cell(row, i, v if v is not None else "")
        cell.font = font_bold if i == 1 else font_cell
        cell.fill = fill(bg)
        cell.alignment = wrap_left
        cell.border = thin
        if v:
            max_lines = max(max_lines, str(v).count("\n") + 1)
            # estimate wrap
            if len(str(v)) > 80:
                max_lines = max(max_lines, min(6, len(str(v)) // 55 + 1))
    ws.row_dimensions[row].height = max(18, min(90, 14 * max_lines + 6))
    return row + 1


def blank(ws, row):
    ws.row_dimensions[row].height = 8
    return row + 1


def note_row(ws, cols, row, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row, 1, text)
    cell.font = Font(name="Calibri", size=10, italic=True, color="334155")
    cell.fill = fill("FEF3C7")
    cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = fill("FEF3C7")
    ws.row_dimensions[row].height = 22 if len(text) < 120 else 36
    return row + 1


def finish_sheet(ws, cols):
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:3"
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "&B SUNOVA — Bản đồ chức năng"
    ws.oddFooter.right.text = "Trang &P / &N"


def write_feature(ws, meta):
    """meta keys: title, subtitle, package, note, fe_entry, apis, flow, models, tables, backend, frontend"""
    cols = 4
    set_widths(ws, [28, 52, 42, 55])
    r = 1
    r = merge_title(ws, cols, r, meta["title"], NAVY, font_title, 30)
    r = merge_title(ws, cols, r, meta["subtitle"], TEAL, font_sub, 22)
    r = note_row(ws, cols, r, f"Package: {meta['package']}    |    {meta.get('note', '')}")
    r = blank(ws, r)

    r = section(ws, cols, r, "1. ENTRY FRONTEND (route / màn hình)", GOLD)
    r = headers(ws, r, ["Hạng mục", "Route / vị trí", "File", "Mô tả"])
    for row in meta["fe_entry"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "2. ENTRY API (mỗi endpoint 1 dòng)", SEC_API)
    r = headers(ws, r, ["Method + Path", "Controller.method()", "File controller", "Mô tả"])
    for row in meta["apis"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "3. LUỒNG XỬ LÝ (controller → service → repository)", SEC_FLOW)
    r = headers(ws, r, ["Bước", "Thành phần", "File / hàm chính", "Việc làm"])
    for row in meta["flow"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "4. MODEL / DTO / ENUM (mỗi model 1 dòng)", SEC_MODEL)
    r = headers(ws, r, ["Tên", "Loại", "File", "Mô tả chức năng"])
    for row in meta["models"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "5. BẢNG DỮ LIỆU", SEC_DB)
    r = headers(ws, r, ["Bảng", "Vai trò", "Migration / cột đáng nhớ", "Ghi chú"])
    for row in meta["tables"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "6. BACKEND — toàn bộ file liên quan", SEC_BE)
    r = headers(ws, r, ["Lớp", "File (từ org.example.templatejava6/)", "Vai trò", "Ghi chú"])
    for row in meta["backend"]:
        r = data_row(ws, r, row)
    r = blank(ws, r)

    r = section(ws, cols, r, "7. FRONTEND — toàn bộ file liên quan", SEC_FE)
    r = headers(ws, r, ["Lớp", "File (từ frontend/src/)", "Vai trò", "Ghi chú"])
    for row in meta["frontend"]:
        r = data_row(ws, r, row)

    finish_sheet(ws, cols)
    ws.auto_filter.ref = None


# =============================================================================
# DATA
# =============================================================================

FEATURES = []

FEATURES.append({
    "sheet": "01. Gio hang",
    "title": "01. CART — Quản lý sản phẩm trong giỏ hàng",
    "subtitle": "Khách login: JWT + bảng gio_hang. Khách vãng lai: localStorage key sunova_cart (không gọi API).",
    "package": "cart",
    "note": "ID khách lấy từ JWT, không lấy từ body. Role: KHACH_HANG.",
    "fe_entry": [
        ["Trang giỏ", "/gio-hang", "views/storefront/GioHang.vue", "Màn hình xem / sửa / xóa / sang đặt hàng"],
        ["Thêm hàng", "Trang chi tiết SP", "views/storefront/ProductDetail.vue", "Gọi useCart().addItem"],
        ["State giỏ", "Composable toàn site", "composables/useCart.js", "Login → API; guest → localStorage sunova_cart"],
        ["Badge navbar", "Mọi trang storefront", "components/storefront/TheNavbar.vue", "Số lượng + link /gio-hang"],
        ["Toast thêm giỏ", "Layout storefront", "components/storefront/CartToast.vue", "Thông báo đã thêm"],
    ],
    "apis": [
        ["GET /api/gio-hang", "GioHangController.getByKhachHang()", "cart/controller/GioHangController.java", "Lấy giỏ theo JWT"],
        ["POST /api/gio-hang/add", "GioHangController.add()", "cart/controller/GioHangController.java", "Thêm / cộng dồn dòng"],
        ["PUT /api/gio-hang/update/{idChiTietGioHang}", "GioHangController.updateSoLuong()", "cart/controller/GioHangController.java", "Đổi số lượng"],
        ["DELETE /api/gio-hang/delete/{idChiTietGioHang}", "GioHangController.deleteItem()", "cart/controller/GioHangController.java", "Xóa 1 dòng"],
        ["DELETE /api/gio-hang/clear", "GioHangController.clear()", "cart/controller/GioHangController.java", "Xóa hết giỏ"],
    ],
    "flow": [
        ["1", "Controller", "GioHangController", "Nhận request, lấy idKhachHang từ SecurityContext"],
        ["2", "Service", "GioHangService.getOrCreateGioHang / add / updateSoLuong / deleteItem / clear", "Nghiệp vụ + validateStock"],
        ["3", "Giá KM", "DotGiamGiaService (VariantSaleInfo)", "Hiện giá đang sale trên dòng giỏ"],
        ["4", "Tồn kho", "ChiTietSanPhamRepository", "Không cho SL > tồn"],
        ["5", "Repo", "GioHangRepository + ChiTietGioHangRepository", "Đọc/ghi gio_hang, chi_tiet_gio_hang"],
    ],
    "models": [
        ["GioHang", "Entity", "cart/entity/GioHang.java", "1 giỏ ↔ 1 khách (bảng gio_hang)"],
        ["ChiTietGioHang", "Entity", "cart/entity/ChiTietGioHang.java", "1 dòng: biến thể + số lượng"],
        ["ThemGioHangRequest", "Request", "cart/model/request/ThemGioHangRequest.java", "Body thêm: idChiTietSanPham, soLuong"],
        ["CapNhatSoLuongGioHangRequest", "Request", "cart/model/request/CapNhatSoLuongGioHangRequest.java", "Body sửa số lượng"],
        ["GioHangResponse", "Response", "cart/model/response/GioHangResponse.java", "JSON giỏ: list dòng + tổng"],
        ["ChiTietGioHangResponse", "Response", "cart/model/response/ChiTietGioHangResponse.java", "JSON 1 dòng (ảnh, giá, tồn, giá sale)"],
    ],
    "tables": [
        ["gio_hang", "Header giỏ", "V1__init_full.sql — FK id_khach_hang", "1 khách 1 giỏ"],
        ["chi_tiet_gio_hang", "Dòng giỏ", "V1 — FK id_gio_hang, id_chi_tiet_san_pham", "Cộng dồn nếu trùng biến thể"],
    ],
    "backend": [
        ["Controller", "cart/controller/GioHangController.java", "5 endpoint giỏ", "Role KHACH_HANG"],
        ["Service", "cart/service/GioHangService.java", "Nghiệp vụ giỏ + validateStock", ""],
        ["Repo", "cart/repository/GioHangRepository.java", "findFirstByKhachHang_Id…", ""],
        ["Repo", "cart/repository/ChiTietGioHangRepository.java", "Dòng giỏ theo HĐ/khách", ""],
        ["Entity", "cart/entity/GioHang.java", "Bảng gio_hang", ""],
        ["Entity", "cart/entity/ChiTietGioHang.java", "Bảng chi_tiet_gio_hang", ""],
        ["DTO", "cart/model/request/ThemGioHangRequest.java", "Body thêm", ""],
        ["DTO", "cart/model/request/CapNhatSoLuongGioHangRequest.java", "Body sửa SL", ""],
        ["DTO", "cart/model/response/GioHangResponse.java", "JSON giỏ", ""],
        ["DTO", "cart/model/response/ChiTietGioHangResponse.java", "JSON dòng", ""],
        ["Phụ", "customer/repository/KhachHangRepository.java", "Load khách", ""],
        ["Phụ", "common/entity/KhachHang.java", "Entity khách", ""],
        ["Phụ", "product/entity/ChiTietSanPham.java", "Biến thể + tồn", ""],
        ["Phụ", "product/repository/ChiTietSanPhamRepository.java", "Check tồn", ""],
        ["Phụ", "product/entity/AnhSanPham.java", "Ảnh dòng giỏ", ""],
        ["Phụ", "product/repository/AnhSanPhamRepository.java", "Load ảnh", ""],
        ["Phụ", "voucher/service/DotGiamGiaService.java", "Giá sale trên giỏ", ""],
        ["Phụ", "voucher/model/response/VariantSaleInfo.java", "DTO giá KM biến thể", ""],
        ["Config", "common/config/SecurityConfig.java", "Khóa /api/gio-hang/**", "hasRole KHACH_HANG"],
        ["SQL", "db/migration/V1__init_full.sql", "Tạo bảng giỏ", "resources/..."],
    ],
    "frontend": [
        ["Router", "router/index.js", "Khai báo path gio-hang", ""],
        ["View", "views/storefront/GioHang.vue", "UI giỏ", ""],
        ["View", "views/storefront/ProductDetail.vue", "Nút thêm giỏ", ""],
        ["Composable", "composables/useCart.js", "State + API/localStorage", "Key sunova_cart"],
        ["API", "api/gioHangApi.js", "Client 5 endpoint", ""],
        ["API", "api/request.js", "Axios + JWT; prefix /gio-hang", ""],
        ["UI", "components/storefront/CartToast.vue", "Toast thêm giỏ", ""],
        ["UI", "components/storefront/TheNavbar.vue", "Badge số lượng", ""],
        ["Layout", "layouts/StorefrontLayout.vue", "Mount CartToast", ""],
        ["Boot", "main.js", "Side-effect import useCart", ""],
    ],
})

FEATURES.append({
    "sheet": "02. Tao don hang",
    "title": "02. CHECKOUT — Tạo đơn hàng từ giỏ hàng",
    "subtitle": "Online: POST /api/online/checkout (login) hoặc /guest/checkout. POS: POST /api/ban-hang/tai-quay. Không có package tên checkout.",
    "package": "order (+ shipping, payment, voucher, cart)",
    "note": "Online chỉ COD + VNPAY. Guest không đụng bảng gio_hang. Có idempotencyKey chống double-click.",
    "fe_entry": [
        ["Checkout online", "/dat-hang", "views/storefront/DatHang.vue", "Địa chỉ, voucher, PTTT, đặt hàng"],
        ["Từ giỏ sang", "/gio-hang → /dat-hang", "views/storefront/GioHang.vue", "Chọn dòng rồi checkout"],
        ["Modal địa chỉ", "Popup trên /dat-hang", "components/storefront/CheckoutRecipientModal.vue", "Chọn tỉnh/xã GHN + địa chỉ"],
        ["Modal voucher", "Popup trên /dat-hang", "components/storefront/CheckoutVoucherModal.vue", "Chọn phiếu giảm giá"],
        ["POS tại quầy", "/admin/pos", "views/admin/pos/PosPage.vue", "Tạo HĐ không qua giỏ online"],
        ["Modal voucher POS", "Popup POS", "components/admin/PosVoucherModal.vue", "Áp voucher tại quầy"],
    ],
    "apis": [
        ["GET /api/online/vouchers", "OnlineCheckoutController.danhSachVoucher()", "order/controller/OnlineCheckoutController.java", "Voucher áp được lúc checkout"],
        ["POST /api/online/tinh-gia", "OnlineCheckoutController.tinhGia()", "order/controller/OnlineCheckoutController.java", "Preview giá user login"],
        ["POST /api/online/checkout", "OnlineCheckoutController.checkout()", "order/controller/OnlineCheckoutController.java", "Tạo đơn từ idsChiTietGioHang"],
        ["POST /api/online/guest/tinh-gia", "OnlineCheckoutController.tinhGiaGuest()", "order/controller/OnlineCheckoutController.java", "Preview giá khách vãng lai (permitAll)"],
        ["POST /api/online/guest/checkout", "OnlineCheckoutController.checkoutGuest()", "order/controller/OnlineCheckoutController.java", "Tạo đơn guest, gửi line items"],
        ["GET /api/shipping/provinces", "ShippingController.provinces()", "shipping/controller/ShippingController.java", "Tỉnh/thành GHN"],
        ["GET /api/shipping/wards", "ShippingController.wards()", "shipping/controller/ShippingController.java", "Phường/xã theo provinceId"],
        ["POST /api/shipping/fee", "ShippingController.fee()", "shipping/controller/ShippingController.java", "Tính phí ship preview"],
        ["POST /api/ban-hang/tinh-gia", "BanHangController.tinhGia()", "order/controller/BanHangController.java", "Tính giá POS"],
        ["GET /api/ban-hang/vouchers", "BanHangController.vouchers()", "order/controller/BanHangController.java", "Voucher POS"],
        ["POST /api/ban-hang/tai-quay", "BanHangController.taoDonTaiQuay()", "order/controller/BanHangController.java", "Tạo HĐ tại quầy"],
        ["POST /api/ban-hang/cho", "BanHangController.giuDonCho()", "order/controller/BanHangController.java", "Giữ đơn chờ POS"],
        ["GET /api/ban-hang/cho", "BanHangController.danhSachDonCho()", "order/controller/BanHangController.java", "List đơn chờ"],
        ["GET /api/ban-hang/cho/{id}", "BanHangController.chiTietDonCho()", "order/controller/BanHangController.java", "Chi tiết đơn chờ"],
        ["DELETE /api/ban-hang/cho/{id}", "BanHangController.huyDonCho()", "order/controller/BanHangController.java", "Hủy đơn chờ"],
    ],
    "flow": [
        ["1", "Controller", "OnlineCheckoutController.checkout / checkoutGuest", "Nhận body + JWT (hoặc guest)"],
        ["2", "Giá", "CheckoutPricingService.loadActiveSales / resolveDonGia / tinhTienGiamPhieu", "Giá sale + voucher"],
        ["3", "Ship", "ShippingService.calcFee", "Phí GHN, fallback 30.000 nếu lỗi"],
        ["4", "Ghi đơn", "OnlineCheckoutService → HoaDon + HoaDonChiTiet", "Transaction; idempotency_key"],
        ["5", "Thanh toán", "COD: taoThanhToanCod | VNPay: PaymentService.taoThanhToan", "Trả paymentUrl nếu VNPay"],
        ["6", "Giỏ", "Trừ/xóa dòng giỏ khi COD thành công (VNPay trừ sau IPN)", "OnlineOrderLifecycleService.truGioHangTheoDon"],
        ["7", "POS", "BanHangService.taoDonTaiQuay", "Tạo HĐ loaiDon=TAI_QUAY"],
    ],
    "models": [
        ["OnlineCheckoutRequest", "Request", "order/model/request/OnlineCheckoutRequest.java", "idsChiTietGioHang, PTTT, voucher, địa chỉ GHN, idempotencyKey"],
        ["OnlineTinhGiaRequest", "Request", "order/model/request/OnlineTinhGiaRequest.java", "Body preview giá login"],
        ["GuestCheckoutRequest", "Request", "order/model/request/GuestCheckoutRequest.java", "Body đặt hàng guest + người nhận"],
        ["GuestCheckoutItemRequest", "Request", "order/model/request/GuestCheckoutItemRequest.java", "1 dòng guest: id biến thể + SL"],
        ["GuestTinhGiaRequest", "Request", "order/model/request/GuestTinhGiaRequest.java", "Body preview giá guest"],
        ["OnlineTinhGiaResponse", "Response", "order/model/response/OnlineTinhGiaResponse.java", "Tạm, giảm, ship, thành tiền"],
        ["OnlineCheckoutResponse", "Response", "order/model/response/OnlineCheckoutResponse.java", "Đơn vừa tạo + paymentUrl + trackingToken"],
        ["HuyDonOnlineRequest", "Request", "order/model/request/HuyDonOnlineRequest.java", "Lý do hủy (API cancel)"],
        ["TaoDonTaiQuayRequest", "Request", "order/model/request/TaoDonTaiQuayRequest.java", "Body tạo HĐ POS"],
        ["PosTinhGiaRequest", "Request", "order/model/request/PosTinhGiaRequest.java", "Body tính giá POS"],
        ["PosTinhGiaResponse", "Response", "order/model/response/PosTinhGiaResponse.java", "JSON giá POS"],
        ["GiuDonChoRequest", "Request", "order/model/request/GiuDonChoRequest.java", "Body giữ đơn chờ"],
        ["GiuDonChoResponse", "Response", "order/model/response/GiuDonChoResponse.java", "JSON đơn chờ vừa giữ"],
        ["DonChoListItemResponse", "Response", "order/model/response/DonChoListItemResponse.java", "1 dòng list đơn chờ"],
        ["DonChoDetailResponse", "Response", "order/model/response/DonChoDetailResponse.java", "Chi tiết đơn chờ"],
        ["BanHangHoaDonResponse", "Response", "order/model/response/BanHangHoaDonResponse.java", "HĐ vừa bán tại quầy"],
        ["BienTheBanResponse", "Response", "order/model/response/BienTheBanResponse.java", "Biến thể bán POS"],
        ["HoaDon", "Entity", "order/entity/HoaDon.java", "Hóa đơn (ONLINE / TAI_QUAY)"],
        ["HoaDonChiTiet", "Entity", "order/entity/HoaDonChiTiet.java", "Dòng hàng hóa đơn"],
        ["ThanhToanHoaDon", "Entity", "order/entity/ThanhToanHoaDon.java", "Giao dịch TT (COD ghi luôn)"],
        ["LichSuDonHang", "Entity", "order/entity/LichSuDonHang.java", "Timeline trạng thái lúc tạo đơn"],
    ],
    "tables": [
        ["hoa_don", "Header đơn", "V1 + V2 idempotency_key + V3 GHN names + V5 email + V8 tracking_token", "loaiDon ONLINE/TAI_QUAY"],
        ["hoa_don_chi_tiet", "Dòng hàng", "V1", ""],
        ["thanh_toan_hoa_don", "GD thanh toán", "V1 + provider cols", "COD tạo ngay"],
        ["lich_su_don_hang", "Timeline", "V1", "Ghi lúc tạo"],
        ["phieu_giam_gia", "Voucher áp", "hoa_don.id_phieu_giam_gia", ""],
        ["phuong_thuc_thanh_toan", "PTTT", "Seed COD, VNPAY, TIEN_MAT…", "Online chỉ COD+VNPAY"],
    ],
    "backend": [
        ["Controller", "order/controller/OnlineCheckoutController.java", "Checkout online + guest", ""],
        ["Controller", "order/controller/BanHangController.java", "POS tạo đơn / đơn chờ", ""],
        ["Controller", "shipping/controller/ShippingController.java", "Tỉnh/xã/phí lúc đặt", ""],
        ["Service", "order/service/OnlineCheckoutService.java", "checkout / tinhGia / taoThanhToanCod", "Lõi"],
        ["Service", "order/service/CheckoutPricingService.java", "Giá sale + phiếu", ""],
        ["Service", "order/service/BanHangService.java", "taoDonTaiQuay", ""],
        ["Service", "order/service/OnlineOrderLifecycleService.java", "Trừ giỏ / hủy", ""],
        ["Service", "order/service/OrderTrackingTokenGenerator.java", "Sinh tracking_token", ""],
        ["Repo", "order/repository/HoaDonRepository.java", "Lưu HĐ", ""],
        ["Repo", "order/repository/HoaDonChiTietRepository.java", "Lưu dòng", ""],
        ["Repo", "order/repository/ThanhToanHoaDonRepository.java", "GD TT", ""],
        ["Repo", "order/repository/LichSuDonHangRepository.java", "Timeline", ""],
        ["Repo", "order/repository/PhuongThucThanhToanRepository.java", "PTTT", ""],
        ["Entity", "order/entity/HoaDon.java", "Bảng hoa_don", ""],
        ["Entity", "order/entity/HoaDonChiTiet.java", "Bảng hoa_don_chi_tiet", ""],
        ["Entity", "order/entity/ThanhToanHoaDon.java", "Bảng thanh_toan_hoa_don", ""],
        ["Entity", "order/entity/LichSuDonHang.java", "Bảng lich_su_don_hang", ""],
        ["DTO", "order/model/request/OnlineCheckoutRequest.java", "Body checkout login", ""],
        ["DTO", "order/model/request/OnlineTinhGiaRequest.java", "Body preview", ""],
        ["DTO", "order/model/request/GuestCheckoutRequest.java", "Body guest", ""],
        ["DTO", "order/model/request/GuestCheckoutItemRequest.java", "Dòng guest", ""],
        ["DTO", "order/model/request/GuestTinhGiaRequest.java", "Preview guest", ""],
        ["DTO", "order/model/request/HuyDonOnlineRequest.java", "Hủy", ""],
        ["DTO", "order/model/request/TaoDonTaiQuayRequest.java", "POS", ""],
        ["DTO", "order/model/request/PosTinhGiaRequest.java", "POS giá", ""],
        ["DTO", "order/model/request/GiuDonChoRequest.java", "Đơn chờ", ""],
        ["DTO", "order/model/response/OnlineCheckoutResponse.java", "Kết quả đặt", ""],
        ["DTO", "order/model/response/OnlineTinhGiaResponse.java", "Preview", ""],
        ["DTO", "order/model/response/BanHangHoaDonResponse.java", "HĐ POS", ""],
        ["DTO", "order/model/response/PosTinhGiaResponse.java", "Giá POS", ""],
        ["DTO", "order/model/response/GiuDonChoResponse.java", "Đơn chờ", ""],
        ["DTO", "order/model/response/DonChoListItemResponse.java", "List chờ", ""],
        ["DTO", "order/model/response/DonChoDetailResponse.java", "Chi tiết chờ", ""],
        ["DTO", "order/model/response/BienTheBanResponse.java", "SP POS", ""],
        ["Phụ", "shipping/service/ShippingService.java", "calcFee", ""],
        ["Phụ", "payment/service/PaymentService.java", "Tạo VNPay", ""],
        ["Phụ", "voucher/service/PhieuGiamGiaService.java", "consumeOne", ""],
        ["Phụ", "voucher/service/DotGiamGiaService.java", "Giá sale", ""],
        ["Phụ", "cart/repository/GioHangRepository.java", "Đọc giỏ", ""],
        ["Phụ", "cart/repository/ChiTietGioHangRepository.java", "Dòng giỏ", ""],
        ["Phụ", "product/repository/ChiTietSanPhamRepository.java", "Tồn", ""],
        ["Phụ", "product/service/LoHangService.java", "Lô (nếu cần)", ""],
        ["Phụ", "notification/service/ThongBaoService.java", "DON_HANG_MOI", ""],
        ["Phụ", "notification/event/DatHangThanhCongMailEvent.java", "Bắn mail", ""],
        ["Entity chung", "common/entity/PhieuGiamGia.java", "Voucher", ""],
        ["Entity chung", "common/entity/PhuongThucThanhToan.java", "PTTT", ""],
        ["Enum", "common/enums/TrangThaiDonHang.java", "CHO_XAC_NHAN lúc tạo online", ""],
        ["Config", "common/config/SecurityConfig.java", "guest permitAll, /api/online KHACH_HANG", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Tạo bảng đơn", ""],
        ["SQL", "db/migration/V2__payment_idempotency.sql", "idempotency_key", ""],
        ["SQL", "db/migration/V3__ghn_address_names_on_hoa_don.sql", "Tên tỉnh/xã GHN", ""],
        ["SQL", "db/migration/V5__hoa_don_email_nguoi_nhan.sql", "email_nguoi_nhan", ""],
        ["SQL", "db/migration/V8__hoa_don_tracking_token.sql", "tracking_token", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "dat-hang, admin/pos", ""],
        ["View", "views/storefront/DatHang.vue", "Checkout online", ""],
        ["View", "views/storefront/GioHang.vue", "Nút sang đặt hàng", ""],
        ["View", "views/admin/pos/PosPage.vue", "Bán tại quầy", ""],
        ["UI", "components/storefront/CheckoutRecipientModal.vue", "Địa chỉ GHN", ""],
        ["UI", "components/storefront/CheckoutVoucherModal.vue", "Chọn voucher", ""],
        ["UI", "components/admin/PosVoucherModal.vue", "Voucher POS", ""],
        ["API", "api/onlineCheckout.js", "tinh-gia / checkout / guest", ""],
        ["API", "api/shipping.js", "provinces / wards / fee", ""],
        ["API", "api/banHangApi.js", "POS", ""],
        ["API", "api/request.js", "Axios", ""],
        ["State", "composables/useCart.js", "ids dòng giỏ gửi checkout", ""],
        ["Auth", "composables/useAuth.js", "Login vs guest", ""],
        ["CSS", "styles/checkoutCss.css", "Style trang đặt", ""],
        ["CSS", "styles/checkoutRecipientModal.css", "Modal địa chỉ", ""],
        ["CSS", "styles/checkoutVoucherModal.css", "Modal voucher", ""],
        ["CSS", "styles/posAdmin.css", "POS", ""],
        ["Asset", "assets/payment/cod.svg", "Logo COD", ""],
        ["Asset", "assets/payment/vnpay.svg", "Logo VNPay", ""],
        ["Menu", "constants/adminMenu.js", "Menu POS", ""],
    ],
})

FEATURES.append({
    "sheet": "03. Thanh toan",
    "title": "03. CHECKOUT & PAYMENT — Thanh toán",
    "subtitle": "Không có route /thanh-toan. Cổng thật chỉ VNPay. MoMo có trong seed DB nhưng không có gateway.",
    "package": "payment (+ order POS)",
    "note": "Callback + IPN permitAll. Timeout đơn VNPay chưa trả: OnlineOrderTimeoutScheduler (mặc định 15 phút).",
    "fe_entry": [
        ["Online", "/dat-hang (chọn COD | VNPAY)", "views/storefront/DatHang.vue", "COD xong tại chỗ; VNPay redirect cổng"],
        ["Return VNPay", "/dat-hang?success=…", "views/storefront/DatHang.vue", "frontend-return-url trong application.properties"],
        ["POS", "/admin/pos", "views/admin/pos/PosPage.vue", "Tiền mặt / CK / VNPay"],
    ],
    "apis": [
        ["POST /api/payments/{provider}/create", "PaymentController.taoThanhToan()", "payment/controller/PaymentController.java", "Tạo URL cổng TT"],
        ["GET /api/payments/vnpay/callback", "PaymentController.vnpayCallback()", "payment/controller/PaymentController.java", "User quay về → redirect FE"],
        ["GET /api/payments/vnpay/ipn", "PaymentController.vnpayIpn()", "payment/controller/PaymentController.java", "VNPay server báo kết quả JSON"],
        ["GET /api/phuong-thuc-thanh-toan", "PhuongThucThanhToanController.hienThiDanhSach()", "payment/controller/PhuongThucThanhToanController.java", "List PTTT (permitAll)"],
        ["GET /api/ban-hang/tai-quay/{id}/thanh-toan", "BanHangController.kiemTraThanhToan()", "order/controller/BanHangController.java", "Poll trạng thái VNPay POS"],
        ["POST /api/ban-hang/tai-quay/{id}/huy-thanh-toan", "BanHangController.huyThanhToan()", "order/controller/BanHangController.java", "Hủy phiên VNPay POS"],
        ["POST /api/ban-hang/tai-quay/{id}/hoan-tat-thanh-toan", "BanHangController.hoanTatThanhToan()", "order/controller/BanHangController.java", "Chốt TT thủ công POS"],
    ],
    "flow": [
        ["1", "Tạo GD", "PaymentService.taoThanhToan → PaymentGatewayRegistry → VnpayGateway", "Sinh URL + lưu thanh_toan_hoa_don"],
        ["2", "User trả", "Cổng VNPay", "Redirect return-url backend"],
        ["3", "Callback", "vnpayCallback → verify chữ ký → redirect /dat-hang", "Không phải nguồn sự thật duy nhất"],
        ["4", "IPN", "vnpayIpn → findByMaGiaoDichForUpdate (pessimistic lock)", "Nguồn cập nhật trạng thái chính"],
        ["5", "Sau OK", "OnlineOrderLifecycleService.truGioHangTheoDon + mail + thông báo", ""],
        ["6", "Timeout", "OnlineOrderTimeoutScheduler", "Hủy đơn VNPay quá hạn (15p)"],
        ["7", "POS", "BanHangService + PosOrderLifecycleService", "TM/CK/VNPay tại quầy"],
    ],
    "models": [
        ["TaoThanhToanRequest", "Request", "payment/model/request/TaoThanhToanRequest.java", "Body: idHoaDon"],
        ["TaoThanhToanResponse", "Response", "payment/model/response/TaoThanhToanResponse.java", "paymentUrl + transactionRef"],
        ["KetQuaThanhToanResponse", "Response", "payment/model/response/KetQuaThanhToanResponse.java", "Kết quả sau callback"],
        ["VnpayIpnResponse", "Response", "payment/model/response/VnpayIpnResponse.java", "RspCode / Message trả VNPay"],
        ["PhuongThucThanhToanResponse", "Response", "payment/model/response/PhuongThucThanhToanResponse.java", "JSON 1 PTTT (mã, tên)"],
        ["PosThanhToanStatusResponse", "Response", "order/model/response/PosThanhToanStatusResponse.java", "Trạng thái poll POS"],
        ["PaymentCreateCommand", "Gateway DTO", "payment/gateway/PaymentCreateCommand.java", "Lệnh nội bộ gửi cổng"],
        ["PaymentCreateResult", "Gateway DTO", "payment/gateway/PaymentCreateResult.java", "Kết quả tạo URL"],
        ["PaymentCallbackResult", "Gateway DTO", "payment/gateway/PaymentCallbackResult.java", "Kết quả verify chữ ký/số tiền"],
        ["PaymentGateway", "Interface", "payment/gateway/PaymentGateway.java", "createPayment / verifyCallback"],
        ["ThanhToanHoaDon", "Entity", "order/entity/ThanhToanHoaDon.java", "Giao dịch: ma_giao_dich, provider_transaction_no"],
        ["PhuongThucThanhToan", "Entity", "common/entity/PhuongThucThanhToan.java", "Danh mục COD, VNPAY, TIEN_MAT, MOMO…"],
    ],
    "tables": [
        ["thanh_toan_hoa_don", "Giao dịch", "V1 + provider_transaction_no, provider_pay_date; V2 index ma_giao_dich", "Lock pesimist khi IPN"],
        ["phuong_thuc_thanh_toan", "Danh mục PTTT", "Seed TIEN_MAT, CHUYEN_KHOAN, VNPAY, MOMO, COD", "MOMO không có gateway"],
        ["hoa_don", "idempotency_key", "V2__payment_idempotency.sql", "Chống tạo đơn trùng"],
    ],
    "backend": [
        ["Controller", "payment/controller/PaymentController.java", "create / callback / ipn", ""],
        ["Controller", "payment/controller/PhuongThucThanhToanController.java", "GET list PTTT", ""],
        ["Service", "payment/service/PaymentService.java", "taoThanhToan, xuLyCallback, xuLyIpn", ""],
        ["Gateway", "payment/gateway/PaymentGateway.java", "Interface cổng", ""],
        ["Gateway", "payment/gateway/PaymentGatewayRegistry.java", "Chọn impl theo mã", "Chỉ VNPAY"],
        ["Gateway", "payment/gateway/PaymentCreateCommand.java", "Command tạo TT", ""],
        ["Gateway", "payment/gateway/PaymentCreateResult.java", "Result URL", ""],
        ["Gateway", "payment/gateway/PaymentCallbackResult.java", "Result verify", ""],
        ["VNPay", "payment/vnpay/VnpayGateway.java", "Impl duy nhất PROVIDER_CODE=VNPAY", ""],
        ["VNPay", "payment/vnpay/VnpayProperties.java", "prefix payment.vnpay", ""],
        ["DTO", "payment/model/request/TaoThanhToanRequest.java", "idHoaDon", ""],
        ["DTO", "payment/model/response/TaoThanhToanResponse.java", "URL", ""],
        ["DTO", "payment/model/response/KetQuaThanhToanResponse.java", "Kết quả", ""],
        ["DTO", "payment/model/response/VnpayIpnResponse.java", "IPN JSON", ""],
        ["DTO", "payment/model/response/PhuongThucThanhToanResponse.java", "PTTT", ""],
        ["Entity", "order/entity/ThanhToanHoaDon.java", "Bảng GD", ""],
        ["Entity", "common/entity/PhuongThucThanhToan.java", "Danh mục", ""],
        ["Repo", "order/repository/ThanhToanHoaDonRepository.java", "findByMaGiaoDichForUpdate", "Pessimistic lock"],
        ["Repo", "order/repository/PhuongThucThanhToanRepository.java", "findByMaIgnoreCase", ""],
        ["Lifecycle", "order/service/OnlineOrderLifecycleService.java", "Trừ giỏ sau OK", ""],
        ["Scheduler", "order/service/OnlineOrderTimeoutScheduler.java", "Hủy đơn quá hạn", "15 phút"],
        ["POS", "order/service/PosOrderLifecycleService.java", "Hoàn thành/hủy VNPay quầy", ""],
        ["POS", "order/service/BanHangService.java", "Poll / hủy / hoàn tất", ""],
        ["POS", "order/controller/BanHangController.java", "3 endpoint TT POS", ""],
        ["DTO POS", "order/model/response/PosThanhToanStatusResponse.java", "Status poll", ""],
        ["Config", "common/config/SecurityConfig.java", "callback/ipn permitAll", ""],
        ["Config", "resources/application.properties", "payment.vnpay.* + timeout", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Bảng PTTT + GD", ""],
        ["SQL", "db/migration/V2__payment_idempotency.sql", "idempotency + index", ""],
        ["Hoàn tiền", "payment/vnpay/VnpayRefundGateway.java", "Dùng ở sheet 14", "Không thuộc luồng trả tiền mua"],
    ],
    "frontend": [
        ["View", "views/storefront/DatHang.vue", "Chọn COD/VNPAY + nhận query return", ""],
        ["View", "views/admin/pos/PosPage.vue", "TM / CK / VNPay", ""],
        ["API", "api/onlineCheckout.js", "Checkout trả paymentUrl", ""],
        ["API", "api/banHangApi.js", "Poll / hủy / hoàn tất POS", ""],
        ["API", "api/request.js", "Axios", ""],
        ["Asset", "assets/payment/cod.svg", "Logo COD", ""],
        ["Asset", "assets/payment/vnpay.svg", "Logo VNPay", ""],
        ["Router", "router/index.js", "Landing return /dat-hang", ""],
    ],
})

FEATURES.append({
    "sheet": "04. Trang thai don",
    "title": "04. ORDER — Theo dõi và cập nhật trạng thái đơn hàng",
    "subtitle": "Không chỉ 1 dropdown: có nhánh gán lô, từ chối, sync GHN, khách hủy.",
    "package": "order + realtime",
    "note": "Enum TrangThaiDonHang: CHO, CHO_XAC_NHAN, DA_XAC_NHAN, DANG_CHUAN_BI, DANG_GIAO, HOAN_THANH, TRA_HANG, DA_HUY.",
    "fe_entry": [
        ["List admin", "/admin/hoa-don", "views/admin/orders/OrderList.vue", "Tab / badge trạng thái"],
        ["Chi tiết admin", "/admin/hoa-don/chi-tiet/:id", "views/admin/orders/HoaDonDetailPage.vue", "Nút chuyển TT, GHN, từ chối"],
        ["Gán lô", "Modal trên chi tiết", "components/admin/orders/GanLoDonHangModal.vue", "Xác nhận → DA_XAC_NHAN"],
        ["Khách theo dõi", "/tra-cuu-don", "views/storefront/TraCuuDon.vue", "Xem TT + hủy nếu được"],
    ],
    "apis": [
        ["PATCH /api/hoa-don/{id}/status", "HoaDonController.changeStatus()", "order/controller/HoaDonController.java", "Admin chuyển trạng thái"],
        ["GET /api/hoa-don/{id}/goi-y-gan-lo", "HoaDonController.goiYGanLo()", "order/controller/HoaDonController.java", "Gợi ý lô FIFO"],
        ["POST /api/hoa-don/{id}/xac-nhan-gan-lo", "HoaDonController.xacNhanGanLo()", "order/controller/HoaDonController.java", "Xác nhận + trừ lô"],
        ["POST /api/hoa-don/{id}/tu-choi", "HoaDonController.tuChoiDon()", "order/controller/HoaDonController.java", "Shop từ chối/hủy"],
        ["POST /api/hoa-don/{id}/dong-bo-ghn", "HoaDonController.dongBoGhn()", "order/controller/HoaDonController.java", "Kéo TT từ GHN"],
        ["POST /api/hoa-don/{id}/gia-lap-webhook-ghn", "HoaDonController.giaLapWebhookGhn()", "order/controller/HoaDonController.java", "Giả lập webhook (dev)"],
        ["GET /api/hoa-don/ghn-trang-thai", "HoaDonController.ghnTrangThai()", "order/controller/HoaDonController.java", "Map mã GHN → label"],
        ["POST /api/hoa-don/cua-toi/{id}/huy", "HoaDonController.huyDonCuaToi()", "order/controller/HoaDonController.java", "Khách hủy"],
        ["PATCH /api/online/orders/{idHoaDon}/cancel", "OnlineCheckoutController.huyDonHang()", "order/controller/OnlineCheckoutController.java", "Hủy từ checkout API"],
        ["GET /api/lich-su-don-hang/hoa-don", "LichSuDonHangController.getByHoaDon()", "order/controller/LichSuDonHangController.java", "Timeline 1 đơn"],
        ["GET /api/lich-su-don-hang", "LichSuDonHangController.hienThiDanhSach()", "order/controller/LichSuDonHangController.java", "List toàn bộ log"],
        ["GET /api/lich-su-don-hang/detail", "LichSuDonHangController.detail()", "order/controller/LichSuDonHangController.java", "Chi tiết 1 log"],
        ["POST /api/lich-su-don-hang/add", "LichSuDonHangController.add()", "order/controller/LichSuDonHangController.java", "Thêm log (admin)"],
        ["PUT /api/lich-su-don-hang/update/{id}", "LichSuDonHangController.update()", "order/controller/LichSuDonHangController.java", "Sửa log"],
        ["DELETE /api/lich-su-don-hang/delete", "LichSuDonHangController.delete()", "order/controller/LichSuDonHangController.java", "Xóa log"],
    ],
    "flow": [
        ["1", "Đổi TT tay", "HoaDonService.chuyenTrangThai", "Validate chuyển hợp lệ + ghiLichSuTrangThai"],
        ["2", "Xác nhận", "HoaDonService.xacNhanGanLo", "Gán hoa_don_chi_tiet_lo, trừ tồn lô"],
        ["3", "GHN", "GhnOrderSyncService.apDungTrangThaiGhn", "Map status GHN → enum nội bộ"],
        ["4", "Hủy online", "OnlineOrderLifecycleService.huyDonOnline", "Hoàn kho / voucher tùy trạng thái TT"],
        ["5", "Realtime", "OrderRealtimeService.publishStatusChanged", "Đẩy chuông khách + badge admin"],
    ],
    "models": [
        ["TrangThaiDonHang", "Enum", "common/enums/TrangThaiDonHang.java", "8 trạng thái đơn"],
        ["HoaDonChuyenTrangThaiRequest", "Request", "order/model/request/HoaDonChuyenTrangThaiRequest.java", "Body PATCH status + ghi chú"],
        ["XacNhanDonGanLoRequest", "Request", "order/model/request/XacNhanDonGanLoRequest.java", "Map dòng HĐ → lô"],
        ["GoiYGanLoResponse", "Response", "order/model/response/GoiYGanLoResponse.java", "JSON gợi ý lô"],
        ["HoaDonTuChoiRequest", "Request", "order/model/request/HoaDonTuChoiRequest.java", "Lý do từ chối"],
        ["HoaDonGhnWebhookRequest", "Request", "order/model/request/HoaDonGhnWebhookRequest.java", "Body sync/giả lập GHN"],
        ["GhnTrangThaiOptionResponse", "Response", "order/model/response/GhnTrangThaiOptionResponse.java", "1 option TT GHN"],
        ["LichSuDonHang", "Entity", "order/entity/LichSuDonHang.java", "Log chuyển trạng thái"],
        ["LichSuDonHangRequest", "Request", "order/model/request/LichSuDonHangRequest.java", "CRUD tạo log"],
        ["LichSuDonHangCapNhatRequest", "Request", "order/model/request/LichSuDonHangCapNhatRequest.java", "CRUD sửa log"],
        ["LichSuDonHangResponse", "Response", "order/model/response/LichSuDonHangResponse.java", "JSON 1 mốc timeline"],
        ["HuyDonOnlineRequest", "Request", "order/model/request/HuyDonOnlineRequest.java", "Lý do hủy phía khách"],
        ["HoaDonChiTietLo", "Entity", "order/entity/HoaDonChiTietLo.java", "Gán lô vào dòng HĐ"],
    ],
    "tables": [
        ["hoa_don.trang_thai", "Trạng thái hiện tại", "VARCHAR enum string", ""],
        ["lich_su_don_hang", "Audit trail", "V1", "Mỗi lần chuyển 1 dòng"],
        ["hoa_don_chi_tiet_lo", "Lô khi xác nhận", "V1 patch", "Trừ tồn theo lô"],
    ],
    "backend": [
        ["Controller", "order/controller/HoaDonController.java", "status / gán lô / từ chối / GHN / hủy khách", ""],
        ["Controller", "order/controller/LichSuDonHangController.java", "Timeline CRUD", ""],
        ["Controller", "order/controller/OnlineCheckoutController.java", "cancel", ""],
        ["Service", "order/service/HoaDonService.java", "chuyenTrangThai, xacNhanGanLo, tuChoiDon", ""],
        ["Service", "order/service/HoaDonStorefrontService.java", "huyDonCuaToi", ""],
        ["Service", "order/service/LichSuDonHangService.java", "getByHoaDon", ""],
        ["Service", "order/service/OnlineOrderLifecycleService.java", "huyDonOnline", ""],
        ["Service", "order/service/PosOrderLifecycleService.java", "Hoàn thành/hủy POS VNPay", ""],
        ["Service", "order/service/GhnOrderSyncService.java", "apDungTrangThaiGhn", ""],
        ["Service", "order/service/GhnTrackingService.java", "Map mã GHN", ""],
        ["Repo", "order/repository/HoaDonRepository.java", "Lưu TT", ""],
        ["Repo", "order/repository/LichSuDonHangRepository.java", "Log", ""],
        ["Repo", "order/repository/HoaDonChiTietLoRepository.java", "Lô", ""],
        ["Entity", "order/entity/HoaDon.java", "trang_thai", ""],
        ["Entity", "order/entity/LichSuDonHang.java", "Log", ""],
        ["Entity", "order/entity/HoaDonChiTietLo.java", "Lô", ""],
        ["DTO", "order/model/request/HoaDonChuyenTrangThaiRequest.java", "PATCH", ""],
        ["DTO", "order/model/request/XacNhanDonGanLoRequest.java", "Gán lô", ""],
        ["DTO", "order/model/request/HoaDonTuChoiRequest.java", "Từ chối", ""],
        ["DTO", "order/model/request/HoaDonGhnWebhookRequest.java", "GHN", ""],
        ["DTO", "order/model/request/HuyDonOnlineRequest.java", "Hủy", ""],
        ["DTO", "order/model/request/LichSuDonHangRequest.java", "Tạo log", ""],
        ["DTO", "order/model/request/LichSuDonHangCapNhatRequest.java", "Sửa log", ""],
        ["DTO", "order/model/response/GoiYGanLoResponse.java", "Gợi ý lô", ""],
        ["DTO", "order/model/response/GhnTrangThaiOptionResponse.java", "Option GHN", ""],
        ["DTO", "order/model/response/LichSuDonHangResponse.java", "Timeline JSON", ""],
        ["Enum", "common/enums/TrangThaiDonHang.java", "8 TT", ""],
        ["Realtime", "realtime/service/OrderRealtimeService.java", "publishStatusChanged", ""],
        ["Realtime", "realtime/event/OrderRealtimeAppEvent.java", "Event nội bộ", ""],
        ["Realtime", "realtime/model/OrderRealtimeEvent.java", "Payload WS", ""],
        ["Realtime", "realtime/listener/RealtimeEventListener.java", "Đẩy STOMP", ""],
        ["Noti", "notification/listener/CustomerNotificationListener.java", "Chuông khách", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "admin/hoa-don, chi-tiet, tra-cuu-don", ""],
        ["View", "views/admin/orders/OrderList.vue", "List + tab TT", ""],
        ["View", "views/admin/orders/HoaDonDetailPage.vue", "Action chuyển TT", ""],
        ["View", "views/storefront/TraCuuDon.vue", "Khách theo dõi / hủy", ""],
        ["UI", "components/admin/orders/GanLoDonHangModal.vue", "Xác nhận gán lô", ""],
        ["UI", "components/storefront/OrderCard.vue", "Card đơn + TT", ""],
        ["API", "api/hoaDonApi.js", "changeStatus, xacNhanGanLo, …", ""],
        ["API", "api/donHangApi.js", "hủy / tra cứu khách", ""],
        ["API", "api/onlineCheckout.js", "cancel", ""],
        ["Util", "utils/orderStatus.js", "Label / màu TT", ""],
        ["Badge", "composables/useAdminBadges.js", "Badge chờ xác nhận", ""],
        ["Menu", "constants/adminMenu.js", "Menu hóa đơn", ""],
    ],
})

FEATURES.append({
    "sheet": "05. Tim kiem loc don",
    "title": "05. ORDER — Tìm kiếm & lọc đơn hàng",
    "subtitle": "Màn admin /admin/hoa-don. Filter: keyword, loaiDon, trangThai, from/to, tab counts.",
    "package": "order",
    "note": "searchVisibleForAdmin ẩn đơn POS trạng thái CHO (đơn chờ).",
    "fe_entry": [
        ["Trang list", "/admin/hoa-don", "views/admin/orders/OrderList.vue", "Ô tìm, filter, tab All / Chờ xác nhận"],
        ["Menu", "Sidebar admin", "constants/adminMenu.js", "Mục Hóa đơn"],
    ],
    "apis": [
        ["GET /api/hoa-don/search", "HoaDonController.search()", "order/controller/HoaDonController.java", "keyword, loaiDon, trangThai, from, to, page, size"],
        ["GET /api/hoa-don/admin-counts", "HoaDonController.adminCounts()", "order/controller/HoaDonController.java", "Số lượng theo tab"],
        ["GET /api/hoa-don/paging", "HoaDonController.paging()", "order/controller/HoaDonController.java", "Phân trang thô"],
        ["GET /api/hoa-don", "HoaDonController.getAll()", "order/controller/HoaDonController.java", "List admin (ẩn đơn CHO)"],
    ],
    "flow": [
        ["1", "Controller", "HoaDonController.search / adminCounts", "Nhận query params"],
        ["2", "Service", "HoaDonService.searchAdmin / adminTabCounts / timKiem / phanTrang", "Ghép điều kiện"],
        ["3", "Repo", "HoaDonRepository.searchVisibleForAdmin", "JPQL/filter visible admin"],
    ],
    "models": [
        ["HoaDonResponse", "Response", "order/model/response/HoaDonResponse.java", "1 dòng list: mã, KH, tiền, TT, loại đơn, ngày"],
        ["HoaDon", "Entity", "order/entity/HoaDon.java", "Filter maHoaDon, loaiDon, trangThai, ngayTao"],
        ["TrangThaiDonHang", "Enum", "common/enums/TrangThaiDonHang.java", "Giá trị filter trạng thái"],
    ],
    "tables": [
        ["hoa_don", "Nguồn lọc", "ma_hoa_don, loai_don, trang_thai, ngay_tao", "ONLINE / TAI_QUAY"],
    ],
    "backend": [
        ["Controller", "order/controller/HoaDonController.java", "search, admin-counts, paging, getAll", ""],
        ["Service", "order/service/HoaDonService.java", "searchAdmin, adminTabCounts, timKiem, phanTrang", ""],
        ["Repo", "order/repository/HoaDonRepository.java", "searchVisibleForAdmin, countVisibleForAdmin", ""],
        ["Entity", "order/entity/HoaDon.java", "Các cột filter", ""],
        ["DTO", "order/model/response/HoaDonResponse.java", "JSON list", ""],
        ["Enum", "common/enums/TrangThaiDonHang.java", "Filter TT", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "path admin/hoa-don", ""],
        ["Menu", "constants/adminMenu.js", "Link menu", ""],
        ["View", "views/admin/orders/OrderList.vue", "UI filter + bảng", ""],
        ["API", "api/hoaDonApi.js", "searchHoaDon, getHoaDonAdminCounts", ""],
        ["Util", "utils/orderStatus.js", "Label tab/filter", ""],
        ["Badge", "composables/useAdminBadges.js", "Số chờ xác nhận trên header", ""],
        ["UI", "components/admin/AdminHeader.vue", "Badge", ""],
    ],
})

FEATURES.append({
    "sheet": "06. Chi tiet don",
    "title": "06. ORDER — Xem chi tiết đơn hàng",
    "subtitle": "3 kênh: admin theo id, khách login cua-toi, public token hoặc mã+SĐT.",
    "package": "order",
    "note": "Public tra cứu có PublicOrderLookupRateLimiter. Guest theo dõi bằng tracking_token (V8), không nhét email lên URL.",
    "fe_entry": [
        ["Admin", "/admin/hoa-don/chi-tiet/:id", "views/admin/orders/HoaDonDetailPage.vue", "Full: dòng, TT, GHN, in"],
        ["Admin rỗng", "/admin/hoa-don/chi-tiet", "views/admin/orders/HoaDonEmptyPage.vue", "Chưa chọn id"],
        ["Khách", "/tra-cuu-don", "views/storefront/TraCuuDon.vue", "List + lookup + card"],
        ["Redirect cũ", "/don-hang → /tra-cuu-don", "router/index.js", "Giữ query"],
    ],
    "apis": [
        ["GET /api/hoa-don/{id}", "HoaDonController.detail()", "order/controller/HoaDonController.java", "Chi tiết admin"],
        ["GET /api/hoa-don/cua-toi", "HoaDonController.donCuaToi()", "order/controller/HoaDonController.java", "List đơn khách login"],
        ["GET /api/hoa-don/cua-toi/{id}", "HoaDonController.chiTietCuaToi()", "order/controller/HoaDonController.java", "Chi tiết khách"],
        ["GET /api/hoa-don/tra-cuu?token=", "HoaDonController.traCuuBangToken()", "order/controller/HoaDonController.java", "Guest bằng tracking_token"],
        ["POST /api/hoa-don/tra-cuu", "HoaDonController.traCuuCongKhai()", "order/controller/HoaDonController.java", "Mã HĐ + SĐT"],
        ["GET /api/online/orders", "OnlineCheckoutController.danhSachDonHang()", "order/controller/OnlineCheckoutController.java", "List qua checkout API"],
        ["GET /api/online/orders/{idHoaDon}", "OnlineCheckoutController.chiTietDonHang()", "order/controller/OnlineCheckoutController.java", "Chi tiết checkout API"],
        ["GET /api/hoa-don-chi-tiet/{id}", "HoaDonChiTietController.detail()", "order/controller/HoaDonChiTietController.java", "1 dòng hàng"],
        ["GET /api/lich-su-don-hang/hoa-don", "LichSuDonHangController.getByHoaDon()", "order/controller/LichSuDonHangController.java", "Timeline trên trang chi tiết"],
    ],
    "flow": [
        ["1 Admin", "HoaDonService.detail", "HoaDonDetailResponse", "Header + dòng + TT + GHN"],
        ["2 Khách", "HoaDonStorefrontService.donCuaToi / chiTietCuaToi", "Chỉ đơn của JWT", ""],
        ["3 Token", "HoaDonStorefrontService.traCuuBangToken", "tracking_token", ""],
        ["4 Public", "traCuuCongKhai + PublicOrderLookupRateLimiter", "Chống spam dò đơn", ""],
    ],
    "models": [
        ["HoaDonDetailResponse", "Response", "order/model/response/HoaDonDetailResponse.java", "JSON full admin"],
        ["HoaDonChiTietResponse", "Response", "order/model/response/HoaDonChiTietResponse.java", "JSON 1 dòng hàng"],
        ["HoaDonResponse", "Response", "order/model/response/HoaDonResponse.java", "Tóm tắt"],
        ["StorefrontOrderSummaryResponse", "Response", "order/model/response/StorefrontOrderSummaryResponse.java", "Card list khách"],
        ["StorefrontOrderDetailResponse", "Response", "order/model/response/StorefrontOrderDetailResponse.java", "Chi tiết khách"],
        ["StorefrontOrderLineResponse", "Response", "order/model/response/StorefrontOrderLineResponse.java", "1 dòng storefront"],
        ["TraCuuDonRequest", "Request", "order/model/request/TraCuuDonRequest.java", "Body public: mã HĐ + SĐT"],
        ["HoaDonChiTietRequest", "Request", "order/model/request/HoaDonChiTietRequest.java", "CRUD dòng (ít dùng UI)"],
        ["HoaDonRequest", "Request", "order/model/request/HoaDonRequest.java", "CRUD HĐ generic"],
        ["HoaDon", "Entity", "order/entity/HoaDon.java", "Header + tracking_token + email"],
        ["HoaDonChiTiet", "Entity", "order/entity/HoaDonChiTiet.java", "Dòng"],
        ["HoaDonChiTietLo", "Entity", "order/entity/HoaDonChiTietLo.java", "Lô trên dòng"],
        ["ThanhToanHoaDon", "Entity", "order/entity/ThanhToanHoaDon.java", "GD TT hiện trên chi tiết"],
    ],
    "tables": [
        ["hoa_don", "Header", "tracking_token (V8), email_nguoi_nhan (V5)", ""],
        ["hoa_don_chi_tiet", "Dòng", "V1", ""],
        ["hoa_don_chi_tiet_lo", "Lô", "V1 patch", ""],
        ["thanh_toan_hoa_don", "GD", "V1", ""],
        ["lich_su_don_hang", "Timeline UI", "V1", ""],
    ],
    "backend": [
        ["Controller", "order/controller/HoaDonController.java", "detail / cua-toi / tra-cuu", ""],
        ["Controller", "order/controller/HoaDonChiTietController.java", "GET dòng", ""],
        ["Controller", "order/controller/OnlineCheckoutController.java", "orders list/detail", ""],
        ["Controller", "order/controller/LichSuDonHangController.java", "timeline", ""],
        ["Service", "order/service/HoaDonService.java", "detail admin", ""],
        ["Service", "order/service/HoaDonStorefrontService.java", "cua-toi + tra cứu", ""],
        ["Service", "order/service/HoaDonChiTietService.java", "detail dòng", ""],
        ["Service", "order/service/OnlineCheckoutService.java", "chiTietDonHang", ""],
        ["Service", "order/service/PublicOrderLookupRateLimiter.java", "Rate limit public", ""],
        ["Service", "order/service/OrderTrackingTokenGenerator.java", "Sinh token", ""],
        ["Repo", "order/repository/HoaDonRepository.java", "", ""],
        ["Repo", "order/repository/HoaDonChiTietRepository.java", "", ""],
        ["Repo", "order/repository/HoaDonChiTietLoRepository.java", "", ""],
        ["Repo", "order/repository/ThanhToanHoaDonRepository.java", "", ""],
        ["Repo", "order/repository/LichSuDonHangRepository.java", "", ""],
        ["Entity", "order/entity/HoaDon.java", "", ""],
        ["Entity", "order/entity/HoaDonChiTiet.java", "", ""],
        ["Entity", "order/entity/HoaDonChiTietLo.java", "", ""],
        ["Entity", "order/entity/ThanhToanHoaDon.java", "", ""],
        ["Entity", "order/entity/LichSuDonHang.java", "", ""],
        ["DTO", "order/model/response/HoaDonDetailResponse.java", "Admin full", ""],
        ["DTO", "order/model/response/HoaDonChiTietResponse.java", "Dòng", ""],
        ["DTO", "order/model/response/HoaDonResponse.java", "Tóm tắt", ""],
        ["DTO", "order/model/response/StorefrontOrderDetailResponse.java", "Khách", ""],
        ["DTO", "order/model/response/StorefrontOrderSummaryResponse.java", "List khách", ""],
        ["DTO", "order/model/response/StorefrontOrderLineResponse.java", "Dòng khách", ""],
        ["DTO", "order/model/response/LichSuDonHangResponse.java", "Timeline", ""],
        ["DTO", "order/model/request/TraCuuDonRequest.java", "Public", ""],
        ["DTO", "order/model/request/HoaDonChiTietRequest.java", "CRUD dòng", ""],
        ["DTO", "order/model/request/HoaDonRequest.java", "CRUD HĐ", ""],
        ["SQL", "db/migration/V5__hoa_don_email_nguoi_nhan.sql", "email", ""],
        ["SQL", "db/migration/V8__hoa_don_tracking_token.sql", "token", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "chi-tiet/:id, tra-cuu-don, redirect don-hang", ""],
        ["View", "views/admin/orders/HoaDonDetailPage.vue", "Chi tiết admin", ""],
        ["View", "views/admin/orders/HoaDonEmptyPage.vue", "Chưa chọn đơn", ""],
        ["View", "views/admin/orders/OrderList.vue", "Click vào chi tiết", ""],
        ["View", "views/storefront/TraCuuDon.vue", "Khách + public", ""],
        ["View", "views/storefront/DatHang.vue", "Sau đặt hiện mã/token", ""],
        ["UI", "components/storefront/OrderCard.vue", "Card đơn", ""],
        ["UI", "components/storefront/AccountSidebar.vue", "Link đơn hàng", ""],
        ["UI", "components/storefront/TheNavbar.vue", "Link tra cứu", ""],
        ["API", "api/hoaDonApi.js", "getHoaDonDetail, getLichSu", ""],
        ["API", "api/donHangApi.js", "cua-toi, tra-cuu token/public", ""],
        ["API", "api/onlineCheckout.js", "orders/:id", ""],
        ["Util", "utils/orderStatus.js", "Label TT", ""],
    ],
})

FEATURES.append({
    "sheet": "07. In hoa don",
    "title": "07. ORDER — In hóa đơn bán hàng",
    "subtitle": "KHÔNG có API PDF / Jasper / iText. In bằng window.print() + CSS @media print.",
    "package": "order (đọc) + frontend print",
    "note": "Hội đồng hỏi in PDF: trả lời thẳng là in trình duyệt, dữ liệu lấy GET /api/hoa-don/{id}.",
    "fe_entry": [
        ["Admin", "/admin/hoa-don/chi-tiet/:id — nút In hóa đơn", "views/admin/orders/HoaDonDetailPage.vue", "printInvoice() → #hoa-don-print-area"],
        ["POS", "/admin/pos — nút In hóa đơn", "views/admin/pos/PosPage.vue", "In sau khi bán tại quầy"],
    ],
    "apis": [
        ["GET /api/hoa-don/{id}", "HoaDonController.detail()", "order/controller/HoaDonController.java", "Nguồn dữ liệu in admin (không API print riêng)"],
        ["POST /api/ban-hang/tai-quay", "BanHangController.taoDonTaiQuay()", "order/controller/BanHangController.java", "POS in luôn payload vừa bán"],
    ],
    "flow": [
        ["1", "Load data", "HoaDonService.detail → HoaDonDetailResponse", "Giống xem chi tiết"],
        ["2", "In", "FE window.print() + @media print", "Ẩn chrome admin, chỉ vùng in"],
    ],
    "models": [
        ["HoaDonDetailResponse", "Response", "order/model/response/HoaDonDetailResponse.java", "Nguồn in admin: mã, KH, dòng, tổng, PTTT"],
        ["HoaDonChiTietResponse", "Response", "order/model/response/HoaDonChiTietResponse.java", "Dòng trên mẫu in"],
        ["BanHangHoaDonResponse", "Response", "order/model/response/BanHangHoaDonResponse.java", "Nguồn in POS"],
    ],
    "tables": [
        ["hoa_don", "Đọc header", "Không bảng in riêng", ""],
        ["hoa_don_chi_tiet", "Đọc dòng", "", ""],
    ],
    "backend": [
        ["Controller", "order/controller/HoaDonController.java", "GET detail", "Không endpoint /print"],
        ["Controller", "order/controller/BanHangController.java", "Tạo HĐ POS để in", ""],
        ["Service", "order/service/HoaDonService.java", "detail", ""],
        ["Service", "order/service/BanHangService.java", "taoDonTaiQuay", ""],
        ["DTO", "order/model/response/HoaDonDetailResponse.java", "Data in", ""],
        ["DTO", "order/model/response/HoaDonChiTietResponse.java", "Dòng in", ""],
        ["DTO", "order/model/response/BanHangHoaDonResponse.java", "In POS", ""],
    ],
    "frontend": [
        ["View", "views/admin/orders/HoaDonDetailPage.vue", "printInvoice, #hoa-don-print-area, @media print", "Chính"],
        ["View", "views/admin/pos/PosPage.vue", "Nút in sau bán", ""],
        ["CSS", "styles/posAdmin.css", "Hỗ trợ in POS", ""],
        ["API", "api/hoaDonApi.js", "getHoaDonDetail", ""],
        ["API", "api/banHangApi.js", "Payload POS", ""],
    ],
})

FEATURES.append({
    "sheet": "08. Dot giam gia",
    "title": "08. PROMOTION — Quản lý chương trình khuyến mãi",
    "subtitle": "Trong code = Đợt giảm giá (DotGiamGia), API /api/sale. Không có entity ChuongTrinh.",
    "package": "voucher",
    "note": "DiscountList.vue tồn tại nhưng KHÔNG gắn router — đừng nhắc khi vấn đáp.",
    "fe_entry": [
        ["List đợt", "/admin/sale", "views/admin/voucher/SaleView.vue", "CRUD đợt"],
        ["Chi tiết đợt", "/admin/sale/:id", "views/admin/voucher/saleDetail.vue", "Gắn SP / mức giảm"],
        ["Storefront KM", "/san-pham/khuyen-mai", "views/storefront/ProductList.vue", "List SP đang sale"],
    ],
    "apis": [
        ["GET /api/sale", "DotGiamGiaController.hienThiDanhSach()", "voucher/controller/DotGiamGiaController.java", "List đợt"],
        ["GET /api/sale/search", "DotGiamGiaController.search()", "voucher/controller/DotGiamGiaController.java", "Tìm đợt"],
        ["GET /api/sale/next-ma", "DotGiamGiaController.nextMa()", "voucher/controller/DotGiamGiaController.java", "Sinh mã đợt"],
        ["GET /api/sale/{id}", "DotGiamGiaController.detail()", "voucher/controller/DotGiamGiaController.java", "Chi tiết đợt"],
        ["POST /api/sale", "DotGiamGiaController.add()", "voucher/controller/DotGiamGiaController.java", "Tạo đợt"],
        ["PUT /api/sale/{id}", "DotGiamGiaController.update()", "voucher/controller/DotGiamGiaController.java", "Sửa đợt"],
        ["DELETE /api/sale/{id}", "DotGiamGiaController.delete()", "voucher/controller/DotGiamGiaController.java", "Xóa đợt"],
        ["PUT /api/sale/{id}/stop", "DotGiamGiaController.stop()", "voucher/controller/DotGiamGiaController.java", "Ngưng"],
        ["PUT /api/sale/{id}/activate", "DotGiamGiaController.activate()", "voucher/controller/DotGiamGiaController.java", "Bật lại"],
        ["GET /api/sale/{id}/products", "DotGiamGiaController.getProducts()", "voucher/controller/DotGiamGiaController.java", "SP trong đợt"],
        ["POST /api/sale/{id}/products", "DotGiamGiaController.addProduct()", "voucher/controller/DotGiamGiaController.java", "Gắn SP"],
        ["PUT /api/sale/{id}/products/{detailId}", "DotGiamGiaController.updateProduct()", "voucher/controller/DotGiamGiaController.java", "Sửa %/giá dòng"],
        ["DELETE /api/sale/{id}/products/{detailId}", "DotGiamGiaController.deleteProduct()", "voucher/controller/DotGiamGiaController.java", "Gỡ SP"],
        ["GET /api/chi-tiet-dot-giam-gia", "ChiTietDotGiamGiaController.hienThi…()", "voucher/controller/ChiTietDotGiamGiaController.java", "List dòng đợt"],
        ["GET /api/chi-tiet-dot-giam-gia/{id}", "ChiTietDotGiamGiaController.detail()", "voucher/controller/ChiTietDotGiamGiaController.java", "Chi tiết dòng"],
        ["POST /api/chi-tiet-dot-giam-gia", "ChiTietDotGiamGiaController.add()", "voucher/controller/ChiTietDotGiamGiaController.java", "Thêm dòng"],
        ["PUT /api/chi-tiet-dot-giam-gia/{id}", "ChiTietDotGiamGiaController.update()", "voucher/controller/ChiTietDotGiamGiaController.java", "Sửa dòng"],
        ["DELETE /api/chi-tiet-dot-giam-gia/{id}", "ChiTietDotGiamGiaController.delete()", "voucher/controller/ChiTietDotGiamGiaController.java", "Xóa dòng"],
    ],
    "flow": [
        ["1 Admin", "DotGiamGiaController → DotGiamGiaService / ChiTietDotGiamGiaService", "CRUD đợt + SP", ""],
        ["2 Bán hàng", "CheckoutPricingService + getActiveSaleByVariantId", "Giá KM đè giá gốc", "Giỏ / checkout / POS"],
    ],
    "models": [
        ["DotGiamGia", "Entity", "voucher/entity/DotGiamGia.java", "Đợt: mã, tên, từ/đến, isActive"],
        ["ChiTietDotGiamGia", "Entity", "voucher/entity/ChiTietDotGiamGia.java", "SP/biến thể trong đợt + mức giảm"],
        ["DotGiamGiaRequest", "Request", "voucher/model/request/DotGiamGiaRequest.java", "Body tạo/sửa đợt"],
        ["DotGiamGiaResponse", "Response", "voucher/model/response/DotGiamGiaResponse.java", "JSON đợt"],
        ["ChiTietDotGiamGiaRequest", "Request", "voucher/model/request/ChiTietDotGiamGiaRequest.java", "Body gắn/sửa SP"],
        ["ChiTietDotGiamGiaResponse", "Response", "voucher/model/response/ChiTietDotGiamGiaResponse.java", "JSON 1 dòng SP trong đợt"],
        ["VariantSaleInfo", "DTO", "voucher/model/response/VariantSaleInfo.java", "Giá KM đang active cho 1 biến thể"],
        ["SalePriceAgg", "Projection", "voucher/repository/SalePriceAgg.java", "Aggregate query giá sale"],
        ["VariantSalePriceRow", "Projection", "voucher/repository/VariantSalePriceRow.java", "Row giá sale theo biến thể"],
    ],
    "tables": [
        ["dot_giam_gia", "Header đợt", "V1, is_active", ""],
        ["chi_tiet_dot_giam_gia", "SP trong đợt", "V1", ""],
    ],
    "backend": [
        ["Controller", "voucher/controller/DotGiamGiaController.java", "CRUD /api/sale", ""],
        ["Controller", "voucher/controller/ChiTietDotGiamGiaController.java", "CRUD dòng", ""],
        ["Service", "voucher/service/DotGiamGiaService.java", "CRUD + getActiveSaleByVariantId", ""],
        ["Service", "voucher/service/ChiTietDotGiamGiaService.java", "CRUD SP trong đợt", ""],
        ["Repo", "voucher/repository/DotGiamGiaRepository.java", "", ""],
        ["Repo", "voucher/repository/ChiTietDotGiamGiaRepository.java", "", ""],
        ["Projection", "voucher/repository/SalePriceAgg.java", "Agg giá", ""],
        ["Projection", "voucher/repository/VariantSalePriceRow.java", "Row giá", ""],
        ["Entity", "voucher/entity/DotGiamGia.java", "dot_giam_gia", ""],
        ["Entity", "voucher/entity/ChiTietDotGiamGia.java", "chi_tiet_dot_giam_gia", ""],
        ["DTO", "voucher/model/request/DotGiamGiaRequest.java", "", ""],
        ["DTO", "voucher/model/request/ChiTietDotGiamGiaRequest.java", "", ""],
        ["DTO", "voucher/model/response/DotGiamGiaResponse.java", "", ""],
        ["DTO", "voucher/model/response/ChiTietDotGiamGiaResponse.java", "", ""],
        ["DTO", "voucher/model/response/VariantSaleInfo.java", "Dùng giỏ/checkout", ""],
        ["Phụ bán", "order/service/CheckoutPricingService.java", "Áp giá sale lúc tính tiền", ""],
        ["Phụ giỏ", "cart/service/GioHangService.java", "Hiện giá KM trên giỏ", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Tạo 2 bảng", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "admin/sale, admin/sale/:id, san-pham/khuyen-mai", ""],
        ["View", "views/admin/voucher/SaleView.vue", "List đợt", ""],
        ["View", "views/admin/voucher/saleDetail.vue", "SP trong đợt", ""],
        ["View", "views/storefront/ProductList.vue", "Trang KM", ""],
        ["API", "api/saleApi.js", "Toàn bộ /api/sale", ""],
        ["CSS", "styles/saleCss.css", "", ""],
        ["CSS", "styles/saleDetailCss.css", "", ""],
        ["Menu", "constants/adminMenu.js", "Đợt giảm giá", ""],
        ["KHÔNG DÙNG", "views/admin/discount/DiscountList.vue", "Orphan — không gắn router", "Đừng nhắc vấn đáp"],
    ],
})

FEATURES.append({
    "sheet": "09. Voucher",
    "title": "09. PROMOTION — Quản lý voucher / coupon code",
    "subtitle": "CRUD /api/vouchers. Áp lúc checkout qua CheckoutPricingService. Phạm vi CÔNG KHAI / CÁ NHÂN.",
    "package": "voucher + common.entity.PhieuGiamGia",
    "note": "Loại: PHAN_TRAM, TIEN_MAT, FREE_SHIP. Gán khách: khach_hang_phieu_giam_gia (V4). Điểm min/max: V9.",
    "fe_entry": [
        ["Admin", "/admin/voucher", "views/admin/voucher/VoucherView.vue", "CRUD + gán khách"],
        ["Khách ví voucher", "/tai-khoan?section=vouchers", "views/storefront/TaiKhoan.vue", "Xem phiếu của tôi"],
        ["Checkout", "Modal /dat-hang", "components/storefront/CheckoutVoucherModal.vue", "Áp mã lúc đặt"],
        ["POS", "Modal /admin/pos", "components/admin/PosVoucherModal.vue", "Áp tại quầy"],
    ],
    "apis": [
        ["GET /api/vouchers", "PhieuGiamGiaController.getAllPhieuGiamGia()", "voucher/controller/PhieuGiamGiaController.java", "List"],
        ["GET /api/vouchers/stats", "PhieuGiamGiaController.getStats()", "voucher/controller/PhieuGiamGiaController.java", "Thống kê dashboard"],
        ["GET /api/vouchers/next-ma", "PhieuGiamGiaController.nextMa()", "voucher/controller/PhieuGiamGiaController.java", "Sinh mã"],
        ["GET /api/vouchers/{id}", "PhieuGiamGiaController.getById()", "voucher/controller/PhieuGiamGiaController.java", "Chi tiết"],
        ["POST /api/vouchers", "PhieuGiamGiaController.create()", "voucher/controller/PhieuGiamGiaController.java", "Tạo"],
        ["PUT /api/vouchers/{id}", "PhieuGiamGiaController.update()", "voucher/controller/PhieuGiamGiaController.java", "Sửa"],
        ["DELETE /api/vouchers/{id}", "PhieuGiamGiaController.delete()", "voucher/controller/PhieuGiamGiaController.java", "Xóa"],
        ["PUT /api/vouchers/{id}/stop", "PhieuGiamGiaController.stop()", "voucher/controller/PhieuGiamGiaController.java", "Ngưng"],
        ["PUT /api/vouchers/{id}/activate", "PhieuGiamGiaController.activate()", "voucher/controller/PhieuGiamGiaController.java", "Bật"],
        ["GET /api/vouchers/search", "PhieuGiamGiaController.search()", "voucher/controller/PhieuGiamGiaController.java", "Tìm"],
        ["GET /api/vouchers/{id}/khach-hang", "PhieuGiamGiaController.danhSachKhachDaGan()", "voucher/controller/PhieuGiamGiaController.java", "Khách đã gán CA_NHAN"],
        ["POST /api/vouchers/{id}/khach-hang", "PhieuGiamGiaController.ganChoKhachHang()", "voucher/controller/PhieuGiamGiaController.java", "Gán 1/n khách"],
        ["POST /api/vouchers/{id}/khach-hang/theo-nhom", "PhieuGiamGiaController.ganTheoNhom()", "voucher/controller/PhieuGiamGiaController.java", "Gán theo nhóm điểm"],
        ["POST /api/vouchers/khach-hang/theo-nhom/preview", "PhieuGiamGiaController.previewNhom()", "voucher/controller/PhieuGiamGiaController.java", "Preview nhóm"],
        ["DELETE /api/vouchers/{id}/khach-hang/{idKhachHang}", "PhieuGiamGiaController.boGan()", "voucher/controller/PhieuGiamGiaController.java", "Bỏ gán"],
        ["GET /api/khach-hang/toi/voucher/cong-khai", "KhachHangToiController.voucherCongKhai()", "customer/controller/KhachHangToiController.java", "Voucher public của khách"],
        ["GET /api/khach-hang/toi/voucher/ca-nhan", "KhachHangToiController.voucherCaNhan()", "customer/controller/KhachHangToiController.java", "Voucher đã gán"],
        ["GET /api/online/vouchers", "OnlineCheckoutController.danhSachVoucher()", "order/controller/OnlineCheckoutController.java", "Voucher lúc checkout"],
        ["GET /api/ban-hang/vouchers", "BanHangController.vouchers()", "order/controller/BanHangController.java", "Voucher POS"],
    ],
    "flow": [
        ["1 CRUD", "PhieuGiamGiaController → PhieuGiamGiaService", "Tạo/sửa/ngưng/thống kê", ""],
        ["2 Gán", "VoucherKhachHangService", "CA_NHAN → khach_hang_phieu_giam_gia", ""],
        ["3 Áp đơn", "CheckoutPricingService.tinhTienGiamPhieu + consumeOne", "Check HSD, SL, min đơn, phạm vi, điểm", ""],
    ],
    "models": [
        ["PhieuGiamGia", "Entity", "common/entity/PhieuGiamGia.java", "Mã, loại, giá trị, HSD, SL, min đơn, phạm vi, điểm"],
        ["KhachHangPhieuGiamGia", "Entity", "voucher/entity/KhachHangPhieuGiamGia.java", "Gán voucher cá nhân"],
        ["LoaiPhieuGiamGia", "Enum", "common/enums/LoaiPhieuGiamGia.java", "PHAN_TRAM / TIEN_MAT / FREE_SHIP"],
        ["PhamViPhieuGiamGia", "Enum", "common/enums/PhamViPhieuGiamGia.java", "CONG_KHAI / CA_NHAN"],
        ["PhieuGiamGiaRequest", "Request", "voucher/model/request/PhieuGiamGiaRequest.java", "Body tạo/sửa"],
        ["PhieuGiamGiaResponse", "Response", "voucher/model/response/PhieuGiamGiaResponse.java", "JSON voucher"],
        ["PhieuGiamGiaStatsResponse", "Response", "voucher/model/response/PhieuGiamGiaStatsResponse.java", "JSON thống kê"],
        ["GanVoucherRequest", "Request", "voucher/model/request/GanVoucherRequest.java", "Gán list idKhachHang"],
        ["GanVoucherTheoNhomRequest", "Request", "voucher/model/request/GanVoucherTheoNhomRequest.java", "Gán theo khoảng điểm"],
    ],
    "tables": [
        ["phieu_giam_gia", "Voucher", "V1; pham_vi V4; diem_toi_thieu/da V9", ""],
        ["khach_hang_phieu_giam_gia", "Gán cá nhân", "V4__voucher_ca_nhan.sql", ""],
        ["hoa_don.id_phieu_giam_gia", "Đơn đã áp", "FK", "consumeOne khi đặt"],
    ],
    "backend": [
        ["Controller", "voucher/controller/PhieuGiamGiaController.java", "CRUD + gán khách", ""],
        ["Service", "voucher/service/PhieuGiamGiaService.java", "CRUD, consumeOne, listAvailable*", ""],
        ["Service", "voucher/service/VoucherKhachHangService.java", "Gán / bỏ gán / nhóm", ""],
        ["Repo", "voucher/repository/PhieuGiamGiaRepository.java", "", ""],
        ["Repo", "voucher/repository/KhachHangPhieuGiamGiaRepository.java", "", ""],
        ["Entity", "common/entity/PhieuGiamGia.java", "phieu_giam_gia", ""],
        ["Entity", "voucher/entity/KhachHangPhieuGiamGia.java", "bảng gán", ""],
        ["Enum", "common/enums/LoaiPhieuGiamGia.java", "", ""],
        ["Enum", "common/enums/PhamViPhieuGiamGia.java", "", ""],
        ["DTO", "voucher/model/request/PhieuGiamGiaRequest.java", "", ""],
        ["DTO", "voucher/model/request/GanVoucherRequest.java", "", ""],
        ["DTO", "voucher/model/request/GanVoucherTheoNhomRequest.java", "", ""],
        ["DTO", "voucher/model/response/PhieuGiamGiaResponse.java", "", ""],
        ["DTO", "voucher/model/response/PhieuGiamGiaStatsResponse.java", "", ""],
        ["Khách", "customer/controller/KhachHangToiController.java", "voucher/cong-khai, ca-nhan", ""],
        ["Checkout", "order/controller/OnlineCheckoutController.java", "GET /vouchers", ""],
        ["POS", "order/controller/BanHangController.java", "GET vouchers", ""],
        ["Áp giá", "order/service/CheckoutPricingService.java", "tinhTienGiamPhieu", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Tạo phiếu", ""],
        ["SQL", "db/migration/V4__voucher_ca_nhan.sql", "pham_vi + bảng gán", ""],
        ["SQL", "db/migration/V7__backfill_phieu_giam_gia_pham_vi.sql", "Backfill CONG_KHAI", ""],
        ["SQL", "db/migration/V9__phieu_giam_gia_diem_range.sql", "điểm min/max", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "admin/voucher, tai-khoan", ""],
        ["Menu", "constants/adminMenu.js", "Phiếu giảm giá", ""],
        ["View", "views/admin/voucher/VoucherView.vue", "Trang quản lý", ""],
        ["UI", "components/voucher/VoucherTable.vue", "Bảng list", ""],
        ["UI", "components/voucher/VoucherCreateModal.vue", "Tạo/sửa", ""],
        ["UI", "components/voucher/VoucherToolBar.vue", "Search/filter", ""],
        ["UI", "components/voucher/VoucherStats.vue", "Thống kê", ""],
        ["UI", "components/voucher/DashboardStat.vue", "Thẻ số", ""],
        ["UI", "components/voucher/Pagination.vue", "Phân trang", ""],
        ["View", "views/storefront/TaiKhoan.vue", "Ví voucher khách", ""],
        ["UI", "components/storefront/VoucherCard.vue", "Card phiếu", ""],
        ["UI", "components/storefront/CheckoutVoucherModal.vue", "Áp lúc đặt", ""],
        ["UI", "components/admin/PosVoucherModal.vue", "Áp POS", ""],
        ["View", "views/storefront/DatHang.vue", "Gửi maPhieuGiamGia", ""],
        ["View", "views/admin/pos/PosPage.vue", "Gửi voucher POS", ""],
        ["API", "api/voucherApi.js", "CRUD admin", ""],
        ["API", "api/onlineCheckout.js", "List voucher checkout", ""],
        ["API", "api/banHangApi.js", "List voucher POS", ""],
        ["API", "api/khachHangApi.js", "voucher của tôi", ""],
        ["CSS", "styles/voucherCss.css", "", ""],
        ["CSS", "styles/checkoutVoucherModal.css", "", ""],
    ],
})

FEATURES.append({
    "sheet": "10. TB dat hang",
    "title": "10. NOTIFICATION — Thông báo đặt hàng thành công",
    "subtitle": "In-app CHO ADMIN (DON_HANG_MOI). Khách không có loại này — khách nhận EMAIL (sheet 12).",
    "package": "notification + realtime",
    "note": "Không có API gửi. Trigger trong OnlineCheckoutService (COD) và PaymentService (VNPay OK).",
    "fe_entry": [
        ["Chuông admin", "Header mọi trang /admin", "components/admin/AdminHeader.vue", "List + đánh dấu đọc"],
        ["Realtime", "STOMP /topic/admin/notifications", "composables/useRealtime.js", "Đẩy tức thì"],
    ],
    "apis": [
        ["GET /api/thong-bao", "ThongBaoController.danhSach()", "notification/controller/ThongBaoController.java", "List admin"],
        ["GET /api/thong-bao/chua-doc", "ThongBaoController.demChuaDoc()", "notification/controller/ThongBaoController.java", "Đếm chưa đọc"],
        ["POST /api/thong-bao/{id}/doc", "ThongBaoController.doc()", "notification/controller/ThongBaoController.java", "Đánh dấu 1"],
        ["POST /api/thong-bao/doc-tat-ca", "ThongBaoController.docTatCa()", "notification/controller/ThongBaoController.java", "Đánh dấu hết"],
    ],
    "flow": [
        ["1", "Checkout/TT OK", "OnlineCheckoutService.thongBaoDonMoi / PaymentService.thongBaoDonMoi", ""],
        ["2", "Tạo bản ghi", "ThongBaoService.taoThongBao(DON_HANG_MOI)", "Lưu thong_bao"],
        ["3", "Realtime", "AdminNotificationAppEvent → RealtimeEventListener", "STOMP admin"],
    ],
    "models": [
        ["ThongBao", "Entity", "notification/entity/ThongBao.java", "Bản ghi thông báo (admin: id_khach_hang null)"],
        ["LoaiThongBao.DON_HANG_MOI", "Enum value", "notification/enums/LoaiThongBao.java", "Loại đơn online mới"],
        ["ThongBaoResponse", "Response", "notification/model/response/ThongBaoResponse.java", "JSON 1 TB: title, nội dung, link, đã đọc"],
        ["ThongBaoTongQuanResponse", "Response", "notification/model/response/ThongBaoTongQuanResponse.java", "Tổng / chưa đọc"],
        ["AdminNotificationEvent", "WS payload", "realtime/model/AdminNotificationEvent.java", "Đẩy chuông"],
        ["AdminNotificationAppEvent", "Spring event", "realtime/event/AdminNotificationAppEvent.java", "Cầu nối service → WS"],
    ],
    "tables": [
        ["thong_bao", "In-app", "V1 + link, id_tham_chieu, ma_tham_chieu", "Admin row: id_khach_hang null"],
    ],
    "backend": [
        ["Controller", "notification/controller/ThongBaoController.java", "Đọc/đánh dấu admin", ""],
        ["Service", "notification/service/ThongBaoService.java", "taoThongBao", ""],
        ["Repo", "notification/repository/ThongBaoRepository.java", "", ""],
        ["Entity", "notification/entity/ThongBao.java", "", ""],
        ["Enum", "notification/enums/LoaiThongBao.java", "DON_HANG_MOI", ""],
        ["DTO", "notification/model/response/ThongBaoResponse.java", "", ""],
        ["DTO", "notification/model/response/ThongBaoTongQuanResponse.java", "", ""],
        ["Trigger", "order/service/OnlineCheckoutService.java", "COD xong", ""],
        ["Trigger", "payment/service/PaymentService.java", "VNPay OK", ""],
        ["Realtime", "realtime/event/AdminNotificationAppEvent.java", "", ""],
        ["Realtime", "realtime/model/AdminNotificationEvent.java", "", ""],
        ["Realtime", "realtime/listener/RealtimeEventListener.java", "", ""],
        ["Realtime", "realtime/config/WebSocketConfig.java", "STOMP", ""],
        ["Realtime", "realtime/config/JwtStompChannelInterceptor.java", "Auth WS", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Tạo thong_bao", ""],
    ],
    "frontend": [
        ["UI", "components/admin/AdminHeader.vue", "Chuông", ""],
        ["Layout", "layouts/AdminLayout.vue", "Mount header", ""],
        ["Composable", "composables/useAdminBadges.js", "Load TB + tách DON_HANG_MOI", ""],
        ["Composable", "composables/useRealtime.js", "subscribeAdminNotifications", ""],
        ["API", "api/thongBaoApi.js", "list / đọc", ""],
    ],
})

FEATURES.append({
    "sheet": "11. TB trang thai",
    "title": "11. NOTIFICATION — Thông báo cập nhật trạng thái đơn",
    "subtitle": "In-app CHO KHÁCH (DON_HANG_CAP_NHAT). Bỏ qua TRA_HANG (return có loại riêng).",
    "package": "notification.listener + realtime",
    "note": "Chỉ tạo khi OrderRealtimeAppEvent TYPE_STATUS_CHANGED và idKhachHang != null (guest không có chuông).",
    "fe_entry": [
        ["Chuông khách", "Navbar storefront", "components/storefront/TheNavbar.vue", "List + đếm chưa đọc"],
        ["Deep link", "/tra-cuu-don?ma=…", "views/storefront/TraCuuDon.vue", "Click TB mở đơn"],
        ["Realtime", "STOMP /topic/customers/{id}/notifications", "composables/useRealtime.js", ""],
    ],
    "apis": [
        ["GET /api/khach-hang/toi/thong-bao", "ThongBaoKhachController.danhSach()", "notification/controller/ThongBaoKhachController.java", "List khách"],
        ["GET /api/khach-hang/toi/thong-bao/chua-doc", "ThongBaoKhachController.demChuaDoc()", "notification/controller/ThongBaoKhachController.java", "Đếm"],
        ["POST /api/khach-hang/toi/thong-bao/{id}/doc", "ThongBaoKhachController.doc()", "notification/controller/ThongBaoKhachController.java", "Đọc 1"],
        ["POST /api/khach-hang/toi/thong-bao/doc-tat-ca", "ThongBaoKhachController.docTatCa()", "notification/controller/ThongBaoKhachController.java", "Đọc hết"],
    ],
    "flow": [
        ["1", "Đổi TT đơn", "HoaDonService / GhnOrderSync / OnlineOrderLifecycle", "Gọi publishStatusChanged"],
        ["2", "Event", "OrderRealtimeService.publishStatusChanged", "OrderRealtimeAppEvent"],
        ["3", "Listener", "CustomerNotificationListener.onOrderStatusChanged", "Skip TRA_HANG"],
        ["4", "Lưu + WS", "ThongBaoService.taoThongBaoKhach → CustomerNotificationAppEvent", "Chuông khách"],
    ],
    "models": [
        ["ThongBao", "Entity", "notification/entity/ThongBao.java", "Cùng bảng, có id_khach_hang"],
        ["LoaiThongBao.DON_HANG_CAP_NHAT", "Enum value", "notification/enums/LoaiThongBao.java", "Đổi trạng thái đơn"],
        ["ThongBaoResponse", "Response", "notification/model/response/ThongBaoResponse.java", "JSON chuông khách"],
        ["OrderRealtimeAppEvent", "Spring event", "realtime/event/OrderRealtimeAppEvent.java", "Đổi TT"],
        ["OrderRealtimeEvent", "WS payload", "realtime/model/OrderRealtimeEvent.java", "Đơn realtime"],
        ["CustomerNotificationAppEvent", "Spring event", "realtime/event/CustomerNotificationAppEvent.java", "Cầu nối chuông"],
        ["CustomerNotificationEvent", "WS payload", "realtime/model/CustomerNotificationEvent.java", "Đẩy navbar"],
    ],
    "tables": [
        ["thong_bao", "In-app khách", "id_khach_hang != null", ""],
    ],
    "backend": [
        ["Controller", "notification/controller/ThongBaoKhachController.java", "API chuông khách", ""],
        ["Listener", "notification/listener/CustomerNotificationListener.java", "Bắt đổi TT → tạo TB", "Lõi"],
        ["Service", "notification/service/ThongBaoService.java", "taoThongBaoKhach", ""],
        ["Entity", "notification/entity/ThongBao.java", "", ""],
        ["Enum", "notification/enums/LoaiThongBao.java", "DON_HANG_CAP_NHAT", ""],
        ["Repo", "notification/repository/ThongBaoRepository.java", "", ""],
        ["DTO", "notification/model/response/ThongBaoResponse.java", "", ""],
        ["Realtime", "realtime/service/OrderRealtimeService.java", "publishStatusChanged", ""],
        ["Realtime", "realtime/event/OrderRealtimeAppEvent.java", "", ""],
        ["Realtime", "realtime/event/CustomerNotificationAppEvent.java", "", ""],
        ["Realtime", "realtime/model/OrderRealtimeEvent.java", "", ""],
        ["Realtime", "realtime/model/CustomerNotificationEvent.java", "", ""],
        ["Realtime", "realtime/listener/RealtimeEventListener.java", "STOMP khách", ""],
        ["Realtime", "realtime/config/WebSocketConfig.java", "", ""],
        ["Realtime", "realtime/config/JwtStompChannelInterceptor.java", "", ""],
        ["Nguồn TT", "order/service/HoaDonService.java", "Đổi TT tay", ""],
        ["Nguồn TT", "order/service/GhnOrderSyncService.java", "Đổi TT theo GHN", ""],
        ["Nguồn TT", "order/service/OnlineOrderLifecycleService.java", "Hủy", ""],
    ],
    "frontend": [
        ["UI", "components/storefront/TheNavbar.vue", "Chuông", ""],
        ["Composable", "composables/useCustomerNotifications.js", "Load + đọc", ""],
        ["Composable", "composables/useRealtime.js", "subscribeCustomerNotifications", ""],
        ["View", "views/storefront/TraCuuDon.vue", "Deep link ma=", ""],
        ["API", "api/thongBaoApi.js", "getThongBaoKhach, docThongBaoKhach", ""],
    ],
})

FEATURES.append({
    "sheet": "12. Email xac nhan",
    "title": "12. NOTIFICATION — Gửi email xác nhận đơn hàng",
    "subtitle": "Event-driven. KHÔNG có REST send-mail. SMTP Gmail (JavaMailSender).",
    "package": "notification.event / listener / OrderMailService",
    "note": "To: hoa_don.email_nguoi_nhan, fallback khach_hang.email. Cùng service còn mail trả hàng (sheet 14).",
    "fe_entry": [
        ["Gián tiếp COD", "Hoàn tất /dat-hang", "views/storefront/DatHang.vue", "Không gọi API mail"],
        ["Gián tiếp VNPay", "Return /dat-hang?success=", "views/storefront/DatHang.vue", "Mail sau IPN/callback OK"],
    ],
    "apis": [
        ["(không public)", "publishEvent(DatHangThanhCongMailEvent)", "OnlineCheckoutService / PaymentService", "Trigger sau commit"],
    ],
    "flow": [
        ["1", "Commit đơn/TT", "OnlineCheckoutService hoặc PaymentService", "publish DatHangThanhCongMailEvent"],
        ["2", "Listener", "OrderMailEventListener.onDatHangThanhCong", "Reload HoaDon sau commit"],
        ["3", "Mail", "OrderMailService.guiHoaDonDatHangThanhCong", "HTML dòng hàng + gửi SMTP"],
        ["4", "Lỗi SMTP", "Log, không rollback đơn", "Mail fail không hủy đơn"],
    ],
    "models": [
        ["DatHangThanhCongMailEvent", "Spring event", "notification/event/DatHangThanhCongMailEvent.java", "Chứa idHoaDon"],
        ["HoaDon", "Entity (đọc)", "order/entity/HoaDon.java", "Mã đơn, tổng, email nhận"],
        ["HoaDonChiTiet", "Entity (đọc)", "order/entity/HoaDonChiTiet.java", "Dựng HTML dòng hàng"],
    ],
    "tables": [
        ["hoa_don", "Đọc", "email_nguoi_nhan (V5)", "Fallback email khách"],
        ["hoa_don_chi_tiet", "Đọc dòng cho HTML", "", ""],
    ],
    "backend": [
        ["Event", "notification/event/DatHangThanhCongMailEvent.java", "Payload idHoaDon", ""],
        ["Listener", "notification/listener/OrderMailEventListener.java", "onDatHangThanhCong", ""],
        ["Service", "notification/service/OrderMailService.java", "guiHoaDonDatHangThanhCong + HTML", "JavaMailSender"],
        ["Trigger", "order/service/OnlineCheckoutService.java", "COD", ""],
        ["Trigger", "payment/service/PaymentService.java", "VNPay OK", ""],
        ["Entity", "order/entity/HoaDon.java", "emailNguoiNhan", ""],
        ["Entity", "order/entity/HoaDonChiTiet.java", "Dòng HTML", ""],
        ["Config", "resources/application.properties", "spring.mail.* , app.mail.from-name=SUNOVA", "Gmail 587 STARTTLS"],
        ["SQL", "db/migration/V5__hoa_don_email_nguoi_nhan.sql", "Cột email", ""],
    ],
    "frontend": [
        ["View", "views/storefront/DatHang.vue", "Chỉ hoàn tất đặt; không UI soạn mail", ""],
        ["API", "api/onlineCheckout.js", "Checkout thành công → backend tự gửi", ""],
    ],
})

FEATURES.append({
    "sheet": "13. GHN Shipping",
    "title": "13. SHIPPING — Tính phí và tạo vận đơn GHN",
    "subtitle": "Master data + fee: package shipping. Tạo vận đơn theo HĐ: GhnOrderCreationService trong order.",
    "package": "shipping + order.service.Ghn*",
    "note": "Sandbox dev-online-gateway.ghn.vn. Fallback phí ghn.fallback-fee=30000. Địa chỉ 2 cấp (tỉnh/xã).",
    "fe_entry": [
        ["Tính phí", "/dat-hang", "views/storefront/DatHang.vue", "Chọn tỉnh/xã → POST /fee"],
        ["Modal địa chỉ", "Popup checkout", "components/storefront/CheckoutRecipientModal.vue", "Provinces / wards"],
        ["Tạo vận đơn", "/admin/hoa-don/chi-tiet/:id", "views/admin/orders/HoaDonDetailPage.vue", "Nút tạo / đồng bộ GHN"],
    ],
    "apis": [
        ["GET /api/shipping/provinces", "ShippingController.provinces()", "shipping/controller/ShippingController.java", "Tỉnh GHN (permitAll)"],
        ["GET /api/shipping/wards", "ShippingController.wards()", "shipping/controller/ShippingController.java", "Xã theo provinceId"],
        ["POST /api/shipping/fee", "ShippingController.fee()", "shipping/controller/ShippingController.java", "Tính phí"],
        ["POST /api/shipping/orders", "ShippingController.createOrder()", "shipping/controller/ShippingController.java", "Tạo vận đơn generic"],
        ["POST /api/hoa-don/{id}/tao-van-don-ghn", "HoaDonController.taoVanDonGhn()", "order/controller/HoaDonController.java", "Tạo theo HĐ"],
        ["POST /api/hoa-don/{id}/dong-bo-ghn", "HoaDonController.dongBoGhn()", "order/controller/HoaDonController.java", "Sync TT"],
        ["GET /api/hoa-don/ghn-trang-thai", "HoaDonController.ghnTrangThai()", "order/controller/HoaDonController.java", "Danh sách mã TT GHN"],
        ["POST /api/hoa-don/{id}/gia-lap-webhook-ghn", "HoaDonController.giaLapWebhookGhn()", "order/controller/HoaDonController.java", "Giả lập webhook (dev)"],
    ],
    "flow": [
        ["1 Fee", "ShippingService.calcFee → GhnClient POST /v2/shipping-order/fee", "Fallback 30k nếu lỗi", "Checkout gọi"],
        ["2 Tạo ĐH", "GhnOrderCreationService.taoVanDonNeuCan / taoVanDonTheoId", "POST /v2/shipping-order/create", "Lưu ma_van_don_ghn"],
        ["3 Sync", "GhnOrderSyncService + GhnTrackingService", "POST /v2/shipping-order/detail", "Map sang TrangThaiDonHang"],
    ],
    "models": [
        ["ShippingFeeRequest", "Request", "shipping/model/request/ShippingFeeRequest.java", "Tỉnh/xã, khối lượng để tính phí"],
        ["ShippingFeeResponse", "Response", "shipping/model/response/ShippingFeeResponse.java", "Phí + service GHN"],
        ["CreateShippingOrderRequest", "Request", "shipping/model/request/CreateShippingOrderRequest.java", "Body tạo vận đơn chiều đi"],
        ["CreateShippingOrderResponse", "Response", "shipping/model/response/CreateShippingOrderResponse.java", "Mã vận đơn GHN"],
        ["ReturnShippingOrderRequest", "Request", "shipping/model/request/ReturnShippingOrderRequest.java", "Vận đơn chiều về (sheet 14)"],
        ["GhnProvinceResponse", "Response", "shipping/model/response/GhnProvinceResponse.java", "Tỉnh"],
        ["GhnWardResponse", "Response", "shipping/model/response/GhnWardResponse.java", "Xã/phường"],
        ["GhnDistrictResponse", "Response", "shipping/model/response/GhnDistrictResponse.java", "Quận (legacy)"],
        ["GhnPickShiftResponse", "Response", "shipping/model/response/GhnPickShiftResponse.java", "Ca lấy hàng"],
        ["GhnProperties", "Config", "shipping/config/GhnProperties.java", "token, shop, from-ward, fallback-fee"],
        ["HoaDonGhnWebhookRequest", "Request", "order/model/request/HoaDonGhnWebhookRequest.java", "Body sync/giả lập"],
        ["GhnTrangThaiOptionResponse", "Response", "order/model/response/GhnTrangThaiOptionResponse.java", "Option TT GHN"],
    ],
    "tables": [
        ["hoa_don.phi_van_chuyen", "Phí", "V1", "Snapshot lúc đặt"],
        ["hoa_don.ma_van_don_ghn", "Mã vận đơn", "V1 GHN patch", ""],
        ["hoa_don.ghn_district_id / ghn_ward_code", "Mã địa chỉ", "V1", ""],
        ["hoa_don.ghn_province_name / ghn_ward_name", "Tên 2 cấp", "V3__ghn_address_names_on_hoa_don.sql", "Tạo vận đơn 2 cấp"],
    ],
    "backend": [
        ["Controller", "shipping/controller/ShippingController.java", "provinces/wards/fee/orders", ""],
        ["Service", "shipping/service/ShippingService.java", "calcFee, createOrder, createReturnOrder", ""],
        ["Client", "shipping/client/GhnClient.java", "RestTemplate bọc GHN", ""],
        ["Config", "shipping/config/GhnProperties.java", "ghn.*", ""],
        ["DTO", "shipping/model/request/ShippingFeeRequest.java", "", ""],
        ["DTO", "shipping/model/request/CreateShippingOrderRequest.java", "", ""],
        ["DTO", "shipping/model/request/ReturnShippingOrderRequest.java", "Trả hàng", ""],
        ["DTO", "shipping/model/response/ShippingFeeResponse.java", "", ""],
        ["DTO", "shipping/model/response/CreateShippingOrderResponse.java", "", ""],
        ["DTO", "shipping/model/response/GhnProvinceResponse.java", "", ""],
        ["DTO", "shipping/model/response/GhnWardResponse.java", "", ""],
        ["DTO", "shipping/model/response/GhnDistrictResponse.java", "", ""],
        ["DTO", "shipping/model/response/GhnPickShiftResponse.java", "", ""],
        ["Order API", "order/controller/HoaDonController.java", "tao-van-don / dong-bo / webhook", ""],
        ["Service", "order/service/GhnOrderCreationService.java", "taoVanDonNeuCan, taoVanDonTheoId", ""],
        ["Service", "order/service/GhnOrderSyncService.java", "dongBoTheoId, apDungTrangThaiGhn", ""],
        ["Service", "order/service/GhnTrackingService.java", "track, allStatusOptions", ""],
        ["Service", "order/service/OnlineCheckoutService.java", "Gọi calcFee lúc đặt", ""],
        ["Service", "order/service/HoaDonService.java", "Tạo vận đơn khi xác nhận", ""],
        ["Entity", "order/entity/HoaDon.java", "Cột GHN", ""],
        ["DTO", "order/model/request/HoaDonGhnWebhookRequest.java", "", ""],
        ["DTO", "order/model/response/GhnTrangThaiOptionResponse.java", "", ""],
        ["Config", "resources/application.properties", "ghn.token, shop-id, fallback-fee…", ""],
        ["SQL", "db/migration/V1__init_full.sql", "Cột GHN", ""],
        ["SQL", "db/migration/V3__ghn_address_names_on_hoa_don.sql", "Tên tỉnh/xã", ""],
    ],
    "frontend": [
        ["View", "views/storefront/DatHang.vue", "Gọi fee + lưu mã tỉnh/xã", ""],
        ["UI", "components/storefront/CheckoutRecipientModal.vue", "Chọn địa chỉ 2 cấp", ""],
        ["View", "views/admin/orders/HoaDonDetailPage.vue", "Tạo / sync vận đơn", ""],
        ["API", "api/shipping.js", "provinces, wards, fee", ""],
        ["API", "api/hoaDonApi.js", "taoVanDonGhn, dongBoGhn, giaLapWebhook", ""],
        ["CSS", "styles/checkoutRecipientModal.css", "", ""],
    ],
})

FEATURES.append({
    "sheet": "14. Tra hang hoan tien",
    "title": "14. RETURN / REFUND — Trả hàng và hoàn tiền",
    "subtitle": "Nằm trong package order (không có package return). VNPay hoàn qua VnpayRefundGateway; COD/CK hoàn tay + ảnh.",
    "package": "order (TraHang*/Refund*) + shipping return + payment refund",
    "note": "TrangThaiTraHang: CHO_DUYET → DA_DUYET → DANG_HOAN_HANG → DA_NHAN_HANG → HOAN_TAT / TU_CHOI. LoaiHangTra: TOT / LOI.",
    "fe_entry": [
        ["Khách tạo YC", "/tra-cuu-don", "views/storefront/TraCuuDon.vue", "Mở modal COD / ví"],
        ["Chi tiết YC khách", "/tra-cuu-don/tra-hang/:id", "views/storefront/DonTraHang.vue", "Timeline + tạo vận đơn trả"],
        ["Admin trả", "/admin/tra-hang", "views/admin/returns/ReturnList.vue", "Duyệt / từ chối / nhận hàng"],
        ["Admin hoàn", "/admin/hoan-tien", "views/admin/refunds/RefundList.vue", "Hoàn tất / từ chối hoàn"],
    ],
    "apis": [
        ["POST /api/online/orders/{idHoaDon}/tra-hang", "TraHangKhachController.taoYeuCau()", "order/controller/TraHangKhachController.java", "Tạo YC (multipart ảnh)"],
        ["GET /api/online/tra-hang", "TraHangKhachController.danhSachCuaToi()", "order/controller/TraHangKhachController.java", "List YC khách"],
        ["GET /api/online/tra-hang/{id}", "TraHangKhachController.chiTietCuaToi()", "order/controller/TraHangKhachController.java", "Chi tiết YC khách"],
        ["GET /api/online/tra-hang/ca-lay-hang", "TraHangKhachController.caLayHang()", "order/controller/TraHangKhachController.java", "Ca lấy GHN"],
        ["POST /api/online/tra-hang/{id}/tao-van-don", "TraHangKhachController.taoVanDonTra()", "order/controller/TraHangKhachController.java", "Vận đơn chiều về"],
        ["GET /api/tra-hang", "TraHangController.danhSach()", "order/controller/TraHangController.java", "List admin ?trangThai="],
        ["POST /api/tra-hang/{id}/duyet", "TraHangController.duyet()", "order/controller/TraHangController.java", "Duyệt"],
        ["POST /api/tra-hang/{id}/tu-choi", "TraHangController.tuChoi()", "order/controller/TraHangController.java", "Từ chối"],
        ["POST /api/tra-hang/{id}/da-nhan-hang", "TraHangController.daNhanHang()", "order/controller/TraHangController.java", "Nhận hàng + phân loại lô"],
        ["GET /api/tra-hang/{id}/lo-hang", "TraHangController.danhSachLo()", "order/controller/TraHangController.java", "Lô gắn YC"],
        ["POST /api/tra-hang/{id}/dong-bo-ghn", "TraHangController.dongBoGhn()", "order/controller/TraHangController.java", "Sync vận đơn trả"],
        ["GET /api/hoan-tien", "HoanTienController.danhSach()", "order/controller/HoanTienController.java", "List hoàn ?trangThai="],
        ["POST /api/hoan-tien/{id}/hoan-tat", "HoanTienController.hoanTatJson() / hoanTatMultipart()", "order/controller/HoanTienController.java", "JSON VNPay hoặc multipart chứng từ"],
        ["POST /api/hoan-tien/{id}/tu-choi", "HoanTienController.tuChoi()", "order/controller/HoanTienController.java", "Từ chối hoàn"],
    ],
    "flow": [
        ["1 Tạo YC", "ReturnRequestService.taoYeuCau", "CHO_DUYET + ảnh", ""],
        ["2 Duyệt", "ReturnRequestService.duyet → RefundService.taoHoanTienChoXuLy", "DA_DUYET + phiếu hoàn CHO_XU_LY", ""],
        ["3 Vận đơn về", "ShippingService.createReturnOrder", "DANG_HOAN_HANG", "Scheduler sync"],
        ["4 Nhận hàng", "xacNhanNhanHang + ChiTietTraHangLo TOT/LOI", "DA_NHAN_HANG, nhập lại lô TOT", ""],
        ["5 Hoàn tiền", "RefundService.hoanTat → VnpayRefundGateway hoặc hoàn tay + ảnh", "DA_HOAN", ""],
    ],
    "models": [
        ["YeuCauTraHang", "Entity", "order/entity/YeuCauTraHang.java", "YC trả hàng"],
        ["AnhYeuCauTraHang", "Entity", "order/entity/AnhYeuCauTraHang.java", "Ảnh minh chứng YC"],
        ["ChiTietTraHangLo", "Entity", "order/entity/ChiTietTraHangLo.java", "Phân loại lô TOT/LOI khi nhận"],
        ["HoanTien", "Entity", "order/entity/HoanTien.java", "Phiếu hoàn"],
        ["AnhHoanTien", "Entity", "order/entity/AnhHoanTien.java", "Ảnh chứng từ hoàn tay"],
        ["TaoYeuCauTraHangRequest", "Request", "order/model/request/TaoYeuCauTraHangRequest.java", "Lý do, dòng, PTTT hoàn"],
        ["TaoVanDonTraRequest", "Request", "order/model/request/TaoVanDonTraRequest.java", "Ca lấy + địa chỉ lấy"],
        ["DuyetTraHangRequest", "Request", "order/model/request/DuyetTraHangRequest.java", "Body duyệt"],
        ["NhanHangTraRequest", "Request", "order/model/request/NhanHangTraRequest.java", "Body nhận hàng + map lô"],
        ["HoanTatHoanTienRequest", "Request", "order/model/request/HoanTatHoanTienRequest.java", "Body hoàn tất"],
        ["YeuCauTraHangResponse", "Response", "order/model/response/YeuCauTraHangResponse.java", "JSON YC admin"],
        ["StorefrontReturnDetailResponse", "Response", "order/model/response/StorefrontReturnDetailResponse.java", "JSON YC khách"],
        ["StorefrontReturnTimelineStepResponse", "Response", "order/model/response/StorefrontReturnTimelineStepResponse.java", "Mốc timeline trả"],
        ["HoanTienResponse", "Response", "order/model/response/HoanTienResponse.java", "JSON phiếu hoàn"],
        ["LoHangDonHangResponse", "Response", "order/model/response/LoHangDonHangResponse.java", "Lô trên đơn khi nhận hàng"],
        ["TrangThaiTraHang", "Enum", "common/enums/TrangThaiTraHang.java", "CHO_DUYET…HOAN_TAT / TU_CHOI"],
        ["TrangThaiHoanTien", "Enum", "common/enums/TrangThaiHoanTien.java", "CHO_XU_LY / DA_HOAN / TU_CHOI"],
        ["LoaiHoanTien", "Enum", "common/enums/LoaiHoanTien.java", "HUY_DON / TRA_HANG"],
        ["LoaiHangTra", "Enum", "common/enums/LoaiHangTra.java", "TOT / LOI"],
        ["RefundGateway", "Interface", "payment/gateway/RefundGateway.java", "refund(RefundCommand)"],
        ["RefundCommand", "Gateway DTO", "payment/gateway/RefundCommand.java", "Lệnh hoàn cổng TT"],
        ["RefundResult", "Gateway DTO", "payment/gateway/RefundResult.java", "Kết quả hoàn"],
    ],
    "tables": [
        ["yeu_cau_tra_hang", "YC trả", "pick_shift, ghn_trang_thai_tra, ngay_nhan_hang", ""],
        ["anh_yeu_cau_tra_hang", "Ảnh YC", "", ""],
        ["chi_tiet_tra_hang_lo", "Phân loại lô", "TOT/LOI", ""],
        ["hoan_tien", "Phiếu hoàn", "phan_hoi_ncc", ""],
        ["anh_hoan_tien", "Chứng từ hoàn tay", "", ""],
        ["thanh_toan_hoa_don", "Cần refund VNPay", "provider_transaction_no, provider_pay_date", ""],
    ],
    "backend": [
        ["Controller", "order/controller/TraHangKhachController.java", "API khách", ""],
        ["Controller", "order/controller/TraHangController.java", "API admin trả", ""],
        ["Controller", "order/controller/HoanTienController.java", "API hoàn tiền", ""],
        ["Service", "order/service/ReturnRequestService.java", "taoYeuCau, duyet, tuChoi, taoVanDonTra, xacNhanNhanHang", "Lõi trả"],
        ["Service", "order/service/RefundService.java", "taoHoanTienChoXuLy, hoanTat, tuChoi", "Lõi hoàn"],
        ["Scheduler", "order/service/ReturnShipmentSyncScheduler.java", "Đồng bộ vận đơn hoàn", "return.shipment.sync-ms"],
        ["Repo", "order/repository/YeuCauTraHangRepository.java", "", ""],
        ["Repo", "order/repository/AnhYeuCauTraHangRepository.java", "", ""],
        ["Repo", "order/repository/ChiTietTraHangLoRepository.java", "", ""],
        ["Repo", "order/repository/HoanTienRepository.java", "", ""],
        ["Repo", "order/repository/AnhHoanTienRepository.java", "", ""],
        ["Entity", "order/entity/YeuCauTraHang.java", "", ""],
        ["Entity", "order/entity/AnhYeuCauTraHang.java", "", ""],
        ["Entity", "order/entity/ChiTietTraHangLo.java", "", ""],
        ["Entity", "order/entity/HoanTien.java", "", ""],
        ["Entity", "order/entity/AnhHoanTien.java", "", ""],
        ["DTO", "order/model/request/TaoYeuCauTraHangRequest.java", "", ""],
        ["DTO", "order/model/request/TaoVanDonTraRequest.java", "", ""],
        ["DTO", "order/model/request/DuyetTraHangRequest.java", "", ""],
        ["DTO", "order/model/request/NhanHangTraRequest.java", "", ""],
        ["DTO", "order/model/request/HoanTatHoanTienRequest.java", "", ""],
        ["DTO", "order/model/request/HoaDonTuChoiRequest.java", "Dùng chung từ chối", ""],
        ["DTO", "order/model/response/YeuCauTraHangResponse.java", "", ""],
        ["DTO", "order/model/response/StorefrontReturnDetailResponse.java", "", ""],
        ["DTO", "order/model/response/StorefrontReturnTimelineStepResponse.java", "", ""],
        ["DTO", "order/model/response/HoanTienResponse.java", "", ""],
        ["DTO", "order/model/response/LoHangDonHangResponse.java", "", ""],
        ["Enum", "common/enums/TrangThaiTraHang.java", "", ""],
        ["Enum", "common/enums/TrangThaiHoanTien.java", "", ""],
        ["Enum", "common/enums/LoaiHoanTien.java", "", ""],
        ["Enum", "common/enums/LoaiHangTra.java", "", ""],
        ["GHN", "shipping/service/ShippingService.java", "createReturnOrder", ""],
        ["GHN", "shipping/model/request/ReturnShippingOrderRequest.java", "", ""],
        ["GHN", "shipping/model/response/GhnPickShiftResponse.java", "Ca lấy", ""],
        ["Refund", "payment/gateway/RefundGateway.java", "Interface", ""],
        ["Refund", "payment/gateway/RefundGatewayRegistry.java", "Chọn cổng", ""],
        ["Refund", "payment/gateway/RefundCommand.java", "", ""],
        ["Refund", "payment/gateway/RefundResult.java", "", ""],
        ["Refund", "payment/vnpay/VnpayRefundGateway.java", "vnp_Command=refund", ""],
        ["Noti", "notification/service/ThongBaoService.java", "YEU_CAU_TRA_HANG, HOAN_TIEN_*", ""],
        ["Mail", "notification/service/OrderMailService.java", "Mail duyệt/từ chối/hoàn xong", ""],
        ["Enum TB", "notification/enums/LoaiThongBao.java", "Các loại trả/hoàn", ""],
        ["SQL", "db/migration/V1__init_full.sql", "5 bảng trả/hoàn", ""],
    ],
    "frontend": [
        ["Router", "router/index.js", "tra-cuu-don, tra-hang/:id, admin/tra-hang, admin/hoan-tien", ""],
        ["Menu", "constants/adminMenu.js", "Trả hàng, Hoàn tiền", ""],
        ["View", "views/storefront/TraCuuDon.vue", "Mở tạo YC", ""],
        ["View", "views/storefront/DonTraHang.vue", "Chi tiết YC khách", ""],
        ["View", "views/admin/returns/ReturnList.vue", "Admin trả", ""],
        ["View", "views/admin/refunds/RefundList.vue", "Admin hoàn", ""],
        ["UI", "components/storefront/ReturnRequestCodModal.vue", "YC hoàn COD", ""],
        ["UI", "components/storefront/ReturnRequestWalletModal.vue", "YC hoàn ví/VNPay", ""],
        ["UI", "components/storefront/ReturnPickShiftModal.vue", "Chọn ca lấy", ""],
        ["UI", "components/storefront/ReturnDetailCard.vue", "Card chi tiết", ""],
        ["UI", "components/storefront/OrderCard.vue", "Link trả hàng", ""],
        ["Composable", "composables/useReturnRequest.js", "Logic tạo YC", ""],
        ["Composable", "composables/useAdminBadges.js", "Badge trả/hoàn", ""],
        ["API", "api/traHangApi.js", "", ""],
        ["API", "api/hoanTienApi.js", "", ""],
        ["API", "api/shipping.js", "Ca lấy", ""],
        ["Util", "utils/returnStatus.js", "Label TT trả", ""],
        ["CSS", "styles/returnModal.css", "", ""],
    ],
})


def write_index(wb):
    ws = wb.active
    ws.title = "00. Muc luc"
    cols = 4
    set_widths(ws, [10, 42, 55, 70])
    r = 1
    r = merge_title(ws, cols, r, "SUNOVA — BẢN ĐỒ CHỨC NĂNG BẢO VỆ ĐỒ ÁN", NAVY, font_index_title, 34)
    r = merge_title(ws, cols, r, "Mỗi chức năng = 1 sheet. Cột cố định: Entry FE → API → Luồng → Model → Bảng DB → Backend files → Frontend files.", TEAL, font_sub, 24)
    r = note_row(ws, cols, r, "Gốc backend: backend/src/main/java/org/example/templatejava6/     |     Gốc FE: frontend/src/     |     SQL: backend/src/main/resources/db/migration/     |     Click tên sheet bên dưới hoặc tab Excel để mở.")
    r = blank(ws, r)
    r = section(ws, cols, r, "DANH SÁCH CHỨC NĂNG (bấm tên sheet ở thanh dưới cửa sổ Excel)", TEAL)
    r = headers(ws, r, ["STT", "Chức năng trên giấy", "Sheet Excel", "Package / điểm nhớ vấn đáp"])
    rows = [
        ["01", "Cart — Quản lý sản phẩm trong giỏ hàng", "01. Gio hang", "cart — guest dùng localStorage, login mới gọi API"],
        ["02", "Checkout — Tạo đơn hàng từ giỏ hàng", "02. Tao don hang", "order.OnlineCheckoutService — không package checkout"],
        ["03", "Checkout — Thanh toán", "03. Thanh toan", "payment — chỉ VNPay; MoMo seed nhưng không gateway"],
        ["04", "Order — Theo dõi & cập nhật trạng thái", "04. Trang thai don", "PATCH status + gán lô + GHN + hủy khách"],
        ["05", "Order — Tìm kiếm & lọc đơn hàng", "05. Tim kiem loc don", "GET /api/hoa-don/search"],
        ["06", "Order — Xem chi tiết đơn hàng", "06. Chi tiet don", "Admin id / khách cua-toi / public token"],
        ["07", "Order — In hóa đơn bán hàng", "07. In hoa don", "window.print — không API PDF"],
        ["08", "Promotion — Chương trình khuyến mãi", "08. Dot giam gia", "DotGiamGia /api/sale — không entity ChuongTrinh"],
        ["09", "Promotion — Voucher / coupon", "09. Voucher", "PhieuGiamGia CONG_KHAI/CA_NHAN"],
        ["10", "Notification — Đặt hàng thành công", "10. TB dat hang", "In-app ADMIN DON_HANG_MOI, không phải khách"],
        ["11", "Notification — Cập nhật trạng thái đơn", "11. TB trang thai", "In-app KHÁCH DON_HANG_CAP_NHAT"],
        ["12", "Notification — Email xác nhận đơn", "12. Email xac nhan", "Event DatHangThanhCongMailEvent — không REST"],
        ["13", "Shipping — Tính phí & vận đơn GHN", "13. GHN Shipping", "shipping + GhnOrderCreationService"],
        ["14", "Return / Refund — Trả hàng hoàn tiền", "14. Tra hang hoan tien", "order.ReturnRequestService + VnpayRefundGateway"],
    ]
    start = r
    for row in rows:
        r = data_row(ws, r, row)
        # hyperlink to sheet
        cell = ws.cell(r - 1, 3)
        target = row[2]
        cell.hyperlink = f"#'{target}'!A1"
        cell.font = font_link
    r = blank(ws, r)
    r = section(ws, cols, r, "FILE DÙNG CHUNG NHIỀU CHỨC NĂNG", GOLD)
    r = headers(ws, r, ["Lớp", "File", "Vai trò", "Ghi chú"])
    for row in [
        ["Security", "common/config/SecurityConfig.java", "Ai được gọi API nào", "guest checkout, callback VNPay, shipping fee permitAll"],
        ["Exception", "common/exception/ApiException.java", "Lỗi nghiệp vụ trả FE", ""],
        ["Axios", "frontend/src/api/request.js", "JWT + prefix API", ""],
        ["Router", "frontend/src/router/index.js", "Toàn bộ route", ""],
        ["Config", "backend/src/main/resources/application.properties", "VNPay, GHN, SMTP, timeout đơn", ""],
    ]:
        r = data_row(ws, r, row)
    finish_sheet(ws, cols)
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    return ws


def main():
    wb = Workbook()
    write_index(wb)
    for feat in FEATURES:
        ws = wb.create_sheet(feat["sheet"])
        write_feature(ws, feat)

    # tab colors
    colors = ["0F766E", "1B3A4B", "0F766E", "1D4ED8", "0369A1", "0284C7", "0EA5E9", "B45309",
              "7C3AED", "BE185D", "CA8A04", "EA580C", "0F766E", "7C2D12", "9F1239"]
    wb.worksheets[0].sheet_properties.tabColor = colors[0]
    for i, ws in enumerate(wb.worksheets[1:], 1):
        ws.sheet_properties.tabColor = colors[i % len(colors)]

    wb.save(OUT)
    print("WROTE", OUT)


if __name__ == "__main__":
    main()
