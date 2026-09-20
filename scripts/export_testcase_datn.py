# -*- coding: utf-8 -*-
"""Ghi test case từ testcase_moi.xlsx sang PhamMinhQuang_Testcase_DATN.xlsx theo mẫu sheet."""
import math
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "testcase_moi.xlsx"
DST = ROOT / "PhamMinhQuang_Testcase_DATN.xlsx"

# sheet nguồn: (tiêu đề, prefix ID hợp lệ)
SHEETS = [
    ("Đợt giảm giá", "ĐỢT GIẢM GIÁ (SALE)", "SALE"),
    ("Phiếu giảm giá", "PHIẾU GIẢM GIÁ (VOUCHER)", "VC"),
    ("Tạo đơn từ giỏ hàng", "TẠO ĐƠN HÀNG TỪ GIỎ HÀNG (CHECKOUT ONLINE)", "OC"),
    ("Thanh toán đơn hàng", "THANH TOÁN ĐƠN HÀNG", "PAY"),
    ("Cập nhật trạng thái đơn hàng", "CẬP NHẬT TRẠNG THÁI ĐƠN HÀNG", "TC"),
    ("Thông báo đơn hàng", "THÔNG BÁO ĐƠN HÀNG", "TB"),
    ("Lọc tìm kiếm đơn hàng", "LỌC / TÌM KIẾM ĐƠN HÀNG", "SR"),
    ("Trả hàng hoàn tiền", "TRẢ HÀNG & HOÀN TIỀN", "RT"),
]

HEADERS = [
    "ID",
    "Kỹ thuật áp dụng",
    "Tiêu đề case",
    "Điều kiện tiên quyết",
    "Dữ liệu đầu vào",
    "Các bước thực hiện",
    "Kết quả mong đợi",
    "Kết quả thực tế",
    "Ưu tiên",
    "Loại",
    "Phân loại",
    "Dữ liệu mẫu API",
]

# A–L; H = Kết quả thực tế (trống để tester điền)
COL_WIDTHS = [14, 28, 34, 28, 32, 38, 44, 28, 13, 16, 16, 72]
CENTER_COLS = {1, 9, 10, 11}

THIN = Border(
    left=Side(style="thin", color="8EA9DB"),
    right=Side(style="thin", color="8EA9DB"),
    top=Side(style="thin", color="8EA9DB"),
    bottom=Side(style="thin", color="8EA9DB"),
)
TITLE_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
ACTUAL_FILL = PatternFill("solid", fgColor="FFF2CC")
ROW_FILL_A = PatternFill("solid", fgColor="FFFFFF")
ROW_FILL_B = PatternFill("solid", fgColor="D6E3F8")

TITLE_FONT = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
HEADER_FONT = Font(name="Times New Roman", size=12, bold=True, color="FFFFFF")
DATA_FONT = Font(name="Times New Roman", size=11)
ID_FONT = Font(name="Times New Roman", size=11, bold=True)

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def find_header_row(ws):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip() == "ID":
            return r
    return None


def read_cases(ws, prefix: str):
    """Chỉ lấy bảng test case chính: ID đúng prefix và có tiêu đề. Dừng khi ra khỏi bảng."""
    header_row = find_header_row(ws)
    if not header_row:
        raise SystemExit(f"Không thấy dòng header ID trên sheet {ws.title}")

    pat = re.compile(rf"^{re.escape(prefix)}\d{{2,3}}$")
    cases = []
    started = False
    empty_run = 0

    for r in range(header_row + 1, ws.max_row + 1):
        raw = ws.cell(r, 1).value
        tid = str(raw).strip() if raw is not None else ""
        title = ws.cell(r, 3).value

        if not tid:
            empty_run += 1
            if started and empty_run >= 2:
                break
            continue

        empty_run = 0
        if not pat.match(tid):
            if started:
                break
            continue

        if not title:
            continue

        started = True
        cases.append([ws.cell(r, c).value for c in range(1, 12)])

    return cases


def safe_value(val):
    """Tránh Excel hiểu '-', '@...', '->...' là công thức."""
    if val is None:
        return None
    if not isinstance(val, str):
        return val
    text = val.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return None
    if text[0] in "=+-@":
        return "'" + text
    return text


def estimate_height(values):
    lines = 1
    for idx, val in enumerate(values):
        if val is None:
            continue
        text = str(val)
        width = COL_WIDTHS[idx]
        chars_per_line = max(int(width * 1.05), 8)
        n = 0
        for part in text.split("\n"):
            n += max(1, math.ceil(len(part) / chars_per_line))
        lines = max(lines, n)
    return min(max(36, 16 + lines * 14), 180)


def apply_data_cell(cell, value, *, center=False, fill=None, font=None):
    cell.value = safe_value(value)
    cell.font = font or DATA_FONT
    cell.alignment = CENTER if center else LEFT
    cell.border = THIN
    cell.number_format = "@"
    if fill is not None:
        cell.fill = fill


def write_sheet(wb_out, sheet_name: str, title: str, cases):
    ws = wb_out.create_sheet(title=sheet_name[:31])

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    title_cell = ws.cell(1, 1, title)
    title_cell.font = TITLE_FONT
    title_cell.fill = TITLE_FILL
    title_cell.alignment = CENTER
    for c in range(1, 13):
        cell = ws.cell(1, c)
        cell.fill = TITLE_FILL
        cell.border = THIN
        cell.alignment = CENTER
    ws.row_dimensions[1].height = 32

    for c, header in enumerate(HEADERS, 1):
        cell = ws.cell(2, c, header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN
        cell.number_format = "@"
    ws.row_dimensions[2].height = 28

    for i, src in enumerate(cases):
        r = 3 + i
        dest = list(src[:7]) + [None] + list(src[7:11])
        row_fill = ROW_FILL_A if i % 2 == 0 else ROW_FILL_B
        for c, val in enumerate(dest, 1):
            fill = ACTUAL_FILL if c == 8 else row_fill
            font = ID_FONT if c == 1 else DATA_FONT
            apply_data_cell(ws.cell(r, c), val, center=c in CENTER_COLS, fill=fill, font=font)
        ws.row_dimensions[r].height = estimate_height(dest)

    for i, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    last = 2 + len(cases)
    ws.auto_filter.ref = f"A2:L{last}"
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"

    dv = DataValidation(
        type="list",
        formula1='"Pass,Fail,Blocked,Skip"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=False,
    )
    dv.prompt = "Điền Pass / Fail / Blocked / Skip hoặc mô tả kết quả"
    dv.promptTitle = "Kết quả thực tế"
    if len(cases) > 0:
        dv.add(f"H3:H{last}")
        ws.add_data_validation(dv)

    return len(cases)


def main():
    src_wb = load_workbook(SRC, data_only=False)
    out = Workbook()
    out.remove(out.active)

    summary = []
    for sheet_name, title, prefix in SHEETS:
        if sheet_name not in src_wb.sheetnames:
            raise SystemExit(f"Thiếu sheet nguồn: {sheet_name}")
        cases = read_cases(src_wb[sheet_name], prefix)
        if not cases:
            raise SystemExit(f"Không lấy được case nào từ {sheet_name} (prefix {prefix})")
        n = write_sheet(out, sheet_name, title, cases)
        summary.append((sheet_name, n, cases[0][0], cases[-1][0]))
        print(f"OK  {sheet_name}: {n}  ({cases[0][0]} → {cases[-1][0]})")

    out.save(DST)
    print(f"\nSaved -> {DST}")
    print(f"Total: {sum(n for _, n, __, ___ in summary)} cases / {len(summary)} sheets")


if __name__ == "__main__":
    main()
